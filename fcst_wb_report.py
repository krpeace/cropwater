#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fcst_wb_report.py — 02-Cycle 4단계(G4) 엑셀: 예보 물수지 검증 엑셀과 서비스(관수 전망) 엑셀. 라이브 수식

[검증 엑셀] build_wbverify_workbook(res, out)
  시트: 요약 / 발표별 / 발표요약 / 관측물수지 / 오차표 / 설정 / 방법
  - 값: 예보 일 입력(예보 ETc = Kc × 예보 ETo, 강수·기대 강수·강수확률), 관측 ETc·강수, 오차표(다른 해), 제외 사유
  - 수식: 관측 물수지(관수 규칙 시나리오 스위치 포함), 발표별 8개 경로의 고갈량, 관수 필요 판정, 발표별 예상일·판정·범위 적중, 요약 지표
  - 설정 시트의 토양 파라미터·판정 기준 고갈량·평가 대상일 수를 바꾸면 다시 계산된다
[서비스 엑셀] build_service_workbook(sv, out)
  시트: 관수 전망 / 예보 물수지 / 관측 물수지 / 날짜별 / 오차표 / 편향 점검 / 설정 / 방법
  - 관수 전망의 조회일(노란 칸)을 바꾸면 그 날짜의 토양 상태·앞으로 3일을 '날짜별' 시트에서 찾아온다(없는 날짜는 '자료 없음')
  - 상태에 과습 경고·과습 주의(관측 물수지 기준, 예보는 쓰지 않음)
  - 관측 물수지의 관수량(노란 칸)을 적으면 어제 끝 Dr과 예보 물수지·관수 필요 예상일이 바로 바뀐다

근거: docs/THEORY.md 6·9장, docs/VALIDATION.md G4, 해석: docs/RESULTS_GUIDE.md
"""
import datetime as dt

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

from fao56_core import safe_save
from fcst_report import BLUE, BORDER, BROWN, C, DATE, DTM, F2, FONT, GREEN, H, LIGHT, T, YEL, _wrap_text, widths
from fcst_wb import MAIN_DAYS, NEED_CATS, PATHS, WET_CAUTION_RATIO

GREY = "8A8F87"; REF_FILL = "F2F2F2"; NEED_FILL = "F4C7B8"; WARN_FILL = "FBEED2"; OK_FILL = "D8EAD3"
F1 = "0.0"
WETW_FILL = "9DC3E6"; WETC_FILL = "DDEBF7"      # 과습 경고·과습 주의 바탕
TL = "날짜별"                                   # 조회용 시트(관측 물수지 + 예보 물수지를 날짜순으로 이음)
SEL = "'관수 전망'!$B$5"                        # 조회일 칸


def _v(x):
    """NaN → None(빈 칸)"""
    if x is None:
        return None
    try:
        if pd.isna(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating,)):
        return float(x)
    if isinstance(x, pd.Timestamp):
        return x.to_pydatetime()
    return x


def _sec(ws, r, text, ncol=8, fill=GREEN):
    H(ws, r, 1, text, fill, size=11)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    ws.cell(r, 1).alignment = Alignment("left", "center")


def _note(ws, r, text, ncol=12, width=150, color="555555"):
    """긴 설명을 여러 행으로 나눠 쓴다(병합 칸 줄바꿈이 프로그램마다 달라서)"""
    for ln in _wrap_text(text, width):
        T(ws, r, 1, ln, size=9, color=color)
        r += 1
    return r


# ── 공통: 설정 시트 ──────────────────────────────────────────────────────
def _settings(wb, soil, extra, W, title):
    """토양·판정 설정(노란 칸) → W[이름] = '설정!$B$n'"""
    ws = wb.create_sheet("설정")
    H(ws, 1, 1, "항목"); H(ws, 1, 2, "값 (노란 칸은 바꿀 수 있음)"); H(ws, 1, 3, "근거·설명")
    T(ws, 1, 5, title, bold=True, color=GREEN)
    rows = [
        ("FC", "포장용수량 θFC (m³/m³)", soil["fc"], "0.00", True, "01-Cycle 워크북 설정 값. FAO-56 Table 19"),
        ("WP", "위조점 θWP (m³/m³)", soil["wp"], "0.00", True, "01-Cycle 워크북 설정 값"),
        ("ZR", "근권심도 Zr (m)", soil["zr"], "0.00", True, "FAO-56 Table 22"),
        ("P", "토양수분고갈계수 p", soil["p"], "0.00", True, "FAO-56 Table 22"),
        ("TAW", "★ TAW (mm)", None, F1, False, "총유효수분 = 1000·(θFC−θWP)·Zr [식82]"),
        ("RAW", "★ RAW (mm)", None, F1, False, "쉽게이용가능수분 = p·TAW [식83]. Ks가 1보다 작아지기 시작하는 고갈량"),
        ("DR0", "초기 고갈량 Dr₀ (mm)", soil["dr0"], F1, True, "관측 물수지 첫날의 전날 끝 고갈량"),
        ("EA", "관수효율 Ea", soil["ea"], "0.00", True, "관수 기록(공급량) × Ea = 근권에 들어간 순관수량"),
    ] + extra
    r = 2
    for key, label, val, fmt, edit, note in rows:
        H(ws, r, 1, label, LIGHT, white=False)
        if key == "TAW":
            val = f"=1000*({W['FC']}-{W['WP']})*{W['ZR']}"
        elif key == "RAW":
            val = f"={W['P']}*{W['TAW']}"
        elif isinstance(val, str) and val.startswith("=RAW"):
            val = f"={W['RAW']}{val[4:]}"                # '=RAW' 또는 '=RAW*0.1' 같은 식
        C(ws, r, 2, val, fmt, fill=(YEL if edit else None))
        x = ws.cell(r, 3, note); x.font = Font(name=FONT, size=9, color="555555"); x.border = BORDER
        x.alignment = Alignment("left", "center", wrap_text=True)
        W[key] = f"설정!$B${r}"
        ws.row_dimensions[r].height = 30
        r += 1
    widths(ws, {"A": 30, "B": 22, "C": 90})
    return ws, r


def _wb_formula(prev, p, etc, W, irr=None):
    """하루 물수지 한 칸 수식: Dr = MIN(MAX(Dr,i-1 − P [− I] + Ks(Dr,i-1)·ETc, 0), TAW)
       (DP = max(P + I − ETc_adj − Dr,i-1, 0)을 넣고 정리한 식 — fao56_core.wb_step과 같은 값)"""
    ks = f"IF({prev}<={W['RAW']},1,({W['TAW']}-{prev})/({W['TAW']}-{W['RAW']}))"
    pp = f"-{p}" if p not in (None, "0") else ""
    ii = f"-{irr}" if irr else ""
    return f"=MIN(MAX({prev}{pp}{ii}+{ks}*{etc},0),{W['TAW']})"


# ── 관측 물수지 시트 (검증·서비스 공용) ─────────────────────────────────────
OWB_HEAD = ["일자", "Kc", "ETo (mm)", "ETo 출처", "ETc (mm)", "강수 P (mm)", "관수 기록\n(공급, mm)", "순관수 I\n(mm)",
            "Dr,i-1 (mm)", "Ks", "ETc_adj (mm)", "DP (mm)", "Dr,i (mm)", "관수\n필요", "관측\n결측", "결측 영향"]


def sheet_owb(wb, owb, W, name="관측물수지", scenario_cell=None, note=None):
    """관측 물수지(라이브 수식). 행: owb 날짜. 반환: (시트, 날짜 열 범위, Dr 열 범위, 마지막 행)"""
    ws = wb.create_sheet(name)
    for c, h in enumerate(OWB_HEAD, 1):
        H(ws, 1, c, h, BROWN if c in (7, 8) else GREEN)
    n = len(owb)
    for i, r in enumerate(owb.itertuples(index=False), 2):
        C(ws, i, 1, _v(r.date), DATE)
        C(ws, i, 2, _v(r.Kc), "0.000")
        C(ws, i, 3, _v(r.ETo), F2, fill=(WARN_FILL if r.ETo_src != "관측" else None))
        C(ws, i, 4, r.ETo_src)
        C(ws, i, 5, f"=IF(C{i}=\"\",0,B{i}*C{i})", F2)
        C(ws, i, 6, _v(r.P), F1)
        C(ws, i, 7, _irr_log_value(r, W), F1, fill=YEL)
        prev = W["DR0"] if i == 2 else f"M{i - 1}"
        auto = f"+IF({scenario_cell}=1,IF(I{i}>={W['RAW']},I{i},0),0)" if scenario_cell else ""
        C(ws, i, 8, f"=G{i}*{W['EA']}{auto}", F1)
        C(ws, i, 9, f"={prev}", F1)
        C(ws, i, 10, f"=IF(I{i}<={W['RAW']},1,({W['TAW']}-I{i})/({W['TAW']}-{W['RAW']}))", "0.000")
        C(ws, i, 11, f"=J{i}*E{i}", F2)
        C(ws, i, 12, f"=MAX(F{i}+H{i}-K{i}-I{i},0)", F1)
        C(ws, i, 13, f"=MIN(MAX(I{i}-F{i}-H{i}+K{i}+L{i},0),{W['TAW']})", F1)
        C(ws, i, 14, f'=IF(M{i}>={W["RAW"]},"●","")')
        C(ws, i, 15, "채움" if r.filled else "")
        C(ws, i, 16, "영향" if r.tainted else "")
    ws.freeze_panes = "B2"
    widths(ws, {"A": 11, "B": 7, "C": 8, "D": 8, "E": 8, "F": 8, "G": 10, "H": 9, "I": 9, "J": 7, "K": 9, "L": 8,
                "M": 9, "N": 6, "O": 6, "P": 7})
    if note:
        T(ws, 1, 18, note, size=9, color="555555")
    last = n + 1
    return ws, f"{name}!$A$2:$A${last}", f"{name}!$M$2:$M${last}", last


def _irr_log_value(r, W):
    """관수 기록 칸 값: 관수 규칙 시나리오의 자동 관수는 수식이 더하므로, 기록(공급량)만 넣는다"""
    return float(W.get("_irr_log", {}).get(pd.Timestamp(r.date), 0.0))


# ── 검증 엑셀 ────────────────────────────────────────────────────────────
RUN_HEAD = ["키", "발표ID", "구분", "발표시각", "순번", "대상일", "월", "주/참고", "유효", "제외 사유", "Kc",
            "예보 ETc", "ETc 상대오차 r", "예보 강수", "기대 강수", "강수확률(최대)", "관측 ETc", "관측 강수", "기준선 ETc",
            "출발 Dr\n(D−1, 관측)"]
PATH_ORDER = ["center", "early", "late", "fcst", "none", "obs", "pers", "true"]
PATH_HEAD = {"center": "Dr 중심\n(기대 강수)", "early": "Dr 빠르면\n(비 없음·ETc+)", "late": "Dr 늦으면\n(예보 강수·ETc−)",
             "fcst": "Dr 예보 강수\n그대로", "none": "Dr 비 무시", "obs": "Dr 관측 강수", "pers": "Dr 기준선",
             "true": "Dr 참값\n(관측)"}
NEED_HEAD = [("center", "필요\n중심"), ("early", "필요\n빠르면"), ("late", "필요\n늦으면"), ("true", "필요\n참값")]


def _path_inputs(path, r):
    """경로별 (강수 셀, ETc 셀) — r은 행 번호. 열: L 예보 ETc, M r, N 예보 강수, O 기대 강수, Q 관측 ETc, R 관측 강수, S 기준선 ETc"""
    return {"center": (f"O{r}", f"L{r}"), "early": (None, f"L{r}*(1+M{r})"), "late": (f"N{r}", f"L{r}*MAX(1-M{r},0)"),
            "fcst": (f"N{r}", f"L{r}"), "none": (None, f"L{r}"), "obs": (f"R{r}", f"L{r}"), "pers": (None, f"S{r}"),
            "true": (f"R{r}", f"Q{r}")}[path]


def sheet_runs(wb, res, W, owb_dates, owb_dr):
    """발표별: 발표마다 (저녁은 그날 D 행 + ) 대상일 행. 경로별 고갈량은 수식 사슬"""
    ws = wb.create_sheet("발표별")
    heads = RUN_HEAD + [PATH_HEAD[k] for k in PATH_ORDER] + [h for _, h in NEED_HEAD]
    for c, h in enumerate(heads, 1):
        fill = GREEN if c <= 10 else (BLUE if c <= 19 else (BROWN if c <= 20 + len(PATH_ORDER) else GREEN))
        H(ws, 1, c, h, fill)
    col = {k: CL(21 + i) for i, k in enumerate(PATH_ORDER)}
    ncol = {k: CL(21 + len(PATH_ORDER) + i) for i, (k, _) in enumerate(NEED_HEAD)}
    runs = res["runs"]
    r = 2
    ids, blocks = {}, []
    for rid, ((rn, run), g) in enumerate(runs.groupby(["run_name", "run"], sort=False), 1):
        g = g.sort_values("order")
        ids[(rn, run)] = rid
        first = r
        rows = []
        # 발표 전체가 빠진 경우(출발 관측 없음·예보 ETc 없음·그날 아침 예보 없음): 입력만 적고 경로 수식은 쓰지 않음
        dead = g.Dr_center.isna().all()
        if rn == "저녁" and g.pre_etc.notna().iloc[0]:
            x = g.iloc[0]
            rows.append(dict(order=0, target=pd.Timestamp(run).normalize(), month=pd.Timestamp(run).month, kind="그날", ok=0,
                             reason="그날(D) — 아침 발표 D+0 예보로 진행(채점 안 함)", Kc=None, ETc=x.pre_etc, rel=x.pre_rel,
                             rain=x.pre_rain, rain_exp=x.pre_rain_exp, pop=None, ETc_t=x.pre_etc_t, rain_t=x.pre_rain_t,
                             ETc_p=x.pre_etc_p))
        for x in g.itertuples():
            rows.append(dict(order=int(x.order), target=x.target, month=int(x.month), kind="주" if x.main else "참고",
                             ok=int(x.ok), reason=x.reason, Kc=x.Kc, ETc=x.ETc_fcst, rel=x.rel, rain=x.rain_fcst,
                             rain_exp=x.rain_exp, pop=x.pop_max, ETc_t=x.ETc_true, rain_t=x.rain_true, ETc_p=x.ETc_pers))
        for j, z in enumerate(rows):
            vals = [rid * 10 + z["order"], rid, rn, _v(run), z["order"], _v(z["target"]), z["month"], z["kind"], z["ok"],
                    z["reason"] or "", _v(z["Kc"]), _v(z["ETc"]), _v(z["rel"]), _v(z["rain"]), _v(z["rain_exp"]),
                    _v(z["pop"]), _v(z["ETc_t"]), _v(z["rain_t"]), _v(z["ETc_p"])]
            fmts = [None, None, None, DTM, None, DATE, None, None, None, None, "0.000", F2, "0.000", F1, F1, "0%", F2, F1, F2]
            grey = z["kind"] in ("참고", "그날")
            for c, (v, f) in enumerate(zip(vals, fmts), 1):
                C(ws, r, c, v, f, color=(GREY if grey else "1A1A1A"), fill=(REF_FILL if z["kind"] == "그날" else None))
            C(ws, r, 20, f"=IFERROR(INDEX({owb_dr},MATCH(INT(D{r})-1,{owb_dates},0)),\"\")", F1)
            if not dead:
                for k in PATH_ORDER:
                    prev = f"T{r}" if j == 0 else f"{col[k]}{r - 1}"
                    p, e = _path_inputs(k, r)
                    C(ws, r, 21 + PATH_ORDER.index(k), _wb_formula(prev, p, e, W), F1, bold=(k in ("center", "true")))
                for k, _ in NEED_HEAD:
                    C(ws, r, 21 + len(PATH_ORDER) + [a for a, _ in NEED_HEAD].index(k),
                      f"=IF(E{r}=0,\"\",IF({col[k]}{r}>={W['TH']},1,0))")
            r += 1
        blocks.append((rid, rn, run, first, r - 1, rn == "저녁" and rows[0]["order"] == 0, dead))
    last = r - 1
    ws.freeze_panes = "F2"
    widths(ws, {"A": 7, "B": 6, "C": 5, "D": 15, "E": 5, "F": 11, "G": 4, "H": 6, "I": 5, "J": 20, "K": 6, "L": 7, "M": 7,
                "N": 7, "O": 7, "P": 7, "Q": 7, "R": 7, "S": 7, "T": 8})
    for i in range(21, 21 + len(PATH_ORDER) + len(NEED_HEAD)):
        ws.column_dimensions[CL(i)].width = 9
    return ws, last, col, ncol, blocks


def sheet_runsum(wb, res, W, last, col, ncol, blocks):
    """발표요약: 발표마다 출발 고갈량, 관수 필요 예상일(중심·빠르면·늦으면·참값 순번), 판정, 범위 적중"""
    ws = wb.create_sheet("발표요약")
    heads = ["발표ID", "구분", "발표시각", "유효 대상일 수\n(평가 기간)", "출발 Dr\n중심", "출발 Dr\n참값", "평가 대상",
             "예상일 중심\n(순번)", "빠르면", "늦으면", "참값", "판정", "범위 적중\n(빠르면~늦으면)", "사건\n(누구라도 필요)",
             "참값이 '빠르면'\n보다 앞"]
    for c, h in enumerate(heads, 1):
        H(ws, 1, c, h, GREEN if c <= 7 else BLUE)
    R = lambda c: f"발표별!${c}$2:${c}${last}"
    days = W["DAYS"]
    for i, (rid, rn, run, first, lastrow, has_pre, dead) in enumerate(blocks, 2):
        C(ws, i, 1, rid); C(ws, i, 2, rn); C(ws, i, 3, _v(run), DTM)
        crit = f"{R('B')},A{i},{R('I')},1,{R('E')},\">=1\",{R('E')},\"<=\"&{days}"
        C(ws, i, 4, f"=COUNTIFS({crit})")
        if dead:
            C(ws, i, 7, 0)
            continue
        if has_pre:
            C(ws, i, 5, f"=INDEX({R(col['center'])},MATCH(A{i}*10,{R('A')},0))", F1)
            C(ws, i, 6, f"=INDEX({R(col['true'])},MATCH(A{i}*10,{R('A')},0))", F1)
        else:
            C(ws, i, 5, f"=INDEX({R('T')},MATCH(A{i}*10+1,{R('A')},0))", F1)
            C(ws, i, 6, f"=E{i}", F1)
        C(ws, i, 7, f"=IF(AND(D{i}>=MIN({days},3),E{i}<{W['TH']},F{i}<{W['TH']}),1,0)")
        for c, k in ((8, "center"), (9, "early"), (10, "late"), (11, "true")):
            nc = ncol[k]
            # 처음 필요 순번 = 순번 1~4 중 평가 기간 안에서 유효 행이 '필요'인 가장 작은 순번 (MINIFS 없이 — 구형 엑셀 호환)
            f = '""'
            for kk in (4, 3, 2, 1):
                f = f"IF(AND({kk}<={days},COUNTIFS({R('B')},A{i},{R('I')},1,{R('E')},{kk},{R(nc)},1)>0),{kk},{f})"
            C(ws, i, c, f"=IF(G{i}=0,\"\",{f})")
        C(ws, i, 12, (f'=IF(G{i}=0,"",IF(AND(H{i}="",K{i}=""),"둘 다 기간 안 필요 없음",IF(H{i}="","실제만 필요(놓침)",'
                      f'IF(K{i}="","예보만 필요(헛경보)",IF(H{i}=K{i},"같은 날",IF(H{i}-K{i}=-1,"하루 빠름",'
                      f'IF(H{i}-K{i}=1,"하루 늦음",IF(H{i}<K{i},"이틀 이상 빠름","이틀 이상 늦음"))))))))'))
        C(ws, i, 13, f'=IF(G{i}=0,"",IF(AND(IF(I{i}="",99,I{i})<=IF(K{i}="",99,K{i}),IF(K{i}="",99,K{i})<=IF(J{i}="",99,J{i})),1,0))')
        C(ws, i, 14, f'=IF(G{i}=0,"",IF(OR(H{i}<>"",K{i}<>"",I{i}<>""),1,0))')
        C(ws, i, 15, f'=IF(G{i}=0,"",IF(IF(K{i}="",99,K{i})<IF(I{i}="",99,I{i}),1,0))')
    ws.freeze_panes = "D2"
    widths(ws, {"A": 7, "B": 6, "C": 16, "D": 12, "E": 9, "F": 9, "G": 8, "H": 11, "I": 8, "J": 8, "K": 8, "L": 20, "M": 12, "N": 10, "O": 12})
    return ws, len(blocks) + 1


def sheet_errtab(wb, res, title="오차표"):
    """예보 ETo 오차표(범위·± 오차의 근거): 일(발표 × 선행일 × 월)과 3일 누적"""
    ws = wb.create_sheet(title)
    err = res["err"]
    T(ws, 1, 1, "예보 ETo 오차표 — 서비스의 '± 오차'와 관수 필요 예상일 범위(r = RMSE ÷ 관측 평균)에 씀", bold=True, color=GREEN)
    yrs = sorted(set(",".join(err.years.astype(str)).split(","))) if len(err) else []
    T(ws, 2, 1, f"자료: {res.get('err_path') or '(없음)'} — 지점 {err.stn_src.iloc[0] if len(err) else '-'}, 해 {', '.join(yrs) or '-'}"
                + (f" (검증 연도 {res['year']}는 뺌: 다른 해 오차로 채점)" if res.get("year") else "")
                + f". 월 칸 표본이 15보다 적으면 '전체 월(0)' 값을 씀", size=9, color="555555")
    heads = ["구분", "발표", "선행일", "월 (0=전체)", "표본 n", "RMSE (mm/일)", "MBE (mm/일)", "관측 평균 (mm/일)", "상대 오차 r", "해"]
    for c, h in enumerate(heads, 1):
        H(ws, 4, c, h)
    for i, x in enumerate(err.sort_values(["kind", "run_name", "lead_day", "month"]).itertuples(index=False), 5):
        vals = ["일" if x.kind == "day" else "3일 누적", x.run_name, int(x.lead_day), int(x.month), int(x.n), x.rmse, x.mbe,
                x.obs_mean, f"=IF(H{i}>0,F{i}/H{i},\"\")", x.years]
        for c, (v, f) in enumerate(zip(vals, [None, None, None, None, None, F2, F2, F2, "0%", None]), 1):
            C(ws, i, c, _v(v), f, fill=(LIGHT if x.month == 0 else None))
    widths(ws, {"A": 9, "B": 6, "C": 7, "D": 10, "E": 7, "F": 11, "G": 11, "H": 13, "I": 10, "J": 11})
    return ws


def build_wbverify_workbook(res, out):
    """G4 검증 엑셀"""
    wb = Workbook()
    wb.remove(wb.active)
    ws_sum = wb.create_sheet("요약")
    W = {"_irr_log": {pd.Timestamp(k): v for k, v in (res.get("irrig") or {}).items()}}
    soil = res["soil"]
    extra = [("SCN", "관수 규칙 시나리오 (0 = 무관수, 1 = 규칙)", 1 if res.get("auto_irrigate") else 0, "0", True,
              "1이면 관측 물수지에서 전날 끝 Dr ≥ RAW인 날 그 Dr만큼(순량) 관수해 포장용수량으로 되돌림(FAO-56 기본 관수 규칙). 발표마다 예보·참값은 같은 출발에서 관수 없이 진행"),
             ("TH", "판정 기준 고갈량 (mm)", "=RAW", F1, True,
              "관수 필요 판정(● · 예상일)의 기준. 기본은 RAW. 더 자주 관수하는 농가 기준(예: 30~50 mm)으로 바꿔 민감도를 볼 수 있음(Ks는 RAW 그대로)"),
             ("DAYS", "예상일 평가 대상일 수", MAIN_DAYS, "0", True, "3 = 주 지표 기간(아침 D+0~D+2, 저녁 D+1~D+3), 4 = 마지막 날(참고) 포함")]
    _settings(wb, soil, extra, W, f"G4 예보 물수지 검증 — ASOS {res['stn']} · {res['year']}년")
    ws_owb, owb_dates, owb_dr, owb_last = sheet_owb(wb, res["owb"], W, scenario_cell=W["SCN"],
                                                    note="관측 ETo가 없는 날은 그날 아침 발표 D+0 예보로 채움(노란 바탕, '채움'). 결측 영향 = 채운 날 이후 Dr이 0이 되기 전까지")
    ws_runs, last, col, ncol, blocks = sheet_runs(wb, res, W, owb_dates, owb_dr)
    ws_rs, rs_last = sheet_runsum(wb, res, W, last, col, ncol, blocks)
    sheet_errtab(wb, res)
    _summary(ws_sum, res, W, last, col, ncol, rs_last)
    _method_verify(wb, res)
    _order(wb, ["요약", "발표별", "발표요약", "관측물수지", "오차표", "설정", "방법"])
    return safe_save(wb, out)


def _order(wb, names):
    """시트 순서를 names 순서로"""
    wb._sheets = [wb[n] for n in names if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in names]
    wb.active = 0


def _summary(ws, res, W, last, col, ncol, rs_last):
    soil = res["soil"]
    R = lambda c: f"발표별!${c}$2:${c}${last}"
    S = lambda c: f"발표요약!${c}$2:${c}${rs_last}"
    stn, yr = res["stn"], res["year"]
    T(ws, 1, 1, f"G4 예보 물수지 검증 — ASOS {stn} · {yr}년 · 사과 (주 방법 {'S4' if 'ETo_S4' in res['ft'] else 'S3'}, 편향 보정 없음)",
      size=13, bold=True, color=GREEN)
    ft = res["ft"]
    runs = res["runs"]
    okr = runs[runs.ok]
    T(ws, 2, 1, f"채점 대상일 {pd.to_datetime(okr.target).min():%Y-%m-%d}~{pd.to_datetime(okr.target).max():%Y-%m-%d}, 발표 "
                f"{runs.groupby(['run_name', 'run']).ngroups}회(아침·저녁), 채점 행 {int(runs.ok.sum())}/{len(runs)} — "
                f"TAW {soil['taw']:.0f} mm · RAW {soil['raw']:.0f} mm, 관수 규칙 시나리오 "
                f"{'규칙(전날 끝 Dr ≥ RAW면 관수)' if res.get('auto_irrigate') else '무관수'} (설정 시트에서 바꾸면 다시 계산)", size=9, color="555555")
    r = 4
    _sec(ws, r, "A. 예상 고갈량(Dr) 오차 — 참값: 같은 출발에서 관측 ETc·관측 강수로 진행 (mm)", 14); r += 1
    hd = ["발표", "순번", "선행일", "n", "RMSE 중심\n(기대 강수) ★", "MBE 중심", "RMSE\n예보 강수 그대로", "RMSE\n비 무시",
          "RMSE\n관측 강수", "RMSE\n기준선", "개선율 중심\n(기준선 대비)", "MBE\n예보 강수 그대로", "MBE\n비 무시"]
    for c, h in enumerate(hd, 1):
        H(ws, r, c, h)
    r += 1
    lead = res["lead"]
    for x in lead.itertuples():
        crit = f"({R('C')}=\"{x.run_name}\")*({R('E')}={x.order})*({R('I')}=1)"
        C(ws, r, 1, x.run_name); C(ws, r, 2, x.order); C(ws, r, 3, f"D+{x.lead_day}")
        C(ws, r, 4, f"=SUMPRODUCT({crit})")
        rm = lambda k: f"=SQRT(SUMPRODUCT({crit}*({R(col[k])}-{R(col['true'])})^2)/D{r})"
        mb = lambda k: f"=SUMPRODUCT({crit}*({R(col[k])}-{R(col['true'])}))/D{r}"
        grey = x.order > MAIN_DAYS
        clr = GREY if grey else "1A1A1A"
        C(ws, r, 5, rm("center"), F2, bold=True, color=clr); C(ws, r, 6, mb("center"), F2, color=clr)
        C(ws, r, 7, rm("fcst"), F2, color=clr); C(ws, r, 8, rm("none"), F2, color=clr); C(ws, r, 9, rm("obs"), F2, color=clr)
        C(ws, r, 10, rm("pers"), F2, color=clr); C(ws, r, 11, f"=1-E{r}/J{r}", "0%", color=clr)
        C(ws, r, 12, mb("fcst"), F2, color=clr); C(ws, r, 13, mb("none"), F2, color=clr)
        r += 1
    r = _note(ws, r, "순번 1~3 = 주 지표(아침 D+0~D+2, 저녁 D+1~D+3), 4 = 마지막 날(참고, 회색). 관측 강수 열은 강수 예보가 완벽할 때 — 나머지 오차는 ETc 예보 몫. "
                     "기준선 = 예보 없이 어제 관측 ETc가 이어지고 비가 없다고 볼 때", width=160)
    r += 1
    _sec(ws, r, "B. 대상일별 '관수 필요(Dr ≥ 판정 기준)' 적중 — 중심 경로 vs 참값 (판정 기준: 설정 시트, 기본 RAW)", 14); r += 1
    for c, h in enumerate(["발표", "순번", "n", "적중", "놓침", "헛경보", "정상(둘 다 아님)", "POD\n적중률", "FAR\n오보율", "CSI"], 1):
        H(ws, r, c, h)
    r += 1
    for x in res["cont"].itertuples():
        base = f"{R('C')},\"{x.run_name}\",{R('E')},{x.order},{R('I')},1"
        C(ws, r, 1, x.run_name); C(ws, r, 2, x.order); C(ws, r, 3, f"=COUNTIFS({base})")
        C(ws, r, 4, f"=COUNTIFS({base},{R(ncol['center'])},1,{R(ncol['true'])},1)")
        C(ws, r, 5, f"=COUNTIFS({base},{R(ncol['center'])},0,{R(ncol['true'])},1)")
        C(ws, r, 6, f"=COUNTIFS({base},{R(ncol['center'])},1,{R(ncol['true'])},0)")
        C(ws, r, 7, f"=COUNTIFS({base},{R(ncol['center'])},0,{R(ncol['true'])},0)")
        C(ws, r, 8, f"=IFERROR(D{r}/(D{r}+E{r}),\"\")", F2); C(ws, r, 9, f"=IFERROR(F{r}/(D{r}+F{r}),\"\")", F2)
        C(ws, r, 10, f"=IFERROR(D{r}/(D{r}+E{r}+F{r}),\"\")", F2)
        r += 1
    r += 1
    _sec(ws, r, "C. 관수 필요 예상일 — 출발 고갈량이 판정 기준 미만인 발표, 평가 대상일 수: 설정 시트(기본 3) (발표요약 시트)", 14); r += 1
    cats = NEED_CATS
    for c, h in enumerate(["발표", "평가 발표"] + cats + ["사건\n(예보·참값 중 필요)", "±1일 이내\n비율", "범위 사건\n(빠르면 포함)", "범위 적중", "실제가\n'빠르면'보다 앞"], 1):
        H(ws, r, c, h)
    r += 1
    for rn in ("아침", "저녁"):
        C(ws, r, 1, rn)
        C(ws, r, 2, f"=COUNTIFS({S('B')},\"{rn}\",{S('G')},1)")
        for j, cat in enumerate(cats):
            C(ws, r, 3 + j, f"=COUNTIFS({S('B')},\"{rn}\",{S('L')},\"{cat}\")")
        c0 = 3 + len(cats)
        C(ws, r, c0, f"=B{r}-{CL(c0 - 1)}{r}")
        C(ws, r, c0 + 1, f"=IFERROR(SUM(C{r}:E{r})/{CL(c0)}{r},\"\")", "0%")
        C(ws, r, c0 + 2, f"=COUNTIFS({S('B')},\"{rn}\",{S('N')},1)")
        C(ws, r, c0 + 3, f"=IFERROR(COUNTIFS({S('B')},\"{rn}\",{S('N')},1,{S('M')},1)/{CL(c0 + 2)}{r},\"\")", "0%")
        C(ws, r, c0 + 4, f"=COUNTIFS({S('B')},\"{rn}\",{S('N')},1,{S('O')},1)")
        r += 1
    r = _note(ws, r, "예상일 = 처음으로 Dr ≥ 판정 기준이 되는 대상일(순번). 범위: 빠르면 = 비가 오지 않고 ETc가 오차(r)만큼 많을 때, 늦으면 = 예보 강수가 모두 오고 ETc가 오차만큼 적을 때. "
                     "범위 적중 = 참값 예상일(기간 안에 없으면 '기간 뒤')이 빠르면~늦으면 사이", width=160)
    r += 1
    _sec(ws, r, "D. 대상일 월별 예상 Dr 오차 (주 지표 기간, 중심 경로)", 14); r += 1
    for c, h in enumerate(["발표", "월", "n", "RMSE 중심", "MBE 중심", "RMSE 관측 강수", "예보 강수\n평균 (mm/일)", "기대 강수\n평균", "관측 강수\n평균"], 1):
        H(ws, r, c, h)
    r += 1
    for x in res["month"].itertuples():
        crit = f"({R('C')}=\"{x.run_name}\")*({R('G')}={x.month})*({R('I')}=1)*({R('E')}>=1)*({R('E')}<={MAIN_DAYS})"
        C(ws, r, 1, x.run_name); C(ws, r, 2, x.month); C(ws, r, 3, f"=SUMPRODUCT({crit})")
        C(ws, r, 4, f"=SQRT(SUMPRODUCT({crit}*({R(col['center'])}-{R(col['true'])})^2)/C{r})", F2)
        C(ws, r, 5, f"=SUMPRODUCT({crit}*({R(col['center'])}-{R(col['true'])}))/C{r}", F2)
        C(ws, r, 6, f"=SQRT(SUMPRODUCT({crit}*({R(col['obs'])}-{R(col['true'])})^2)/C{r})", F2)
        C(ws, r, 7, f"=SUMPRODUCT({crit}*{R('N')})/C{r}", F1); C(ws, r, 8, f"=SUMPRODUCT({crit}*{R('O')})/C{r}", F1)
        C(ws, r, 9, f"=SUMPRODUCT({crit}*{R('R')})/C{r}", F1)
        r += 1
    r += 1
    _sec(ws, r, "E. 예보 강수 검증 (대상일 일합계, 채점 행) — 값(Python 계산)", 14); r += 1
    for c, h in enumerate(["발표", "순번", "n", "예보 강수\n평균", "기대 강수\n평균", "관측 강수\n평균", "예보/관측", "기대/관측",
                           "비 예보일\n비율(≥1mm)", "비 온 날\n비율", "POD", "FAR"], 1):
        H(ws, r, c, h)
    r += 1
    for x in res["rain"].itertuples():
        for c, (v, f) in enumerate(zip([x.run_name, x.order, x.n, x.fcst, x.exp, x.obs, x.ratio, x.ratio_exp, x.f_days, x.o_days, x.POD, x.FAR],
                                       [None, None, None, F1, F1, F1, F2, F2, "0%", "0%", F2, F2]), 1):
            C(ws, r, c, _v(v), f)
        r += 1
    r += 1
    _sec(ws, r, "F. 판정 기준별 관수 필요 예상일 (민감도, 주 지표 기간) — 값(Python 계산). 기준만 바꾸고 Ks는 RAW 그대로", 14); r += 1
    for c, h in enumerate(["판정 기준 (mm)", "발표", "평가 발표", "사건", "같은 날", "±1일 이내", "빠름·헛경보", "늦음·놓침", "범위 사건", "범위 적중", "실제가 '빠르면'보다 앞"], 1):
        H(ws, r, c, h)
    r += 1
    for x in res["sens"].itertuples():
        for c, (v, f) in enumerate(zip([x.threshold, x.run_name, x.n_runs, x.n_event, x.same, x.within1, x.early, x.late_or_miss,
                                        x.n_range_event, x.coverage, x.true_before_early],
                                       [None, None, None, None, None, "0%", None, None, None, "0%", None]), 1):
            C(ws, r, c, _v(v), f, fill=(LIGHT if x.run_name == "전체" else None))
        r += 1
    r += 1
    _sec(ws, r, "해석", 14); r += 1
    for ln in res["findings"]:
        r = _note(ws, r, ln, width=170, color="1A1A1A")
    widths(ws, {CL(i): w for i, w in enumerate([8, 7, 8, 7, 12, 9, 12, 10, 10, 10, 11, 11, 10, 11, 10, 10], 1)})


def _method_verify(wb, res):
    ws = wb.create_sheet("방법")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 110
    doc = [("예보 물수지 (G4)", "", True),
           ("목적", "어제까지의 관측 물수지(Dr)에서 출발해 예보 ETc와 예보 강수로 며칠 뒤 Dr을 내다보고, 처음 Dr ≥ RAW가 되는 날(관수 필요 예상일)을 알린다."),
           ("물수지 식", "FAO-56 식(84·85·86·88). Ks는 전날 끝 고갈량 Dr,i-1로 정함(#15). 한 칸 수식: Dr = MIN(MAX(Dr,i-1 − P + Ks·ETc, 0), TAW) — DP 식(88)을 넣고 정리한 식"),
           ("출발", "아침(02시): 관측 Dr(D−1 끝) → D+0~D+3. 저녁(17시): 관측 Dr(D−1 끝) + 그날(D)은 아침 발표 D+0 예보로 진행 → D+1~D+4 (발표별 시트의 '그날' 행)"),
           ("예보 ETc", "Kc × 예보 ETo(주 방법 S4, 없으면 S3). 편향 보정 없음(#10)"),
           ("중심 경로", "기대 강수 = 시각별 예보 강수량 × 강수확률의 하루 합. 예보 강수량은 '비가 온다면'의 양이라(강수확률 60% 이상인 시각에만 값) 확률을 곱해 기댓값으로 씀(#19)"),
           ("범위", "빠르면 = 비 없음 + ETc × (1 + r), 늦으면 = 예보 강수 전부 + ETc × (1 − r). r = 월·발표·선행일별 예보 ETo 상대 오차(오차표 시트, 검증 연도를 뺀 다른 해 자료)"),
           ("비교 경로", "예보 강수 그대로(확률 곱하지 않음), 비 무시(0), 관측 강수(완벽한 강수 예보 — ETc 예보의 몫만 남음), 기준선(어제 관측 ETc 지속·비 없음)"),
           ("참값", "같은 출발에서 관측 ETc(Kc × 관측 ETo)와 관측 강수로 관수 없이 진행한 Dr. 대상일(저녁은 그날 D 포함)에 관측이 빠진 날이 있으면 그 행부터 채점에서 뺌"),
           ("관측 결측", "관측 ETo가 없는 날은 그날 아침 발표 D+0 예보로 채움(운영 규칙). 예보·참값이 같은 출발을 쓰므로 출발 고갈량이 채운 날에 기대더라도 비교는 공정함"),
           ("관수 규칙 시나리오", "설정 1이면 관측 물수지에서 전날 끝 Dr ≥ RAW인 날 그 Dr만큼 관수(순량). Dr이 0~RAW 사이를 오가는 관수 농가에 가까운 출발 상태를 만든다"),
           ("주 지표·참고", "처음 3개 대상일이 주 지표, 마지막 날(아침 D+3·저녁 D+4)은 참고(#18)"),
           ("수식·값", "값: 예보 입력(ETc·강수·기대 강수·강수확률), 관측 ETc·강수, 오차표, E·F표(Python). 수식: 관측 물수지, 경로별 Dr, 관수 필요 판정, 예상일·판정·범위 적중, A~D표"),
           ("근거 문서", "docs/THEORY.md 6·9장, docs/VALIDATION.md G4, 해석: docs/RESULTS_GUIDE.md")]
    for i, (k, v, *hd) in enumerate(doc, 1):
        head = bool(hd and hd[0])
        H(ws, i, 1, k, GREEN if head else LIGHT, white=head)
        x = ws.cell(i, 2, v); x.font = Font(name=FONT, size=10, bold=head); x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 18 if head else 32


# ── 서비스 엑셀 ──────────────────────────────────────────────────────────
def build_service_workbook(sv, out):
    """한 발표의 관수 전망 엑셀. sv: cropwater_fcst의 service 준비 결과(dict: outlook, owb, soil, err, meta, stn, stn_name, grid, bias_note ...)"""
    ol, soil = sv["outlook"], sv["soil"]
    wb = Workbook()
    wb.remove(wb.active)
    ws_main = wb.create_sheet("관수 전망")
    W = {"_irr_log": {pd.Timestamp(k): v for k, v in (sv.get("irrig") or {}).items()}}
    extra = [("WETC", "과습 주의 기준 Dr (mm)", f"=RAW*{WET_CAUTION_RATIO}", F1, True,
              "관측 물수지의 끝 고갈량 Dr이 이 값 이하이면 '과습 주의'(Dr = 0이면 '과습 경고'). 기본 = RAW의 "
              f"{WET_CAUTION_RATIO:.0%} — 포장용수량에 닿은 뒤 하루 이틀 치 증발산만 빠진 상태. 숫자로 바꿔 쓸 수 있음(참고 지표, 검증 전)")]
    _settings(wb, soil, extra, W, f"관수 전망 — ASOS {sv['stn']} {sv.get('stn_name', '')} · {ol['run']:%Y-%m-%d %H시} {ol['run_name']} 발표")
    ws_owb, owb_dates, owb_dr, owb_last = sheet_owb(wb, sv["owb"], W, name="관측 물수지",
                                                    note="노란 칸(관수 기록)에 준 관수량(공급 mm, 10a당 1톤 = 1 mm)을 적으면 어제 끝 Dr과 전망이 다시 계산됩니다")
    # 출발 = 관측 마지막 날 끝 Dr (보통 어제. 어제 관측이 아직 없으면 그 전날 — 어제는 '예보 물수지'에서 예보 하루로 진행)
    dates = [pd.Timestamp(x) for x in sv["owb"].date]
    W["_owb_row"] = {x: i + 2 for i, x in enumerate(dates)}
    owb_last = W["_owb_row"].get(pd.Timestamp(ol["prev"]), owb_last)
    fr = _sheet_service_wb(wb, sv, W, owb_last)
    tl = _sheet_timeline(wb, sv, owb_last, fr)
    sheet_errtab(wb, dict(err=sv["err"], err_path=sv.get("err_path"), year=None))
    _sheet_bias(wb, sv)
    _service_main(ws_main, sv, W, fr, tl)
    _method_service(wb, sv)
    _order(wb, ["관수 전망", "예보 물수지", "관측 물수지", TL, "오차표", "편향 점검", "설정", "방법"])
    return safe_save(wb, out)


# 날짜별 시트의 열 ('관수 전망'의 날짜 조회 수식이 찾아오는 곳)
TL_HEAD = ["날짜", "구분", "선행일", "ETo (mm)", "Kc", "ETc (mm)", "ETc 오차 ± (mm)", "강수 (mm)\n관측·예보", "강수확률\n(최대)",
           "기대 강수 (mm)", "순관수 I (mm)", "심층침투 DP\n(mm)", "Dr 중심 (mm)", "Dr 빠르면 (mm)", "Dr 늦으면 (mm)", "관측 물수지\n(1/0)"]
TL_COL = dict(date="A", kind="B", lead="C", eto="D", kc="E", etc="F", err="G", rain="H", pop="I", rain_exp="J", irr="K", dp="L",
              c="M", e="N", l="O", obs="P")
TL_CELLS = ["sel", "D", "k0", "p0", "pt", "p1", "p2", "p3", "p4", "pw", "ps"]      # 날짜별!S1.. 찾기 칸


def _sheet_timeline(wb, sv, owb_last, fr):
    """날짜별 시트: 관측 물수지(첫날~관측 마지막 날) 뒤에 예보 물수지(관측 전·그날·대상일)를 날짜순으로 이은 표(모두 참조 수식).
       '관수 전망'의 조회일(노란 칸)이 이 표에서 그 날짜의 행을 찾는다.
       반환 dict(last, rng(열 키) → 범위, cell[찾기 칸 이름] → 주소, ix(열 키, 찾기 칸) → INDEX 식, k0, D, first, end)"""
    ol = sv["outlook"]
    ws = wb.create_sheet(TL)
    for c, h in enumerate(TL_HEAD, 1):
        H(ws, 1, c, h, GREEN if c <= 3 else (BLUE if c <= 12 else BROWN))
    o = "'관측 물수지'"
    for i in range(2, owb_last + 1):                     # 관측 물수지와 같은 행 번호
        vals = [(f"={o}!A{i}", DATE), ("관측 물수지", None),
                (f'=IF({o}!D{i}="관측","관측",IF({o}!D{i}="예보","예보로 채움","ETo 없음"))', None),
                (f"={o}!C{i}", F2), (f"={o}!B{i}", "0.000"), (f"={o}!E{i}", F2), ('=""', F2), (f"={o}!F{i}", F1), ('=""', "0%"),
                (f"={o}!F{i}", F1), (f"={o}!H{i}", F1), (f"={o}!L{i}", F1), (f"={o}!M{i}", F1), (f"={o}!M{i}", F1),
                (f"={o}!M{i}", F1), (1, None)]
        for c, (v, f) in enumerate(vals, 1):
            C(ws, i, c, v, f)
    f_ = "'예보 물수지'"
    r = owb_last + 1
    for rr in range(3, fr["last"] + 1):                  # 예보 물수지: 관측 전(있으면)·그날(저녁)·대상일
        blank = lambda col: f'=IF({f_}!{col}{rr}="","",{f_}!{col}{rr})'
        vals = [(f"={f_}!B{rr}", DATE), (f"={f_}!A{rr}", None), (f"={f_}!C{rr}", None), (f"={f_}!D{rr}", F2), (f"={f_}!E{rr}", "0.000"),
                (f"={f_}!F{rr}", F2), (blank("G"), F2), (f"={f_}!I{rr}", F1), (blank("J"), "0%"), (f"={f_}!K{rr}", F1),
                ('=""', F1), ('=""', F1), (f"={f_}!L{rr}", F1), (f"={f_}!M{rr}", F1), (f"={f_}!N{rr}", F1), (0, None)]
        for c, (v, f) in enumerate(vals, 1):
            C(ws, r, c, v, f, fill=REF_FILL)
        r += 1
    last = r - 1
    rng = lambda k: f"{TL}!${TL_COL[k]}$2:${TL_COL[k]}${last}"
    D = pd.Timestamp(ol["run"]).normalize()
    k0 = 1 if ol["run_name"] == "저녁" else 0             # 주 지표 첫날 = 조회일 + k0
    pos = lambda d: f'=IFERROR(MATCH({d},$A$2:$A${last},0),"")'
    helpers = [("조회일 ('관수 전망' 노란 칸)", f"={SEL}", DATE),
               ("발표일", D.to_pydatetime(), DATE),
               ("주 지표 첫날 = 조회일 +", k0, None),
               ("전날 행 (조회일 − 1)", pos("$S$1-1"), None),
               ("조회일 행", pos("$S$1"), None),
               ("대상일 1 행", pos("$S$1+$S$3"), None),
               ("대상일 2 행", pos("$S$1+$S$3+1"), None),
               ("대상일 3 행", pos("$S$1+$S$3+2"), None),
               ("대상일 4 행 (참고)", pos("$S$1+$S$3+3"), None),
               # 과습 판정 행: 전날이 관측 물수지면 그 행, 전날이 '관측 전'(예보로 먼저 진행한 날)이면 관측 마지막 날. 그 밖(예보)은 판정 안 함
               ("과습 판정 행 (관측 물수지)", f'=IF($S$4="","",IF(INDEX($P$2:$P${last},$S$4)=1,$S$4,'
                                       f'IF(LEFT(INDEX($B$2:$B${last},$S$4),4)="관측 전",{owb_last - 1},"")))', None),
               ("주 지표 출발 행", '=IF($S$3=0,$S$4,$S$5)', None)]
    for i, (k, v, f) in enumerate(helpers, 1):
        H(ws, i, 18, k, LIGHT, white=False); C(ws, i, 19, v, f)
    T(ws, len(helpers) + 2, 18, "행 = 이 표에서 위에서 몇 번째 날짜인지(없으면 빈 칸 → '자료 없음'). 이 칸들은 바꾸지 마세요", size=9, color="555555")
    ws.freeze_panes = "B2"
    widths(ws, {"A": 11, "B": 22, "C": 12, "D": 8, "E": 7, "F": 8, "G": 9, "H": 10, "I": 8, "J": 9, "K": 9, "L": 10, "M": 10,
                "N": 10, "O": 10, "P": 9, "Q": 3, "R": 28, "S": 13})
    cell = {n: f"{TL}!$S${i}" for i, n in enumerate(TL_CELLS, 1)}
    return dict(last=last, rng=rng, cell=cell, k0=k0, D=D, first=pd.Timestamp(sv["owb"].date.iloc[0]),
                end=pd.Timestamp(ol["days"].target.max()), ix=lambda k, p: f"INDEX({rng(k)},{cell[p]})")


def _sheet_service_wb(wb, sv, W, owb_last):
    """예보 물수지 시트: (저녁은 그날 D) + 대상일. 반환: 행 정보 dict"""
    ol = sv["outlook"]
    ws = wb.create_sheet("예보 물수지")
    heads = ["구분", "날짜", "선행일", "ETo 예보 (mm)", "Kc", "ETc 예보 (mm)", "ETc 오차 ± (mm)", "상대 오차 r", "예보 강수 (mm)",
             "강수확률 (최대)", "기대 강수 (mm)", "Dr 중심 (mm)", "Dr 빠르면 (mm)", "Dr 늦으면 (mm)", "상태 (중심)"]
    for c, h in enumerate(heads, 1):
        H(ws, 1, c, h, GREEN if c <= 3 else (BLUE if c <= 11 else BROWN))
    D = pd.Timestamp(ol["run"]).normalize()
    lagged = bool(ol.get("lag_days"))
    C(ws, 2, 1, "출발"); C(ws, 2, 2, _v(ol["prev"]), DATE)
    C(ws, 2, 3, "관측 마지막 날 끝 (관측 물수지)" if lagged else "D−1 끝 (관측 물수지)")
    for c in (12, 13, 14):
        C(ws, 2, c, f"='관측 물수지'!M{owb_last}", F1, bold=True)
    r = 3
    rows = []
    for p in ol["pre"]:
        k = p.get("kind", "today")
        irr_row = W.get("_owb_row", {}).get(pd.Timestamp(p["date"])) if k == "lag" else None
        rows.append(dict(kind="관측 전" if k == "lag" else "그날", date=p["date"], lead=0 if k != "lag" else -(D - pd.Timestamp(p["date"])).days,
                         ETo=p["ETo"], Kc=p["Kc"], err=None, rel=p["rel"], rain=p["rain"], pop=p.get("pop_max"),
                         rain_exp=p["rain_exp"], irr=(f"'관측 물수지'!H{irr_row}" if irr_row else None),
                         known=p.get("rain_known", False)))
    for x in ol["days"].itertuples():
        rows.append(dict(kind="주" if x.main else "참고", date=x.target, lead=int(x.lead_day), ETo=x.ETo_main, Kc=x.Kc,
                         err=x.etc_err, rel=x.rel, rain=x.rain, pop=getattr(x, "pop_max", np.nan),
                         rain_exp=(x.rain_exp if not pd.isna(getattr(x, "rain_exp", np.nan)) else x.rain)))
    info = dict(first=r, main=[], ref=None, pre=None, pres=[])
    for z in rows:
        grey = z["kind"] in ("참고", "그날", "관측 전")
        clr = GREY if grey else "1A1A1A"
        fill = REF_FILL if grey else None
        label = {"참고": "참고 (마지막 날)", "그날": "그날 (아침 D+0 예보)",
                 "관측 전": "관측 전 (그날 아침 D+0 예보" + (", 관측 강수)" if z.get("known") else ")")}.get(z["kind"], "주 지표")
        C(ws, r, 1, label, color=clr, fill=fill)
        lead_txt = f"D+{z['lead']}" if z["lead"] >= 0 else f"D−{-z['lead']}"
        C(ws, r, 2, _v(z["date"]), DATE, color=clr, fill=fill); C(ws, r, 3, lead_txt, color=clr, fill=fill)
        C(ws, r, 4, _v(z["ETo"]), F2, color=clr, fill=fill); C(ws, r, 5, _v(z["Kc"]), "0.000", color=clr, fill=fill)
        C(ws, r, 6, f"=D{r}*E{r}", F2, color=clr, fill=fill)
        C(ws, r, 7, f"=F{r}*H{r}" if z["kind"] != "그날" else None, F2, color=clr, fill=fill)
        C(ws, r, 8, _v(z["rel"]), "0%", color=clr, fill=fill)
        C(ws, r, 9, _v(z["rain"]), F1, color=clr, fill=fill); C(ws, r, 10, _v(z["pop"]), "0%", color=clr, fill=fill)
        C(ws, r, 11, _v(z["rain_exp"]), F1, color=clr, fill=fill)
        irr = z.get("irr")                    # 관측 전 날: 관측 물수지 노란 칸(관수 기록)의 순관수량도 모든 경로에 넣음
        early_p = f"K{r}" if z.get("known") else None     # 관측 강수를 아는 날은 '빠르면'도 그 강수
        C(ws, r, 12, _wb_formula(f"L{r - 1}", f"K{r}", f"F{r}", W, irr), F1, bold=not grey, color=clr, fill=fill)
        C(ws, r, 13, _wb_formula(f"M{r - 1}", early_p, f"F{r}*(1+H{r})", W, irr), F1, color=clr, fill=fill)
        C(ws, r, 14, _wb_formula(f"N{r - 1}", f"I{r}", f"F{r}*MAX(1-H{r},0)", W, irr), F1, color=clr, fill=fill)
        C(ws, r, 15, f'=IF(L{r}>={W["RAW"]},"관수 필요",IF(L{r}>={W["RAW"]}*0.5,"주의","안전"))', color=clr, fill=fill)
        if z["kind"] == "주":
            info["main"].append(r)
        elif z["kind"] == "참고":
            info["ref"] = r
        else:
            info["pre"] = r
            info["pres"].append((r, z["kind"], pd.Timestamp(z["date"])))
        r += 1
    info["last"] = r - 1
    r += 1
    r = _note(ws, r, "Dr 중심 = 기대 강수(예보 강수량 × 강수확률) · 예보 ETc, 빠르면 = 비가 오지 않고 ETc가 오차(r)만큼 많을 때, "
                     "늦으면 = 예보 강수가 모두 오고 ETc가 오차만큼 적을 때. 관수하지 않을 때의 고갈량입니다. 한 칸 수식은 FAO-56 식(84·85·88)과 같은 값", width=150)
    widths(ws, {"A": 20, "B": 11, "C": 16, "D": 10, "E": 7, "F": 10, "G": 11, "H": 9, "I": 10, "J": 10, "K": 10, "L": 11, "M": 11, "N": 11, "O": 11})
    return info


def _need_text(W, tl, k):
    """예상일 문자열 수식(조회일 기준): 출발 ≥ RAW면 '지금 필요', 주 지표 3일에 처음 Dr ≥ RAW가 되는 날, 참고 날, 없으면 '3일 안에 없음'.
       k: 경로 열(c 중심 / e 빠르면 / l 늦으면). 자료가 없는 날짜면 '자료 없음'"""
    ix, cell, raw = tl["ix"], tl["cell"], W["RAW"]
    lab = lambda p, tail="": f'TEXT({ix("date", p)},"m/d")&" ("&{ix("lead", p)}&"{tail})"'
    f = f'IF({cell["p4"]}="","3일 안에 없음",IF({ix(k, "p4")}>={raw},{lab("p4", ", 참고")},"3일 안에 없음"))'
    for p in ("p3", "p2", "p1"):
        f = f"IF({ix(k, p)}>={raw},{lab(p)},{f})"
    return f'=IF(OR({cell["ps"]}="",{cell["p3"]}=""),"자료 없음",IF({ix(k, "ps")}>={raw},"지금 필요",{f}))'


def _service_main(ws, sv, W, fr, tl):
    """관수 전망 시트. 조회일(노란 칸, 기본 = 발표일)을 바꾸면 그 날짜의 '지금 토양 상태'와 '앞으로 3일'을 날짜별 시트에서 찾아온다"""
    ol, soil = sv["outlook"], sv["soil"]
    rn, run = ol["run_name"], ol["run"]
    ix, cell, raw, D, k0 = tl["ix"], tl["cell"], W["RAW"], tl["D"], tl["k0"]
    T(ws, 1, 1, f"관수 전망 — {sv.get('stn_name', '')}(ASOS {sv['stn']}) · 사과 · {run:%Y-%m-%d %H시} {rn} 발표", size=14, bold=True, color=GREEN)
    ops = sv.get("ops")
    method = (ops or {}).get("rs_method") or sv.get("main", "S4")
    T(ws, 2, 1, f"격자 {sv.get('grid', '')} 단기예보 + ASOS 관측(어제까지) · 예보 ETo 주 방법 {method} · 편향 보정 없음 · "
                f"대상일 {'오늘~D+3' if rn == '아침' else '내일~D+4'} (마지막 날은 참고)", size=9, color="555555")
    if ops and (ops.get("backup") or "S3" in (ops.get("rs_method") or "") or ops.get("obs_prev") != "관측"):
        T(ws, 3, 1, "※ " + " · ".join(ops_flags(ops, rn)), size=9, bold=True, color="A8681B")
    # ── 날짜 선택 ──
    r = 4
    _sec(ws, r, "날짜 선택", 8); r += 1
    H(ws, r, 1, "조회일 (노란 칸: 목록에서 고르거나 날짜 입력)", LIGHT, white=False)
    C(ws, r, 2, D.to_pydatetime(), DATE, bold=True, size=12, fill=YEL)
    dv = DataValidation(type="list", formula1=f"={TL}!$A$3:$A${tl['last']}", allow_blank=False, showErrorMessage=False)
    dv.prompt, dv.promptTitle, dv.showInputMessage = (f"기본 = 발표일 {D:%Y-%m-%d}. 조회할 수 있는 날짜: {tl['first'] + pd.Timedelta(days=1):%Y-%m-%d} ~ "
                                                     f"{tl['end']:%Y-%m-%d}"), "조회일", True
    ws.add_data_validation(dv); dv.add(ws.cell(r, 2))
    sel = f"$B${r}"
    C(ws, r, 3, (f'=IF({cell["p0"]}="","자료 없음 — 조회할 수 있는 날짜: {tl["first"] + pd.Timedelta(days=1):%Y-%m-%d} ~ {tl["end"]:%Y-%m-%d}",'
                 f'IF({sel}={cell["D"]},"발표일 (이 발표의 관수 전망)",IF({sel}<{cell["D"]},'
                 f'"지난 날짜 — 관측 물수지 결과 (발표일 뒤로 넘어가는 날은 이 발표의 예보)","발표일 뒤 — 이 발표의 예보로 진행한 값")))'),
      left=True, bold=True, color="A8681B")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    r += 2
    # ── 지금 토양 상태 (조회일 전날 끝) ──
    _sec(ws, r, "지금 토양 상태 (조회일 전날 끝)", 8); r += 1
    r_dr, r_raw, r_pct, r_st, r_dp = r, r + 1, r + 2, r + 3, r + 4
    H(ws, r_dr, 1, f'="전날(" & TEXT({sel}-1,"m/d") & ") 끝 고갈량 Dr"', LIGHT, white=False)
    C(ws, r_dr, 2, f'=IF({cell["p0"]}="","자료 없음",{ix("c", "p0")})', F1, bold=True)
    C(ws, r_dr, 3, (f'=IF({cell["p0"]}="","","mm — "&IF({ix("obs", "p0")}=1,"관측 물수지",{ix("kind", "p0")}&", 범위 "'
                    f'&TEXT({ix("l", "p0")},"0")&" ~ "&TEXT({ix("e", "p0")},"0")&" mm"))'), left=True, size=9, color="555555")
    H(ws, r_raw, 1, "RAW (이만큼 빠지면 관수)", LIGHT, white=False); C(ws, r_raw, 2, f"={raw}", F1, bold=True)
    T(ws, r_raw, 3, "mm", size=9, color="555555")
    H(ws, r_pct, 1, "RAW 대비", LIGHT, white=False); C(ws, r_pct, 2, f'=IF(ISNUMBER(B{r_dr}),B{r_dr}/B{r_raw},"")', "0%", bold=True)
    # 과습(관측 물수지만, 생육기 Kc > 0): 경고 = 끝 Dr 0(0.05 mm 미만), 주의 = 끝 Dr ≤ 과습 주의 기준(설정). 그 밖은 관수 쪽 상태
    wet_dr = ix("c", "pw")
    grow = f'{ix("kc", "pw")}>0'
    dry = f'IF(B{r_dr}>=B{r_raw},"관수 필요",IF(B{r_dr}>=0.5*B{r_raw},"주의","안전"))'
    H(ws, r_st, 1, "상태", LIGHT, white=False)
    C(ws, r_st, 2, (f'=IF({cell["p0"]}="","자료 없음",IF({cell["pw"]}="",{dry},IF(AND({grow},ROUND({wet_dr},1)<=0),"과습 경고",'
                    f'IF(AND({grow},{wet_dr}<={W["WETC"]}),"과습 주의",{dry}))))'), None, bold=True, size=12)
    C(ws, r_st, 3, (f'=IF(OR(B{r_st}="과습 경고",B{r_st}="과습 주의"),"관측 물수지 "&TEXT({ix("date", "pw")},"m/d")&" 끝 Dr "'
                    f'&TEXT({wet_dr},"0.0")&" mm (과습 주의 기준 "&TEXT({W["WETC"]},"0.0")&" mm 이하) — 관수를 미루세요","")'),
      left=True, size=9, color="1F4E79")
    for txt, fill in (("과습 경고", WETW_FILL), ("과습 주의", WETC_FILL), ("관수 필요", NEED_FILL), ("주의", WARN_FILL), ("안전", OK_FILL)):
        ws.conditional_formatting.add(f"B{r_st}", FormulaRule(formula=[f'$B${r_st}="{txt}"'], fill=PatternFill("solid", bgColor=fill, fgColor=fill)))
    wd = ix("date", "pw")
    H(ws, r_dp, 1, "심층침투 DP 최근 7일 합 (관측 물수지)", LIGHT, white=False)
    C(ws, r_dp, 2, (f'=IF({cell["pw"]}="","",SUMIFS({tl["rng"]("dp")},{tl["rng"]("date")},">="&({wd}-6),{tl["rng"]("date")},"<="&{wd},'
                    f'{tl["rng"]("obs")},1))'), F1)
    C(ws, r_dp, 3, (f'=IF({cell["pw"]}="","","mm — 그중 끝 Dr이 0인 날 "&COUNTIFS({tl["rng"]("c")},"<0.05",{tl["rng"]("date")},">="&({wd}-6),'
                    f'{tl["rng"]("date")},"<="&{wd},{tl["rng"]("obs")},1)&"일 (근권 아래로 빠진 비·관수)")'), left=True, size=9, color="555555")
    r = r_dp + 1
    if k0:                                               # 저녁 발표: 주 지표는 조회일 다음 날부터 → 조회일 끝 Dr을 함께 보임
        H(ws, r, 1, f'="조회일(" & TEXT({sel},"m/d") & ") 끝 예상 Dr"', LIGHT, white=False)
        C(ws, r, 2, f'=IF({cell["pt"]}="","자료 없음",{ix("c", "pt")})', F1, bold=True)
        C(ws, r, 3, f'=IF({cell["pt"]}="","","mm — "&IF({ix("obs", "pt")}=1,"관측 물수지",{ix("kind", "pt")})&" (저녁 발표의 출발)")',
          left=True, size=9, color="555555")
        r += 1
    for rr in range(r_dr, r):
        if rr not in (r_raw, r_pct):
            ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=8)
    r += 1
    # ── 앞으로 3일 ──
    _sec(ws, r, "★ 앞으로 3일 — 주 지표 (조회일" + (" 다음 날" if k0 else "") + "부터)", 8); r += 1
    ok3 = f'OR({cell["p1"]}="",{cell["p3"]}="")'
    sum3 = lambda k: f'SUM({ix(k, "p1")}:{ix(k, "p3")})'
    H(ws, r, 1, "작물 증발산 ETc 3일 합", LIGHT, white=False)
    C(ws, r, 2, f'=IF({ok3},"자료 없음",{sum3("etc")})', F1, bold=True, size=14)
    C(ws, r, 3, (f'=IF(ISNUMBER(B{r}),IF({sel}={cell["D"]},"± "&TEXT(B{r}*D{r},"0.0")&" mm",IF({ix("obs", "p3")}=1,"(관측값)","(± 는 발표일만)")),"")'),
      None, bold=True)
    C(ws, r, 4, _v(ol["cum3"][3]), "0%", fill=LIGHT)
    T(ws, r, 5, f"← 상대 오차 r₃: 같은 달·같은 발표의 3일 누적 예보 ETo 오차 (RMSE ÷ 관측 평균, {ol['cum3'][2]}, 오차표 시트). 발표일 조회에만 적용",
      size=9, color="555555")
    r += 1
    H(ws, r, 1, "강수 3일 합 (예보 · 지난 날짜는 관측)", LIGHT, white=False)
    C(ws, r, 2, f'=IF({ok3},"자료 없음",{sum3("rain")})', F1)
    C(ws, r, 3, f'=IF({ok3},"","(기대 강수 "&TEXT({sum3("rain_exp")},"0.0")&" mm)")')
    T(ws, r, 5, "기대 강수 = 시각별 예보 강수량 × 강수확률. 예보 강수량은 관측보다 많게 나오는 경향이 있음(두 해 검증, 처음 이틀 1.1~2.1배)",
      size=9, color="555555")
    r += 1
    H(ws, r, 1, "관수 필요 예상일 (관수하지 않으면)", LIGHT, white=False)
    C(ws, r, 2, _need_text(W, tl, "c"), None, bold=True, size=13, fill=NEED_FILL)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
    H(ws, r, 1, "  빠르면 (비가 오지 않으면)", LIGHT, white=False)
    C(ws, r, 2, _need_text(W, tl, "e"), None); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
    H(ws, r, 1, "  늦으면 (예보 비가 모두 오면)", LIGHT, white=False)
    C(ws, r, 2, _need_text(W, tl, "l"), None); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
    H(ws, r, 1, "권장 관수량 (예상일에 포장용수량까지)", LIGHT, white=False)
    # 예상일의 중심 Dr: 출발이 이미 필요면 출발 값, 아니면 주 지표 기간에 처음 Dr ≥ RAW인 날의 값
    f = "\"\""
    for p in ("p3", "p2", "p1"):
        f = f'IF({ix("c", p)}>={raw},{ix("c", p)},{f})'
    C(ws, r, 2, f'=IF(OR({cell["ps"]}="",{cell["p3"]}=""),"",IF({ix("c", "ps")}>={raw},{ix("c", "ps")},{f}))', F1, bold=True)
    C(ws, r, 3, f"=IF(B{r}=\"\",\"\",\"순 \"&TEXT(B{r},\"0\")&\" mm → 공급 \"&TEXT(B{r}/{W['EA']},\"0\")&\" mm (10a당 \"&TEXT(B{r}/{W['EA']},\"0\")&\"톤)\")")
    r += 2
    # ── 날짜별 전망 ──
    _sec(ws, r, "날짜별 전망 (마지막 날은 참고)", 13); r += 1
    hd = ["날짜", "선행일", "구분", "ETo (mm)", "Kc", "ETc (mm)", "± 오차 (mm)", "강수 (mm)", "강수확률", "기대 강수 (mm)",
          "Dr 중심 (mm)", "범위: 늦으면~빠르면 (mm)", "상태"]
    for c, h in enumerate(hd, 1):
        H(ws, r, c, h)
    r += 1
    for p in (["pt"] if k0 else []) + ["p1", "p2", "p3", "p4"]:
        grey = p in ("pt", "p4")
        clr, fill = (GREY, REF_FILL) if grey else ("1A1A1A", None)
        g = lambda k, p=p: f'=IF({cell[p]}="","",{ix(k, p)})'
        C(ws, r, 1, f'=IF({cell[p]}="","자료 없음",{ix("date", p)})', DATE, color=clr, fill=fill)
        C(ws, r, 2, g("lead"), color=clr, fill=fill); C(ws, r, 3, g("kind"), color=clr, fill=fill)
        C(ws, r, 4, g("eto"), F2, color=clr, fill=fill); C(ws, r, 5, g("kc"), "0.00", color=clr, fill=fill)
        C(ws, r, 6, g("etc"), F1, color=clr, fill=fill, bold=not grey); C(ws, r, 7, g("err"), F1, color=clr, fill=fill)
        C(ws, r, 8, g("rain"), F1, color=clr, fill=fill); C(ws, r, 9, g("pop"), "0%", color=clr, fill=fill)
        C(ws, r, 10, g("rain_exp"), F1, color=clr, fill=fill)
        C(ws, r, 11, g("c"), F1, color=clr, fill=fill, bold=not grey)
        C(ws, r, 12, f'=IF({cell[p]}="","",IF({ix("obs", p)}=1,"관측",TEXT({ix("l", p)},"0")&" ~ "&TEXT({ix("e", p)},"0")))', None, color=clr, fill=fill)
        # 상태: 관측 물수지 날(지난 날짜)에는 과습도 판정, 예보 날은 관수 쪽만
        C(ws, r, 13, (f'=IF({cell[p]}="","",IF(AND({ix("obs", p)}=1,{ix("kc", p)}>0,ROUND({ix("c", p)},1)<=0),"과습 경고",'
                      f'IF(AND({ix("obs", p)}=1,{ix("kc", p)}>0,{ix("c", p)}<={W["WETC"]}),"과습 주의",'
                      f'IF({ix("c", p)}>={raw},"관수 필요",IF({ix("c", p)}>={raw}*0.5,"주의","안전")))))'), color=clr, fill=fill)
        r += 1
    r += 1
    if ops:
        _sec(ws, r, "수집 정보 (운영 수집기, THEORY 9장 ◆ 운영 수집)", 8); r += 1
        src = pd.Timestamp(ops["src_run"])
        info = [("예보 발표", f"{src:%Y-%m-%d %H}시 발표" + (" — 서비스 발표를 마감까지 받지 못해 직전 발표로 대신(#12)" if ops.get("backup") else " (서비스 발표)")),
                ("일사 추정", {"S4": "S4 (하늘상태·강수확률 포함)", "S3": "S3 (하늘상태·강수확률 없음 — 식50 + 강수유무, #12)"}.get(
                    ops.get("rs_method"), f"{ops.get('rs_method')} (날마다 다름: '예보 물수지' 참고)")),
                ("어제 관측", "ASOS 관측" if ops.get("obs_prev") == "관측" else
                 "아직 조회되지 않음 → 어제를 어제 아침 발표 D+0 예보로 먼저 진행(세 경로 범위 포함, #12). '관측 물수지'에는 예보로 채워 보임"),
                ("대상일", f"{ops.get('n_days')}일" + (" (대체 발표가 마지막 날을 담지 않음)" if ops.get("dropped") else "")),
                ("만든 시각", f"{pd.Timestamp(ops['made_at']):%Y-%m-%d %H:%M}" if ops.get("made_at") is not None else "")]
        for k, v in info:
            H(ws, r, 1, k, LIGHT, white=False); C(ws, r, 2, v, left=True)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            r += 1
        r += 1
    notes = ["읽는 법: 고갈량 Dr은 포장용수량에서 빠진 물(mm)입니다. Dr이 RAW에 닿으면 관수가 필요합니다. 예상 Dr은 '관수하지 않을 때'의 값입니다.",
             "날짜 선택: 조회일(노란 칸)을 바꾸면 그 날짜의 토양 상태와 그 뒤 3일을 보여 줍니다. 지난 날짜는 관측 물수지 결과(실제 비·관수 기록 반영), "
             "발표일 뒤는 이 발표의 예보입니다. 자료가 없는 날짜는 '자료 없음'으로 나옵니다. 발표일로 되돌리려면 발표일을 다시 고르세요.",
             "과습: 관측 물수지에서 전날 끝 Dr이 0(포장용수량까지 참)이면 '과습 경고', 과습 주의 기준(설정 시트, 기본 RAW의 10%) 이하이면 '과습 주의'입니다. "
             "예보는 쓰지 않고, 생육기(Kc > 0)에만 판정합니다. 전날 관측이 아직 없으면 관측 마지막 날 기준입니다. 배수 상태는 계산에 없는 참고 지표입니다.",
             "3일 합(주 지표)이 하루 값보다 믿을 만합니다(FAO-56 권고: 추정 일사로 계산한 ETo는 여러 날 합계로 쓰기). 마지막 날은 참고로만 보세요.",
             "범위: 빠르면 = 예보 비가 오지 않고 증발산이 예보보다 많을 때, 늦으면 = 예보 비가 모두 오고 증발산이 적을 때. 두 해 검증에서 실제 날짜가 '빠르면'보다 앞선 적은 없었습니다.",
             "관수를 했으면 '관측 물수지' 시트의 노란 칸에 적으세요. 어제 끝 Dr과 전망이 바로 다시 계산됩니다.",
             f"예보 ETo는 편향 보정을 하지 않습니다(#10). 월·선행일별 예보 오차는 '편향 점검' 시트에서 봅니다. {sv.get('bias_note', '')}"]
    for n_ in notes:
        r = _note(ws, r, n_, width=150)
    widths(ws, {"A": 40, "B": 14, "C": 16, "D": 10, "E": 7, "F": 10, "G": 10, "H": 12, "I": 9, "J": 12, "K": 14, "L": 20, "M": 11})
    ws.sheet_view.showGridLines = False


def ops_flags(ops, run_name):
    """서비스 엑셀 머리의 운영 표시(직전 발표 대체·S3·전날 관측 채움)"""
    out = []
    if ops.get("backup"):
        out.append(f"{'02' if run_name == '아침' else '17'}시 발표를 받지 못해 {pd.Timestamp(ops['src_run']):%m-%d %H}시 발표로 계산")
    if "S3" in (ops.get("rs_method") or ""):
        out.append("하늘상태·강수확률 없음 → 일사 S3")
    if ops.get("obs_prev") != "관측":
        out.append("어제 ASOS 관측이 아직 없어 어제를 아침 예보로 진행(범위 포함)")
    return out


def _sheet_bias(wb, sv):
    """편향 점검(#10 A): 오차표의 월·발표·선행일별 MBE, 최근 기간(가능하면) 예보−관측"""
    ws = wb.create_sheet("편향 점검")
    T(ws, 1, 1, "편향 점검 — 예보 ETo(주 방법) − 관측 ETo. 보정은 하지 않고(#10) 크기를 매달 확인합니다", bold=True, color=GREEN)
    T(ws, 2, 1, "해마다 편향이 달라(2026년 −5.2%, 2025년 +1.7%) 한 해에서 구한 고정 보정값은 다른 해로 옮겨 가지 않았습니다(VALIDATION #10).",
      size=9, color="555555")
    err = sv["err"]
    d = err[(err.kind == "day") & (err.month > 0)]
    r = 4
    H(ws, r, 1, "발표"); H(ws, r, 2, "선행일")
    months = sorted(d.month.unique())
    for j, m in enumerate(months):
        H(ws, r, 3 + j, f"{int(m)}월 MBE")
    H(ws, r, 3 + len(months), "전체 MBE (mm/일)")
    r += 1
    for (rn, k), g in d.groupby(["run_name", "lead_day"]):
        C(ws, r, 1, rn); C(ws, r, 2, f"D+{int(k)}")
        gm = g.set_index("month")
        for j, m in enumerate(months):
            v = gm.mbe.get(m, np.nan)
            C(ws, r, 3 + j, _v(v), F2, fill=(WARN_FILL if not pd.isna(v) and abs(v) >= 0.3 else None))
        a = err[(err.kind == "day") & (err.run_name == rn) & (err.lead_day == k) & (err.month == 0)]
        C(ws, r, 3 + len(months), _v(a.mbe.iloc[0]) if len(a) else None, F2, bold=True)
        r += 1
    r = _note(ws, r + 1, "노란 칸: |MBE| ≥ 0.3 mm/일. 오차표 자료(검증 연도들)의 값입니다.", width=120)
    rec = sv.get("recent")
    if rec is not None and len(rec):
        r += 1
        _sec(ws, r, f"최근 {sv.get('recent_days', 30)}일 (이 발표 이전, 관측이 있는 대상일) — 예보 − 관측", 8); r += 1
        for c, h in enumerate(["발표", "선행일", "n", "ETo 예보 평균", "ETo 관측 평균", "MBE (mm/일)", "편향 (%)"], 1):
            H(ws, r, c, h)
        r += 1
        for x in rec.itertuples():
            for c, (v, f) in enumerate(zip([x.run_name, f"D+{int(x.lead_day)}", x.n, x.fcst, x.obs, x.mbe, x.pct],
                                           [None, None, None, F2, F2, F2, "0.0%"]), 1):
                C(ws, r, c, _v(v), f)
            r += 1
    widths(ws, {"A": 8, "B": 8, **{CL(3 + j): 10 for j in range(len(months) + 1)}})


def _method_service(wb, sv):
    ws = wb.create_sheet("방법")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 110
    doc = [("관수 전망 계산 방법", "", True),
           ("관측 물수지", "01-Cycle 워크북의 ASOS 관측 ETo·강수와 Kc로 날마다 Dr을 계산(FAO-56 식84·85·86·88, Ks는 전날 끝 Dr). 관수 기록(공급량 × Ea)을 빼 줌"),
           ("예보 ETo", "기상청 단기예보(격자) 기온·습도·풍속 + 추정 일사(S4: 식50형 + 강수확률 + 하늘상태, 선행일별 운영 계수). 편향 보정 없음(#10)"),
           ("예보 물수지", "어제 끝 Dr에서 출발해 날마다 Dr = MIN(MAX(Dr,i-1 − P + Ks·ETc, 0), TAW). 저녁 발표는 오늘을 아침 발표 D+0 예보로 먼저 진행"),
           ("세 경로", "중심 = 기대 강수(강수량 × 강수확률), 빠르면 = 비 없음 + ETc × (1 + r), 늦으면 = 예보 강수 전부 + ETc × (1 − r). r = 오차표의 월·발표·선행일별 상대 오차"),
           ("주 지표", "처음 3일(아침 오늘~D+2, 저녁 내일~D+3) 합계. 마지막 날(아침 D+3·저녁 D+4)은 참고(#18)"),
           ("검증", "2025·2026년 과거 단기예보로 검증(VALIDATION G4): 예보 물수지 오차의 대부분은 강수 예보에서 오며, 범위(빠르면~늦으면)가 실제 관수 필요일을 대부분 포함"),
           ("한계", "예보 강수량은 관측보다 많게 나오는 경향이 있습니다(특히 장마철). 비 예보가 있으면 '빠르면'도 함께 보세요. 한 지점(춘천)·두 해 검증 결과입니다"),
           ("날짜 선택", "'관수 전망'의 조회일(노란 칸, 기본 = 발표일)을 바꾸면 '날짜별' 시트(관측 물수지 + 예보 물수지를 날짜순으로 이은 표)에서 "
                      "전날 끝 Dr과 그 뒤 3일을 찾아옵니다. 지난 날짜는 관측 물수지 결과이고 ± 오차·범위는 없습니다. 표에 없는 날짜는 '자료 없음'"),
           ("과습", "관측 물수지의 전날 끝 Dr로만 판정(예보는 쓰지 않음). 과습 경고 = Dr 0(0.05 mm 미만 — 비·관수가 근권을 포장용수량까지 채움, "
                  "넘친 물은 심층침투 DP), 과습 주의 = Dr ≤ 과습 주의 기준(설정 시트, 기본 RAW × 10%). 생육기(Kc > 0)에만 판정. "
                  "전날 관측이 아직 없으면 관측 마지막 날 기준. FAO-56 물수지는 포장용수량보다 젖은 상태와 배수 속도를 다루지 않으므로 참고 지표(THEORY 9장)"),
           ("운영 규칙", "서비스 발표를 마감(발표 + 40분)까지 못 받으면 24시간 안의 직전 발표로 계산(그 발표 기준 선행일의 S4 계수·오차 r). "
                      "하늘상태·강수확률이 없으면 S3. 어제 ASOS 관측이 아직 없으면 어제를 어제 아침 발표 D+0 예보로 먼저 진행하고, "
                      "저녁 발표의 '그날'처럼 범위(빠르면·늦으면)에 그날의 불확실성을 담음(#12, VALIDATION G5)"),
           ("근거 문서", "docs/THEORY.md 6·9장, docs/VALIDATION.md G4, docs/RESULTS_GUIDE.md")]
    for i, (k, v, *hd) in enumerate(doc, 1):
        head = bool(hd and hd[0])
        H(ws, i, 1, k, GREEN if head else LIGHT, white=head)
        x = ws.cell(i, 2, v); x.font = Font(name=FONT, size=10, bold=head); x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 18 if head else 32
