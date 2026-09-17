"""
excel_util.py — 엑셀 출력 공통 스타일 및 유틸리티

cropwater 프로젝트의 스타일을 공유하여 farmos-* 시리즈 전체에서
일관된 엑셀 디자인을 유지한다.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 스타일 상수
# ---------------------------------------------------------------------------
FONT_NAME = "맑은 고딕"

# 색상 팔레트
GREEN   = "2E6A4C"     # 헤더 (진한 녹색)
BLUE    = "2B6E86"     # 헤더 (파란)
BROWN   = "A8681B"     # 강조
LIGHT   = "EFF1EC"     # 연한 배경 (헤더)
YEL     = "FBEED2"     # 입력 셀 배경
WARN    = "F4C7B8"     # 경고/주의
WHITE   = "FFFFFF"
LIGHT_BLUE = "DAE8F0"  # 올해 행 배경
LIGHT_YEL  = "FFF8E7"  # 평년 행 배경

# 테두리
_thin = Side(style="thin", color="D0D3CC")
BORDER = Border(_thin, _thin, _thin, _thin)

# 숫자 포맷
FMT_NUM1 = "0.0"       # 소수점 1자리
FMT_NUM2 = "0.00"      # 소수점 2자리
FMT_INT  = "#,##0"     # 정수 (천단위)
FMT_PCT  = "0.0%"      # 퍼센트
FMT_DATE = "YYYY-MM-DD"


# ---------------------------------------------------------------------------
# 셀 스타일 함수
# ---------------------------------------------------------------------------
def H(ws, r, c, v, fill=LIGHT, white=False, size=10):
    """헤더 셀: 굵은 글꼴, 중앙정렬, 배경색"""
    x = ws.cell(r, c, v)
    x.font = Font(
        name=FONT_NAME, bold=True, size=size,
        color=(WHITE if white else "1A1A1A"),
    )
    x.fill = PatternFill("solid", fgColor=fill)
    x.alignment = Alignment("center", "center", wrap_text=True)
    x.border = BORDER
    return x


def C(ws, r, c, v, fmt=None, left=False, fill=None):
    """데이터 셀: 일반 글꼴, 선택적 포맷/정렬/배경"""
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT_NAME, size=10)
    x.border = BORDER
    x.alignment = Alignment("left" if left else "center", "center")
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    return x


def set_col_widths(ws, widths: list[float], start_col: int = 1):
    """열 너비 일괄 설정.

    Args:
        widths: 열 너비 리스트
        start_col: 시작 열 번호 (1-based)
    """
    for i, w in enumerate(widths):
        col_letter = get_column_letter(start_col + i)
        ws.column_dimensions[col_letter].width = w


def write_header_row(ws, row: int, headers: list[str],
                     fill: str = LIGHT, start_col: int = 1):
    """한 행에 헤더를 일괄 입력"""
    for i, header in enumerate(headers):
        H(ws, row, start_col + i, header, fill=fill)


def write_data_row(ws, row: int, values: list,
                   fmt: str | None = None, fill: str | None = None,
                   start_col: int = 1, left_cols: set | None = None):
    """한 행에 데이터를 일괄 입력.

    Args:
        left_cols: 왼쪽 정렬할 열 인덱스(0-based) 집합
    """
    for i, val in enumerate(values):
        is_left = left_cols and i in left_cols
        C(ws, row, start_col + i, val, fmt=fmt, left=is_left, fill=fill)


def apply_title(ws, row: int, col: int, title: str,
                merge_end_col: int | None = None, size: int = 12):
    """시트 제목 셀"""
    x = H(ws, row, col, title, fill=GREEN, white=True, size=size)
    if merge_end_col and merge_end_col > col:
        ws.merge_cells(
            start_row=row, start_column=col,
            end_row=row, end_column=merge_end_col,
        )
