# -*- coding: utf-8 -*-
"""02-Cycle 4단계(G4) 예보 물수지 단위 테스트 (pytest)
  fcst_archive : 기대 강수(시각별 강수량 × 강수확률)
  fcst_wb      : 관측 물수지(관수 기록·관수 규칙·결측 채움), 발표별 예보 물수지(아침·저녁 출발, 세 경로, 참값),
                 관수 필요 예상일·판정·범위, 오차표(합치기·다른 해·월 칸 대체), 편향 보정 자리, 서비스 전망
  fcst_wb_report: 한 칸 물수지 수식 = fao56_core.wb_step
"""
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fao56_core import wb_step
from fcst_archive import daily_inputs
import fcst_wb as W

TS = pd.Timestamp
SOIL = dict(fc=0.22, wp=0.10, zr=1.0, p=0.5, taw=120.0, raw=60.0, dr0=0.0, ea=0.95)
KP1 = dict(bud=dt.date(2026, 1, 1), L=(1, 1, 500, 1), kc_ini=1.0, kc_mid=1.0, kc_end=1.0)   # Kc = 1 고정


def _obs(days, eto, rain):
    return pd.DataFrame({"date": pd.date_range(days[0], periods=len(eto)), "ETo_obs": eto, "rain": rain})


# ── 기대 강수 ─────────────────────────────────────────────────────────
def test_rain_exp_is_pcp_times_pop():
    from test_fcst import _archive, _hourly
    a, run = _archive()
    lead = a.hourly["PCP"][a.hourly["PCP"].issue == run].lead.tolist()
    a.hourly["POP"] = _hourly(str(run), lead, [60.0] * len(lead))
    out = daily_inputs(a, run, [run.normalize() + pd.Timedelta(days=k) for k in range(4)], 37.9, 127.74).set_index("lead_day")
    assert out.loc[0, "rain"] == pytest.approx(1.5) and out.loc[0, "rain_exp"] == pytest.approx(0.9)    # 1.5 mm × 60%
    assert out.loc[3, "rain"] == pytest.approx(4.5) and out.loc[3, "rain_exp"] == pytest.approx(2.7)    # 코드 1 → 4.5 mm × 60%
    assert out.loc[1, "rain_exp"] == 0.0


# ── 관측 물수지 ───────────────────────────────────────────────────────
def test_observed_wb_matches_wb_step_with_irrigation_log():
    eto, rain = [5.0] * 10, [0, 0, 0, 0, 30, 0, 0, 0, 0, 0]
    obs = _obs(["2026-05-01"], eto, rain)
    irr = {dt.date(2026, 5, 8): 20 / 0.95}                          # 공급 21.05 mm → 순 20 mm
    w = W.observed_wb(obs, KP1, SOIL, irrig=irr)
    dr = 0.0
    for r, p, d in zip(w.itertuples(), rain, obs.date):
        inet = 20.0 if d == TS("2026-05-08") else 0.0
        ks, ea_, dp, dr = wb_step(dr, p, 5.0, 120.0, 60.0, inet)
        assert r.Dr == pytest.approx(dr) and r.Ks == pytest.approx(ks)
    assert w.loc[w.date == TS("2026-05-08"), "I_net"].iloc[0] == pytest.approx(20.0)
    assert not w.filled.any() and not w.tainted.any()


def test_observed_wb_auto_irrigation_rule():
    soil = dict(SOIL, taw=40.0, raw=20.0)
    w = W.observed_wb(_obs(["2026-05-01"], [5.0] * 6, [0.0] * 6), KP1, soil, auto_irrigate=True)
    assert list(w.Dr.round(6)[:4]) == [5.0, 10.0, 15.0, 20.0]       # 4일째 끝 Dr = RAW
    assert w.I_net.iloc[4] == pytest.approx(20.0)                   # 5일째 Dr(20)만큼 관수
    assert w.Dr.iloc[4] == pytest.approx(5.0) and w.I_gross.iloc[4] == pytest.approx(20 / 0.95)


def test_observed_wb_fill_and_taint():
    obs = _obs(["2026-05-01"], [5.0, 5.0, np.nan, 5.0, 5.0, 5.0], [0, 0, 0, 0, 50, 0])
    obs = obs[obs.date != TS("2026-05-04")]                        # 5/4 행 자체가 없음
    fill = {TS("2026-05-03"): (4.0, 0.0), TS("2026-05-04"): (3.0, 2.0)}
    w = W.observed_wb(obs, KP1, SOIL, fill=fill).set_index("date")
    assert w.loc["2026-05-03", "ETo"] == 4.0 and w.loc["2026-05-03", "ETo_src"] == "예보"
    assert w.loc["2026-05-04", "P"] == 2.0 and w.loc["2026-05-04", "filled"]      # 행이 없는 날은 강수도 예보로
    assert w.loc["2026-05-02", "tainted"] == False and w.loc["2026-05-04", "tainted"]
    assert w.loc["2026-05-05", "Dr"] == 0.0 and not w.loc["2026-05-05", "tainted"]   # 큰 비로 Dr 0 → 영향 끝


# ── 발표별 예보 물수지 ────────────────────────────────────────────────
def _ft(D, etc=4.0):
    rows = []
    for rn, h, leads in (("아침", 2, range(0, 4)), ("저녁", 17, range(1, 5))):
        run = TS(D) + pd.Timedelta(hours=h)
        for k in leads:
            rain = 10.0 if k == 1 else 0.0
            rows.append(dict(run_name=rn, run=run, lead_day=k, target=TS(D) + pd.Timedelta(days=k), Kc=1.0,
                             ETc_main=etc, ETo_main=etc, rain=rain, rain_exp=0.6 * rain, pop_max=0.6 if k == 1 else 0.1,
                             ETc_pers=5.0, ETo_obs=5.0))
    return pd.DataFrame(rows)


def test_forecast_runs_morning_and_evening_paths():
    D = "2026-06-10"
    obs = _obs(["2026-06-01"], [5.0] * 15, [0.0] * 15)
    owb = W.observed_wb(obs, KP1, SOIL)
    ft = _ft(D)
    runs = W.forecast_runs(ft, owb, SOIL)
    dr0 = float(owb.set_index("date").loc[TS(D) - pd.Timedelta(days=1), "Dr"])     # 9일 × 5 = 45
    assert dr0 == pytest.approx(45.0)
    m = runs[runs.run_name == "아침"].sort_values("order")
    exp = W.path(dr0, [4.0] * 4, [0.0, 6.0, 0.0, 0.0], SOIL)                         # 중심: 기대 강수
    assert list(m.Dr_center) == pytest.approx(exp)
    assert list(m.Dr_true) == pytest.approx(W.path(dr0, [5.0] * 4, [0.0] * 4, SOIL))
    assert list(m.Dr_none) == pytest.approx(W.path(dr0, [4.0] * 4, [0.0] * 4, SOIL))
    rel = W.ERR_DEFAULT_REL
    assert list(m.Dr_early) == pytest.approx(W.path(dr0, [4.0 * (1 + rel)] * 4, [0.0] * 4, SOIL))
    assert list(m.Dr_late) == pytest.approx(W.path(dr0, [4.0 * (1 - rel)] * 4, [0.0, 10.0, 0.0, 0.0], SOIL))
    e = runs[runs.run_name == "저녁"].sort_values("order")
    start = wb_step(dr0, 0.0, 4.0, 120.0, 60.0)[3]                                  # 그날 D는 아침 D+0 예보로
    assert e.dr_start.iloc[0] == pytest.approx(start)
    assert list(e.Dr_center) == pytest.approx(W.path(dr0, [4.0] * 5, [0.0, 6.0, 0.0, 0.0, 0.0], SOIL)[1:])
    assert e.dr_start_true.iloc[0] == pytest.approx(wb_step(dr0, 0.0, 5.0, 120.0, 60.0)[3])
    assert runs.ok.all() and list(m.main) == [True, True, True, False]


def test_forecast_runs_excludes_rows_after_filled_target():
    D = "2026-06-10"
    eto = [5.0] * 15
    eto[10] = np.nan                                                                 # 6/11 관측 없음
    obs = _obs(["2026-06-01"], eto, [0.0] * 15)
    ft = _ft(D)
    owb = W.observed_wb(obs, KP1, SOIL, fill=W.fill_from_forecast(ft))
    runs = W.forecast_runs(ft, owb, SOIL)
    m = runs[runs.run_name == "아침"].sort_values("order")
    assert list(m.ok) == [True, False, False, False]
    assert set(m.reason[~m.ok]) == {"예보 기간에 관측 결측일(예보로 채움)"}


def test_first_need_and_categories():
    assert W.first_need([10, 70, 80], 60) == 2 and W.first_need([10, 20], 60) is None
    assert W.first_need([70, 80], 60, orders=[2, 3]) == 2
    assert W.need_category(2, 2) == "같은 날" and W.need_category(1, 2) == "하루 빠름"
    assert W.need_category(3, 1) == "이틀 이상 늦음" and W.need_category(None, 2) == "실제만 필요(놓침)"
    assert W.need_category(2, None) == "예보만 필요(헛경보)" and W.need_category(None, None) == "둘 다 기간 안 필요 없음"


def test_first_need_eval_range_coverage():
    base = dict(run_name="아침", ok=True, dr_start=50.0, dr_start_true=50.0, lead_day=0)
    rows = []
    # 발표 1: 중심 없음, 빠르면 2, 참값 2 → 놓침이지만 범위 안
    for k, (c, e, l, t) in enumerate([(52, 55, 50, 55), (55, 62, 51, 61), (58, 70, 52, 66)], 1):
        rows.append(dict(base, run=TS("2026-06-01 02:00"), order=k, Dr_center=c, Dr_early=e, Dr_late=l, Dr_true=t))
    # 발표 2: 참값이 빠르면보다 앞
    for k, (c, e, l, t) in enumerate([(52, 55, 50, 61), (55, 62, 51, 66), (58, 70, 52, 70)], 1):
        rows.append(dict(base, run=TS("2026-06-02 02:00"), order=k, Dr_center=c, Dr_early=e, Dr_late=l, Dr_true=t))
    fe = W.first_need_eval(pd.DataFrame(rows), 60.0)
    a, b = fe.iloc[0], fe.iloc[1]
    assert (a.fcst, a.early, a.late, a.true, a.category, a.covered) == (None, 2, None, 2, "실제만 필요(놓침)", True)
    assert b.true == 1 and b.early == 2 and b.true_before_early and not b.covered
    s = W.first_need_summary(fe).iloc[0]
    assert s.n_range_event == 2 and s.coverage == pytest.approx(0.5) and s.true_before_early == 1


# ── 오차표 ────────────────────────────────────────────────────────────
def _ft_err(year, bias):
    rows = []
    for d in pd.date_range(f"{year}-05-01", f"{year}-06-30"):
        for rn, h, leads in (("아침", 2, range(0, 4)), ("저녁", 17, range(1, 5))):
            for k in leads:
                obs = 4.0 + (d.day % 3)
                rows.append(dict(run_name=rn, run=d + pd.Timedelta(hours=h), lead_day=k, target=d + pd.Timedelta(days=k),
                                 ETo_main=obs + bias + (0.5 if k % 2 else -0.5), ETo_obs=obs))
    return pd.DataFrame(rows)


def test_error_table_pool_exclude_and_lookup(tmp_path):
    rows = pd.concat([W.eto_error_rows(_ft_err(2025, 0.2), "101", 2025), W.eto_error_rows(_ft_err(2026, -0.2), "101", 2026)])
    p = tmp_path / "err.csv"
    W.save_error_table(rows, p)
    W.save_error_table(W.eto_error_rows(_ft_err(2026, -0.2), "101", 2026), p)       # 같은 지점·해는 바꿔 넣음
    back = W.load_error_rows(str(p))
    assert len(back) == len(rows)
    pool = W.pooled_errors(back, "101")
    one = pool[(pool.kind == "day") & (pool.run_name == "아침") & (pool.lead_day == 1) & (pool.month == 0)].iloc[0]
    assert one.mbe == pytest.approx(0.5, abs=1e-9) and one.years == "2025,2026"     # +0.2+0.5, −0.2+0.5 → 평균 0.5
    other = W.pooled_errors(back, "101", exclude_year=2026)
    o1 = other[(other.kind == "day") & (other.run_name == "아침") & (other.lead_day == 1) & (other.month == 0)].iloc[0]
    assert o1.mbe == pytest.approx(0.7) and o1.years == "2025"
    rmse, rel, n, src = W.err_lookup(pool, "day", "아침", 1, 5)
    assert src == "월별" and n >= 15 and rel == pytest.approx(rmse / pool[(pool.kind == "day") & (pool.run_name == "아침")
                                                                           & (pool.lead_day == 1) & (pool.month == 5)].obs_mean.iloc[0])
    assert W.err_lookup(pool, "day", "아침", 1, 9)[3] == "전체 월"                   # 9월 칸 없음 → 전체 월
    assert W.err_lookup(None, "day", "아침", 1, 5)[1] == W.ERR_DEFAULT_REL
    assert W.pooled_errors(back, "999").stn_src.iloc[0].startswith("전 지점")        # 지점 행이 없으면 다른 지점 오차
    c3 = pool[(pool.kind == "cum3") & (pool.run_name == "저녁")]
    assert set(c3.lead_day) == {1}


def test_apply_bias_hook():
    ft = _ft("2026-06-10")
    assert W.apply_bias(ft) is ft                                                   # #10: 기본은 보정 없음
    t = pd.DataFrame(dict(run_name=["아침"], lead_day=[0], factor=[1.1]))
    b = W.apply_bias(ft, t)
    m = (b.run_name == "아침") & (b.lead_day == 0)
    assert np.allclose(b.loc[m, "ETc_main"], 4.4) and np.allclose(b.loc[~m, "ETc_main"], 4.0)


# ── 서비스 전망 ───────────────────────────────────────────────────────
def test_service_outlook_need_and_advice():
    D = "2026-06-10"
    obs = _obs(["2026-05-25"], [5.0] * 20, [0.0] * 20)                              # 6/9 끝 Dr ≥ RAW (Ks < 1로 80보다 조금 작음)
    owb = W.observed_wb(obs, KP1, SOIL)
    ol = W.service_outlook(_ft(D), owb, SOIL, None, "아침", TS(D) + pd.Timedelta(hours=2))
    start = float(owb.set_index("date").loc[TS("2026-06-09"), "Dr"])
    assert start > 60 and ol["start"] == pytest.approx(start) and ol["need"]["center"] == 0
    assert ol["advice"][0] == pytest.approx(start) and ol["advice"][1] == pytest.approx(start / 0.95)
    obs2 = _obs(["2026-05-31"], [5.0] * 20, [0.0] * 20)                             # 6/9 끝 Dr 50 → 54, 52(비 6), 56, 60
    ol2 = W.service_outlook(_ft(D), W.observed_wb(obs2, KP1, SOIL), SOIL, None, "저녁", TS(D) + pd.Timedelta(hours=17))
    d = ol2["days"]
    k = ol2["need"]["center"]
    assert k == 3 and d.Dr_center.iloc[2] == pytest.approx(60.0) and ol2["need"]["early"] < 3
    assert ol2["cum3"][0] == pytest.approx(12.0) and len(ol2["pre"]) == 1


# ── 엑셀 한 칸 물수지 수식 = wb_step ─────────────────────────────────────
def _eval(expr, cells):
    import re
    py = expr.lstrip("=")
    for k, v in cells.items():
        py = re.sub(rf"\b{k}\b", repr(v), py)
    return eval(py.replace("^", "**"), {"MIN": min, "MAX": max, "IF": lambda c, a, b: a if c else b})


@pytest.mark.parametrize("prev,p,etc", [(0.0, 0.0, 5.0), (55.0, 3.0, 4.0), (70.0, 30.0, 6.0), (90.0, 0.0, 5.0), (10.0, 40.0, 3.0)])
def test_excel_wb_formula_matches_wb_step(prev, p, etc):
    from fcst_wb_report import _wb_formula
    f = _wb_formula("PREV", "RAIN", "ETC", {"RAW": "60", "TAW": "120"})
    assert _eval(f, {"PREV": prev, "RAIN": p, "ETC": etc}) == pytest.approx(wb_step(prev, p, etc, 120.0, 60.0)[3])


# ── 서비스 엑셀: 과습 상태·날짜 선택 ──────────────────────────────────────
def test_soil_status_wet_and_dry():
    s = W.soil_status
    assert s(0.0, 60, kc=1.0) == "과습 경고" and s(8.9e-16, 60, kc=1.0) == "과습 경고"       # 끝 Dr 0 (부동소수 오차 포함)
    assert s(0.3, 60, kc=1.0) == "토양 수분 충분" and s(29.9, 60, kc=1.0) == "토양 수분 충분"          # < RAW ÷ 2 (포장용수량과 RAW의 중간)
    assert s(30.0, 60, kc=1.0) == "주의" and s(60.0, 60, kc=1.0) == "관수 필요"
    assert s(6.1, 60, kc=1.0, wet_caution=6.0) == "안전" and s(6.0, 60, kc=1.0, wet_caution=6.0) == "안전"   # 기준을 낮추면 '안전'이 다시 나옴
    assert s(0.0, 60, kc=0.0) == "안전" and s(0.0, 60, kc=None) == "안전"                    # 휴면기(Kc 0)는 과습 판정 안 함
    assert s(0.0, 60, kc=1.0, observed=False) == "안전"                                     # 예보로 진행한 날은 판정 안 함
    assert s(10.0, 60, kc=1.0, wet_caution=12.0) == "토양 수분 충분" and s(np.nan, 60, kc=1.0) == "자료 없음"


def _service_sv(D, name, rain_day=None, lag=False, days=20):
    """서비스 엑셀용 최소 sv: ETo 5·Kc 1, rain_day에 큰 비(Dr → 0). lag면 전날 관측이 아직 없음"""
    hour = 2 if name == "아침" else 17
    start = TS(D) - pd.Timedelta(days=days)
    rain = [200.0 if (rain_day and start + pd.Timedelta(days=i) == TS(rain_day)) else 0.0 for i in range(days)]
    obs = _obs([start], [5.0] * days, rain)
    ft = _ft(D)
    if lag:
        ft = pd.concat([_ft(str((TS(D) - pd.Timedelta(days=1)).date())), ft], ignore_index=True)
        obs = obs[obs.date < TS(D) - pd.Timedelta(days=1)]
        owb = W.observed_wb(obs, KP1, SOIL, fill=W.fill_from_forecast(ft), end=TS(D) - pd.Timedelta(days=1))
    else:
        owb = W.observed_wb(obs, KP1, SOIL)
    last = W.last_observed(owb, TS(D))
    ol = W.service_outlook(ft, owb, SOIL, None, name, TS(D) + pd.Timedelta(hours=hour),
                           obs_last=last if last < TS(D) - pd.Timedelta(days=1) else None)
    err = W.pooled_errors(W.load_error_rows(os.path.join(os.path.dirname(__file__), "..", "fcst_error_table.csv")), "101")
    return dict(outlook=ol, owb=owb, soil=SOIL, err=err, stn="101", stn_name="춘천", grid="73_134", irrig={}, recent=None,
                wet=W.wet_info(owb, ol["prev"], SOIL["raw"]))


def test_wet_info_uses_last_observed_day():
    sv = _service_sv("2026-06-10", "아침", rain_day="2026-06-09")
    w = sv["wet"]
    assert w["date"] == TS("2026-06-09") and w["status"] == "과습 경고" and w["dr"] == 0.0 and w["zero_days"] == 1
    o = sv["owb"].set_index("date")
    assert w["dp"] == pytest.approx(float(o.loc[TS("2026-06-09"), "DP"])) and w["dp"] > 100 and w["dp7"] == pytest.approx(w["dp"])
    sv2 = _service_sv("2026-06-10", "아침", rain_day="2026-06-08")                           # 하루 지남: Dr 5 ≤ 6
    assert sv2["wet"]["status"] == "토양 수분 충분" and sv2["wet"]["dr"] == pytest.approx(5.0)
    sv3 = _service_sv("2026-06-10", "아침", rain_day="2026-06-08", lag=True)                 # 6/9 관측 없음 → 6/8(관측 마지막 날) 기준
    assert sv3["wet"]["date"] == TS("2026-06-08") and sv3["wet"]["status"] == "과습 경고"


def test_service_workbook_date_picker_layout(tmp_path):
    import openpyxl
    from fcst_wb_report import SEL, TL, build_service_workbook
    out = build_service_workbook(_service_sv("2026-06-10", "저녁", rain_day="2026-06-09"), str(tmp_path / "s.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:4] == ["관수 전망", "예보 물수지", "관측 물수지", TL]
    ws = wb["관수 전망"]
    assert SEL == "'관수 전망'!$B$5" and ws["B5"].value == dt.datetime(2026, 6, 10)             # 기본 조회일 = 발표일
    dv = ws.data_validations.dataValidation[0]
    assert dv.type == "list" and str(dv.sqref) == "B5" and dv.formula1.startswith(f"={TL}!$A$3:") and not dv.showErrorMessage
    labels = [ws.cell(r, 1).value for r in range(1, 30)]
    assert labels.index("날짜 선택") < labels.index("지금 토양 상태 (조회일 전날 끝)")            # 날짜 선택이 토양 상태 위
    st = labels.index("상태") + 1
    assert "과습 경고" in ws.cell(st, 2).value and "토양 수분 충분" in ws.cell(st, 2).value and "자료 없음" in ws.cell(st, 2).value
    assert any(str(wb["설정"].cell(r, 1).value).startswith("토양 수분 충분 기준") for r in range(2, 14))
    tl = wb[TL]
    assert tl["P2"].value == 1 and tl.cell(tl.max_row, 16).value == 0 and tl["S3"].value == 1    # 관측 행 → 예보 행, 저녁 = +1


def _recalc(path, sel, tmp_path, tag):
    """조회일을 바꿔 LibreOffice로 다시 계산한 '관수 전망' 시트 {A열 라벨: B열 값}과 통합 문서"""
    import shutil
    import subprocess
    import openpyxl
    wb = openpyxl.load_workbook(path)
    if sel:
        wb["관수 전망"]["B5"] = dt.datetime.strptime(sel, "%Y-%m-%d")
    d = tmp_path / tag
    d.mkdir()
    wb.save(str(d / "in.xlsx"))
    subprocess.run([shutil.which("soffice"), "--headless", "--convert-to", "xlsx", "--outdir", str(d / "o"), str(d / "in.xlsx")],
                   capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(str(d / "o" / "in.xlsx"), data_only=True)
    ws = wb["관수 전망"]
    return {str(ws.cell(r, 1).value).strip(): ws.cell(r, 2).value for r in range(4, 30) if ws.cell(r, 1).value is not None}, wb


@pytest.mark.skipif(__import__("shutil").which("soffice") is None, reason="LibreOffice 없음")
def test_service_workbook_selected_date_matches_python(tmp_path):
    from fcst_wb_report import build_service_workbook
    sv = _service_sv("2026-06-10", "아침", rain_day="2026-06-06")
    out = build_service_workbook(sv, str(tmp_path / "s.xlsx"))
    o, ol = sv["owb"].set_index("date"), sv["outlook"]
    key = lambda v, head: next(x for k, x in v.items() if k.startswith(head))
    for sel in (None, "2026-06-07", "2026-06-08", "2026-06-09", "2026-06-12", "2026-05-30", "2026-07-15", "2026-05-21"):
        v, wb = _recalc(out, sel, tmp_path, f"t{sel}")
        assert not [c.value for ws in wb for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
        S = TS(sel or "2026-06-10")
        prev = S - pd.Timedelta(days=1)
        dr, st = key(v, "전날("), v["상태"]
        if prev < o.index.min() or prev > ol["days"].target.max():         # 표에 없는 날짜
            assert dr == "자료 없음" and st == "자료 없음"
        elif prev in o.index:                                              # 지난 날짜·발표일: 관측 물수지
            assert dr == pytest.approx(float(o.loc[prev, "Dr"]), abs=1e-9)
            assert st == W.soil_status(float(o.loc[prev, "Dr"]), SOIL["raw"], 1.0)
        else:                                                              # 발표일 뒤: 이 발표의 예보(과습 판정 안 함)
            d = ol["days"].set_index("target")
            assert dr == pytest.approx(float(d.loc[prev, "Dr_center"]), abs=1e-9)
            assert st == W.soil_status(float(d.loc[prev, "Dr_center"]), SOIL["raw"], 1.0, observed=False)
        etc3 = v["작물 증발산 ETc 3일 합"]
        win = [S + pd.Timedelta(days=i) for i in range(3)]
        etc = {**{t: float(o.loc[t, "ETc"]) for t in o.index}, **{t: 4.0 for t in ol["days"].target}}
        assert etc3 == (pytest.approx(sum(etc[t] for t in win)) if all(t in etc for t in win) else "자료 없음")
    v, _ = _recalc(out, None, tmp_path, "d")                               # 발표일: 기존 전망과 같은 값
    assert key(v, "전날(") == pytest.approx(ol["dr_obs_prev"]) and v["작물 증발산 ETc 3일 합"] == pytest.approx(ol["cum3"][0])
    assert v["상태"] == sv["wet"]["status"] == "토양 수분 충분"                   # 6/6 비 뒤 3일: Dr 15 < 30
    assert _recalc(out, "2026-06-07", tmp_path, "w")[0]["상태"] == "과습 경고"
    assert _recalc(out, "2026-06-08", tmp_path, "c")[0]["상태"] == "토양 수분 충분"


@pytest.mark.skipif(__import__("shutil").which("soffice") is None, reason="LibreOffice 없음")
def test_service_workbook_obs_lag_wet_from_last_observed(tmp_path):
    """전날 관측이 아직 없을 때: 전날 끝 Dr은 예보로 먼저 진행한 값, 과습은 관측 마지막 날 기준"""
    from fcst_wb_report import build_service_workbook
    sv = _service_sv("2026-06-10", "아침", rain_day="2026-06-08", lag=True)
    v, _ = _recalc(build_service_workbook(sv, str(tmp_path / "s.xlsx")), None, tmp_path, "lag")
    dr = next(x for k, x in v.items() if k.startswith("전날("))
    assert dr == pytest.approx(sv["outlook"]["pre"][0]["Dr_center"]) and dr > 0
    assert v["상태"] == sv["wet"]["status"] == "과습 경고"
    ev = _service_sv("2026-06-10", "저녁", rain_day="2026-06-09")            # 저녁: 주 지표는 조회일 다음 날부터
    v, _ = _recalc(build_service_workbook(ev, str(tmp_path / "e.xlsx")), None, tmp_path, "ev")
    assert v["상태"] == "과습 경고" and v["작물 증발산 ETc 3일 합"] == pytest.approx(ev["outlook"]["cum3"][0])
    assert next(x for k, x in v.items() if k.startswith("조회일(")) == pytest.approx(ev["outlook"]["start"])
