#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ops_report.py — H3(운영 수집) 보고서 엑셀 (02-Cycle 5단계, G5)

  시트: 요약 / 슬롯 / 수집 시도 / 전날 관측
  - 슬롯: 예정 슬롯(아침 02시·저녁 17시)마다 처음 실행 결과(정시 여부)와 마지막 결과, 대체 발표, Rs 방법, 전날 관측
  - 요약의 정시 수집 성공률·백업 자동 전환율은 슬롯 시트를 세는 수식(COUNTIF)이라 슬롯 시트를 고치면 다시 계산됨
  - 기준(노란 칸): 성공률 ≥ 99%, 백업 자동 전환 100%, 판정 표본 ≥ 100회 (THEORY 9장 ◆ 운영 수집)
"""
import pandas as pd
from openpyxl import Workbook

from fao56_core import safe_save
from fcst_report import C, DATE, DTM, GREEN, H, LIGHT, T, YEL, widths

OKF, BADF = "D8EAD3", "F4C7B8"


def build_h3_workbook(summ, tab, obs, collect_log, out, cfg=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    # 슬롯 시트 (요약 수식이 참조)
    ss = wb.create_sheet("슬롯")
    heads = ["발표 (슬롯)", "구분", "처음 결과", "정시 성공", "백업 엑셀 자동", "마지막 결과", "계산한 발표", "Rs 방법", "전날 관측", "시도", "오류"]
    for c, h in enumerate(heads, 1):
        H(ss, 1, c, h)
    for i, r in enumerate(tab.itertuples(index=False), 2):
        C(ss, i, 1, pd.Timestamp(r.slot_run).to_pydatetime(), DTM)
        C(ss, i, 2, r.run_name)
        C(ss, i, 3, r.first_result, fill=(OKF if r.on_time else BADF))
        C(ss, i, 4, "예" if r.on_time else "아니오")
        C(ss, i, 5, "" if r.on_time else ("예" if r.backup_excel else "아니오"))
        C(ss, i, 6, r.final_result)
        for c, k in ((7, "src_run"), (8, "rs_method"), (9, "obs_prev"), (10, "attempts"), (11, "error")):
            v = getattr(r, k, "")
            C(ss, i, c, "" if pd.isna(v) else v)
    n = len(tab) + 1
    ss.freeze_panes = "A2"
    widths(ss, {"A": 17, "B": 6, "C": 18, "D": 9, "E": 12, "F": 18, "G": 17, "H": 8, "I": 12, "J": 6, "K": 50})

    T(ws, 1, 1, f"H3 운영 수집 점검 — ASOS {(cfg or {}).get('stn', '')} · 격자 {(cfg or {}).get('grid', '')} · "
                f"{summ['since']:%Y-%m-%d} ~ {summ['until']:%Y-%m-%d}", size=14, bold=True, color=GREEN)
    T(ws, 2, 1, "정시 수집 성공 = 마감(발표 + 40분) 안에 서비스 발표(02·17시)를 8요소 완결로 받음. 실행 기록이 없는 슬롯은 '실행 안 됨'(실패)으로 셈",
      size=9, color="555555")
    r = 4
    rows = [("예정 슬롯 (회)", f"=COUNTA(슬롯!A2:A{n})", "0", None),
            ("정시 수집 성공 (회)", f"=COUNTIF(슬롯!D2:D{n},\"예\")", "0", None),
            ("★ 정시 수집 성공률", f"=IF(B{r}=0,\"\",B{r + 1}/B{r})", "0.0%", None),
            ("  기준", 0.99, "0%", YEL),
            ("정시 수집 실패 (회)", f"=B{r}-B{r + 1}", "0", None),
            ("  그중 백업 엑셀 자동 생성 (회)", f"=COUNTIF(슬롯!E2:E{n},\"예\")", "0", None),
            ("★ 백업 자동 전환율", f"=IF(B{r + 4}=0,\"실패 없음\",B{r + 5}/B{r + 4})", "0%", None),
            ("  기준", 1.0, "0%", YEL),
            ("판정 표본 기준 (회)", 100, "0", YEL),
            ("★ H3 판정", f"=IF(B{r}=0,\"\",IF(AND(B{r + 2}>=B{r + 3},OR(B{r + 4}=0,B{r + 6}>=B{r + 7})),\"충족\",\"미달\")"
                          f"&IF(B{r}<B{r + 8},\" (표본 \"&B{r}&\"회: 중간 점검 — 실패 0회여야 기준과 맞음)\",\"\"))", None, None)]
    for k, v, f, fill in rows:
        H(ws, r, 1, k, LIGHT, white=False)
        C(ws, r, 2, v, f, bold=k.startswith("★"), fill=fill)
        r += 1
    r += 1
    H(ws, r, 1, "처음 결과별 슬롯 수", GREEN); H(ws, r, 2, "회"); r += 1
    for k, v in sorted(summ.get("counts", {}).items(), key=lambda x: -x[1]):
        H(ws, r, 1, k, LIGHT, white=False); C(ws, r, 2, f"=COUNTIF(슬롯!C2:C{n},\"{k}\")", "0")
        r += 1
    r += 1
    H(ws, r, 1, "전날 ASOS 일자료", GREEN); H(ws, r, 2, "값"); r += 1
    items = [("아침 슬롯 시점에 조회된 비율", summ.get("obs_morning_ok"), "0%"),
             ("처음 조회된 시각 (자정 뒤 시간, 중앙값)", summ.get("obs_lag_median_h"), "0.0")]
    for k, v, f in items:
        H(ws, r, 1, k, LIGHT, white=False); C(ws, r, 2, None if v is None or pd.isna(v) else float(v), f)
        r += 1
    T(ws, r + 1, 1, "처음 조회된 시각은 수집(asos_d1)·점검(asos_probe) 기록 중 가장 이른 성공입니다. probe-asos를 매시 실행하면 정확해집니다.",
      size=9, color="555555")
    widths(ws, {"A": 40, "B": 26})

    # 수집 시도
    sc = wb.create_sheet("수집 시도")
    if len(collect_log):
        cols = list(collect_log.columns)
        for c, h in enumerate(cols, 1):
            H(sc, 1, c, h)
        for i, row in enumerate(collect_log.itertuples(index=False), 2):
            for c, v in enumerate(row, 1):
                C(sc, i, c, "" if pd.isna(v) else v)
        sc.freeze_panes = "A2"
        widths(sc, {chr(64 + i): w for i, w in enumerate([19, 17, 10, 17, 7, 6, 6, 22, 8, 8, 10, 10, 10, 9, 40], 1)})
    # 전날 관측
    so = wb.create_sheet("전날 관측")
    for c, h in enumerate(["관측일 (D−1)", "처음 조회된 시각", "자정 뒤 (시간)", "아침 슬롯 시점 결과"], 1):
        H(so, 1, c, h)
    for i, r_ in enumerate(obs.itertuples(index=False) if len(obs) else [], 2):
        C(so, i, 1, pd.Timestamp(r_.target).to_pydatetime(), DATE)
        C(so, i, 2, pd.Timestamp(r_.first_ok).to_pydatetime() if not pd.isna(r_.first_ok) else None, DTM)
        C(so, i, 3, None if pd.isna(r_.lag_h) else float(r_.lag_h), "0.0")
        C(so, i, 4, "" if pd.isna(r_.morning_status) else r_.morning_status)
    widths(so, {"A": 13, "B": 19, "C": 13, "D": 16})
    return safe_save(wb, out)
