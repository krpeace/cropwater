"""
solar_multi.py — 다지점 일사·일조 종합 분석 엑셀 생성

하나의 엑셀 파일 안에 6개 시트로 모든 분석 결과를 출력한다:
  1. 설정          — 조회 조건 요약 (노란칸 편집 가능)
  2. 권역별    — 권역 선택 콤보 + 동기간 비교
  3. 지점별 집계    — 지점 선택 콤보 + 동기간 비교
  4. 기간별(월별)   — 권역별 열 구성 + 연간/월별 행 구성
  5. 년도별         — 지표별 권역 행 + 월별 열 구성
  6. 원데이터       — 지점별 일별 전체 데이터 (맨 뒤)

사용법:
  python solar_multi.py                    # 전체 66개 지점
  python solar_multi.py --region 서울경기   # 특정 권역만
  python solar_multi.py --station 108      # 특정 지점만
"""

import argparse
import sys
from datetime import date, timedelta
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import stations
import solar_core
from openpyxl.styles import Font, PatternFill, Alignment

from excel_util import (
    H, C, FONT_NAME, GREEN, BLUE, BROWN, LIGHT, YEL, WARN,
    LIGHT_BLUE, LIGHT_YEL, BORDER, FMT_NUM1, FMT_NUM2, FMT_INT,
    WHITE, set_col_widths, write_header_row, write_data_row, apply_title,
)

# ---------------------------------------------------------------------------
# 동기간 비교 유틸
# ---------------------------------------------------------------------------

def _filter_same_period(parsed_rows, start_md, end_md):
    """MM-DD 기반 동기간 필터.

    start_md, end_md: (month, day) 튜플
    연도에 관계없이 start_md <= MM-DD <= end_md 인 행만 반환.
    """
    filtered = []
    for r in parsed_rows:
        md = (r["month"], r["day"])
        if start_md <= md <= end_md:
            filtered.append(r)
    return filtered


def _period_md_from_dates(start_date, end_date):
    """date → (month, day) 튜플 변환"""
    return (start_date.month, start_date.day), (end_date.month, end_date.day)


# ---------------------------------------------------------------------------
# 집계 유틸
# ---------------------------------------------------------------------------

def _daily_average_rows(rows):
    """다중 지점 일별 데이터를 일별 평균 행으로 변환.

    같은 날의 여러 지점 값을 평균하여 하루에 한 행으로 만든다.
    → 권역/전국 적산 계산 시 지점 수에 비례하여 과다 합산되는 문제 방지.
    → 기상청 '전국' 통계와 동일한 방식 (일별 지점평균 → 월 집계).
    """
    metric_keys = ["sum_gsr", "sum_ss_hr", "ss_dur", "ss_rate",
                    "hr1_max_icsr", "avg_ta", "max_ta", "min_ta", "sum_rn"]

    daily = defaultdict(lambda: defaultdict(list))
    for r in rows:
        day_key = (r["year"], r["month"], r["day"])
        for key in metric_keys:
            v = r.get(key)
            if v is not None:
                daily[day_key][key].append(v)

    result = []
    for (year, month, day), metrics in sorted(daily.items()):
        row = {"year": year, "month": month, "day": day}
        for key, vals in metrics.items():
            row[key] = sum(vals) / len(vals)
        result.append(row)

    return result


def _monthly_metric_by_year(rows, metric_key, agg_type):
    """일별 데이터 → {year: {month: value}} 형태로 집계.

    agg_type "avg": 월 내 일평균
    agg_type "cumsum": 월 내 합계

    주의: 다중 지점 데이터를 넘기면 지점 수만큼 과다 합산됨.
    → 권역/전국 집계 시 반드시 _daily_average_rows()로
      일별 평균 행으로 변환한 후 호출할 것.
    """
    buckets = defaultdict(lambda: defaultdict(list))
    for r in rows:
        v = r.get(metric_key)
        if v is not None:
            buckets[r["year"]][r["month"]].append(v)

    result = {}
    for year, months in buckets.items():
        result[year] = {}
        for month, vals in months.items():
            if agg_type == "avg":
                result[year][month] = sum(vals) / len(vals)
            else:  # cumsum
                result[year][month] = sum(vals)
    return result


def _period_metric_by_year(rows, metric_key, agg_type):
    """일별 데이터 → {year: value} 기간 전체 집계.

    agg_type "avg": 기간 내 일평균
    agg_type "cumsum": 기간 내 합계

    주의: 다중 지점 데이터를 넘기면 지점 수만큼 과다 합산됨.
    → 권역/전국 집계 시 반드시 _daily_average_rows()로
      일별 평균 행으로 변환한 후 호출할 것.
    """
    buckets = defaultdict(list)
    for r in rows:
        v = r.get(metric_key)
        if v is not None:
            buckets[r["year"]].append(v)

    result = {}
    for year, vals in buckets.items():
        if agg_type == "avg":
            result[year] = sum(vals) / len(vals)
        else:
            result[year] = sum(vals)
    return result


def _calc_normal_period(year_vals, normal_start, normal_end):
    """연도별 값 딕셔너리에서 평년 평균 계산"""
    vals = [year_vals[y] for y in range(normal_start, normal_end + 1)
            if y in year_vals]
    return sum(vals) / len(vals) if vals else None


def _calc_normal_monthly(monthly_by_year, normal_start, normal_end):
    """연도별 월별 값에서 평년 월별 평균 계산"""
    normal_vals = {}
    for month in range(1, 13):
        yr_vals = []
        for yr in range(normal_start, normal_end + 1):
            v = monthly_by_year.get(yr, {}).get(month)
            if v is not None:
                yr_vals.append(v)
        normal_vals[month] = (sum(yr_vals) / len(yr_vals)) if yr_vals else None
    return normal_vals


# ---------------------------------------------------------------------------
# 시트 1: 설정
# ---------------------------------------------------------------------------
def _sheet_settings(wb, start_date, end_date, ref_year, normal_years,
                    target_regions, target_stations):
    ws = wb.active
    ws.title = "설정"
    set_col_widths(ws, [18, 40, 20])

    apply_title(ws, 1, 1, "일사·일조 분석 설정", merge_end_col=3, size=12)

    # 안내 문구
    C(ws, 2, 1, "(노란칸은 수정 가능)", left=True)

    # 설정 항목 — 올해는 노란칸 (편집 가능)
    row = 4
    H(ws, row, 1, "올해")
    C(ws, row, 2, ref_year, fill=YEL)  # 편집 가능
    row += 1

    # 작년 = 올해 - 1 (수식)
    H(ws, row, 1, "작년")
    C(ws, row, 2, None)
    ws.cell(row, 2).value = "=B4-1"
    ws.cell(row, 2).number_format = "0"
    row += 1

    # 평년 기간 시작 = 올해 - 평년기간
    H(ws, row, 1, "평년 기간(년)")
    C(ws, row, 2, normal_years, fill=YEL)  # 편집 가능
    row += 1

    # 평년 범위 표시 (수식)
    H(ws, row, 1, "평년 범위")
    C(ws, row, 2, None, left=True)
    ws.cell(row, 2).value = f'=B4-B6&" ~ "&B4-1&" ("&B6&"년 평균)"'
    row += 1

    # 분석 기간 — 행 8, 9 고정 (다른 시트에서 설정!B8, 설정!B9 참조)
    # ※ 읽기전용: 기간 변경은 스크립트 재실행 필요
    H(ws, row, 1, "분석 시작일")       # row=8
    C(ws, row, 2, start_date.strftime("%m-%d"))
    ws.cell(row, 2).number_format = "@"   # 텍스트 형식 (날짜 해석 방지)
    C(ws, row, 3, "← 변경 시 스크립트 재실행 필요", left=True)
    ws.cell(row, 3).font = Font(name=FONT_NAME, size=9, color="999999")
    row += 1
    H(ws, row, 1, "분석 종료일")       # row=9
    C(ws, row, 2, end_date.strftime("%m-%d"))
    ws.cell(row, 2).number_format = "@"   # 텍스트 형식 (날짜 해석 방지)
    row += 1

    # 고정 정보
    row += 1
    items = [
        ("데이터 소스", "기상청 ASOS 일자료 (공공데이터포털)"),
        ("대상 권역", ", ".join(target_regions) if target_regions else "전체"),
        ("대상 지점 수", str(len(target_stations))),
        ("생성일", date.today().isoformat()),
    ]
    for label, value in items:
        H(ws, row, 1, label)
        C(ws, row, 2, value, left=True)
        row += 1

    # 지점 목록
    row += 1
    H(ws, row, 1, "지점번호")
    H(ws, row, 2, "지점명")
    H(ws, row, 3, "권역")
    for stn_id in target_stations:
        row += 1
        C(ws, row, 1, stn_id)
        C(ws, row, 2, stations.station_name(stn_id), left=True)
        C(ws, row, 3, stations.region_of(stn_id) or "", left=True)


# ---------------------------------------------------------------------------
# 시트 6 (맨 뒤): 원데이터
# ---------------------------------------------------------------------------
def _sheet_raw_data(wb, all_parsed: dict[int, list[dict]]):
    ws = wb.create_sheet("원데이터")

    headers = [
        "날짜", "지점번호", "지점명", "권역",
        "합계일사(MJ/m²)", "합계일조(hr)", "가조시간(hr)",
        "일조율(%)", "1h최다일사(MJ/m²)",
        "평균기온(℃)", "최고기온(℃)", "최저기온(℃)", "일강수량(mm)",
    ]
    write_header_row(ws, 1, headers)
    set_col_widths(ws, [12, 9, 8, 8, 14, 12, 12, 10, 15, 12, 12, 12, 12])
    ws.freeze_panes = "A2"

    row = 2
    for stn_id in sorted(all_parsed.keys()):
        region = stations.region_of(stn_id) or ""
        for p in all_parsed[stn_id]:
            values = [
                p["date"].isoformat(),
                stn_id,                           # 정수로 출력 (item 2)
                stations.station_name(stn_id),
                region,
                p["sum_gsr"],
                p["sum_ss_hr"],
                p["ss_dur"],
                p["ss_rate"],
                p["hr1_max_icsr"],
                p["avg_ta"],
                p["max_ta"],                      # 최고기온 (item 3)
                p["min_ta"],                      # 최저기온 (item 3)
                p["sum_rn"],
            ]
            write_data_row(ws, row, values, fmt=FMT_NUM1, left_cols={0, 2, 3})
            # 지점번호를 정수 포맷으로 (item 2)
            ws.cell(row, 2).number_format = "0"
            row += 1


# ---------------------------------------------------------------------------
# 읽기전용 비교기간 표시 헬퍼
# ---------------------------------------------------------------------------
def _write_readonly_period(ws, row, col):
    """비교기간을 설정 시트 참조 수식으로 읽기전용 표시.

    설정!B8=분석 시작일, 설정!B9=분석 종료일 을 참조.
    흰색 배경 + 검정 글씨 → 수정 불가 항목임을 시각적으로 표시.
    """
    H(ws, row, col, "비교기간")
    cell = C(ws, row, col + 1, None, left=True)
    cell.value = '=설정!B8&" ~ "&설정!B9'
    cell.font = Font(name=FONT_NAME, size=10, color="1A1A1A")
    cell.fill = PatternFill("solid", fgColor=WHITE)


# ---------------------------------------------------------------------------
# 시트 2: 지역별 집계 — 동적 필터링
# ---------------------------------------------------------------------------
def _build_region_ref_sheet(wb, region_data, ref_year, normal_years):
    """숨긴 _권역별REF 시트 생성.

    복합키 기반으로 모든 권역 × 모든 연도 × 모든 지표 데이터를 기록.
    키 형식: {지표}_{권역}_{연도}_{월}
    예: "평균일사_전국_2026_1" → 전국 2026년 1월 평균 일사량
        "평균일사_전국_2026_연" → 전국 2026년 연간 평균 일사량

    평년: base year별로 사전 계산하여 저장.
    키 형식: {지표}_{권역}_평년{base_year}_{월}
    예: "평균일사_전국_평년2026_1" → 올해=2026 기준 평년 1월 값
        "평균일사_전국_평년2024_1" → 올해=2024 기준 평년 1월 값

    다중 지점 → 일별 평균 후 집계 (기상청 전국 통계 방식).
    설정 시트에서 올해(B4)를 바꾸면 VLOOKUP 키에 포함된 연도가
    바뀌므로 해당 연도 데이터를 동적으로 조회한다.

    열: A=복합키, B=값
    """
    ws = wb.create_sheet("_권역별REF")

    metrics = [
        ("평균일사", "sum_gsr", "avg"),
        ("적산일사", "sum_gsr", "cumsum"),
        ("평균일조", "sum_ss_hr", "avg"),
        ("적산일조", "sum_ss_hr", "cumsum"),
        ("평균일조율", "ss_rate", "avg"),
    ]

    row = 1
    ws.cell(row, 1, "키")
    ws.cell(row, 2, "값")
    row = 2

    for metric_label, metric_key, agg_type in metrics:
        for region in stations.REGION_ORDER:
            reg_rows = region_data.get(region, [])
            if not reg_rows:
                continue

            # ★ 다중 지점 → 일별 평균 변환 (기상청 방식)
            avg_rows = _daily_average_rows(reg_rows)
            monthly = _monthly_metric_by_year(avg_rows, metric_key, agg_type)

            # ── 모든 연도 데이터 (실제 연도 숫자를 키에 사용) ──
            all_years = sorted(monthly.keys())
            for year in all_years:
                vals = monthly[year]

                # 월별
                for m in range(1, 13):
                    key = f"{metric_label}_{region}_{year}_{m}"
                    v = vals.get(m)
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(v, 2) if v is not None else "")
                    row += 1

                # 연간 (평균 or 누계)
                total = 0; cnt = 0
                for m in range(1, 13):
                    v = vals.get(m)
                    if v is not None:
                        total += v; cnt += 1
                key = f"{metric_label}_{region}_{year}_연"
                if cnt:
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(
                        (total if agg_type == "cumsum" else total / cnt), 2))
                else:
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, "")
                row += 1

            # ── 평년 데이터: 가능한 모든 base year별로 저장 ──
            # 올해를 바꾸면 평년 범위도 바뀌므로, 각 base year별 평년을 사전 계산
            # 키: {지표}_{권역}_평년{base_year}_{월}
            for base_year in all_years:
                n_start = base_year - normal_years
                n_end = base_year - 1
                # base_year의 평년 범위에 데이터가 있는 연도가 하나라도 있어야
                if not any(y in monthly for y in range(n_start, n_end + 1)):
                    continue

                normal_vals = _calc_normal_monthly(monthly, n_start, n_end)

                for m in range(1, 13):
                    key = f"{metric_label}_{region}_평년{base_year}_{m}"
                    v = normal_vals.get(m)
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(v, 2) if v is not None else "")
                    row += 1

                # 평년 연간
                total = 0; cnt = 0
                for m in range(1, 13):
                    v = normal_vals.get(m)
                    if v is not None:
                        total += v; cnt += 1
                key = f"{metric_label}_{region}_평년{base_year}_연"
                if cnt:
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(
                        (total if agg_type == "cumsum" else total / cnt), 2))
                else:
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, "")
                row += 1

    # 시트 숨기기
    ws.sheet_state = "hidden"
    return ws


def _sheet_region_summary(wb, all_parsed, ref_year, normal_years,
                          start_date, end_date, target_regions_list):
    """지역별 집계 시트 — 콤보박스+연도 변경으로 동적 필터링."""
    ws = wb.create_sheet("권역별")

    this_year = ref_year
    last_year = ref_year - 1
    normal_start = ref_year - normal_years
    normal_end = ref_year - 1

    # 동기간 MM-DD
    start_md, end_md = _period_md_from_dates(start_date, end_date)

    # 데이터 준비: 권역별로 묶기 + 동기간 필터
    region_data = defaultdict(list)
    for stn_id, parsed in all_parsed.items():
        region = stations.region_of(stn_id) or "기타"
        filtered = _filter_same_period(parsed, start_md, end_md)
        region_data[region].extend(filtered)

    # 전국 = 전체 합산
    all_rows_flat = []
    for rows in region_data.values():
        if rows:
            all_rows_flat.extend(rows)
    region_data["전국"] = all_rows_flat

    # 숨긴 참조 시트 생성
    _build_region_ref_sheet(wb, region_data, ref_year, normal_years)

    # ── 콤보 박스 (노란색) + 비교기간 (읽기전용) ──
    H(ws, 1, 1, "권역 선택")
    C(ws, 1, 2, "전국", fill=YEL)
    # 권역 드롭다운
    region_list = '"' + ",".join(stations.REGION_ORDER) + '"'
    dv_region = DataValidation(type="list", formula1=region_list, allow_blank=False)
    dv_region.prompt = "분석할 권역을 선택하세요"
    dv_region.promptTitle = "권역 선택"
    ws.add_data_validation(dv_region)
    dv_region.add(ws["B1"])

    # 비교기간 — 읽기전용
    _write_readonly_period(ws, 1, 3)

    # 지표별 블록 — VLOOKUP 수식 기반 (설정!B4 참조로 연도 동적)
    metrics = [
        ("< 일 평균 적산일사 (MJ/m²) >", "평균일사"),
        ("< 월 누적 적산일사 (MJ/m²) >", "적산일사"),
        ("< 평균 일조시간 (hr) >", "평균일조"),
        ("< 적산 일조시간 (hr) >", "적산일조"),
        ("< 평균 일조율 (%) >", "평균일조율"),
    ]

    row = 3
    for title_text, metric_label in metrics:
        row = _write_region_vlookup_block(
            ws, row, title_text, metric_label,
        )
        row += 1  # 섹션 간 빈 행

    set_col_widths(ws, [22, 9] + [9] * 12 + [10])


def _write_region_vlookup_block(ws, row_start, title, metric_label):
    """하나의 지표에 대한 VLOOKUP 기반 동적 블록.

    - $B$1: 권역 선택 (콤보)
    - 설정!$B$4: 올해 연도
    - 설정!$B$5: 작년 연도 (=B4-1 수식)

    VLOOKUP 키 예: "평균일사_" & $B$1 & "_" & 설정!$B$4 & "_1"
    → 권역과 연도 모두 동적으로 변경됨.

    Returns: 다음 블록 시작 행
    """
    apply_title(ws, row_start, 1, title, merge_end_col=14, size=11)

    headers = ["비교"] + [f"{m}월" for m in range(1, 13)] + ["연간"]
    write_header_row(ws, row_start + 1, headers)

    row = row_start + 2

    # 비교 행: (A열 레이블 수식, REF키의 연도 부분)
    # 올해: 설정!$B$4, 작년: 설정!$B$5, 평년: "평년" 고정
    compare_rows = [
        ('="올해("&설정!$B$4&")"', '설정!$B$4'),
        ('="작년("&설정!$B$5&")"', '설정!$B$5'),
        ('="평년("&(설정!$B$4-설정!$B$6)&"~"&설정!$B$5&")"', '"평년"&설정!$B$4'),
    ]

    row_refs = {}
    for i, (label_formula, year_ref) in enumerate(compare_rows):
        comp_name = ["올해", "작년", "평년"][i]
        row_refs[comp_name] = row

        # A열: 레이블 (수식으로 동적 표시)
        cell_a = ws.cell(row, 1)
        cell_a.value = label_formula
        cell_a.font = Font(name=FONT_NAME, size=10)
        cell_a.alignment = Alignment("left", "center")
        cell_a.border = BORDER

        # B~M열: 1월~12월 (VLOOKUP 수식, 연도 동적)
        for m in range(1, 13):
            col_idx = m + 1  # B=2, C=3, ..., M=13
            formula = (
                f'=IFERROR(VLOOKUP("{metric_label}_"&$B$1&"_"&{year_ref}&"_{m}",'
                f"'_권역별REF'!$A:$B,2,FALSE),\"\")"
            )
            cell = ws.cell(row, col_idx)
            cell.value = formula
            cell.number_format = FMT_NUM1
            cell.border = BORDER

        # N열: 연간 (VLOOKUP 수식)
        formula = (
            f'=IFERROR(VLOOKUP("{metric_label}_"&$B$1&"_"&{year_ref}&"_연",'
            f"'_권역별REF'!$A:$B,2,FALSE),\"\")"
        )
        cell = ws.cell(row, 14)
        cell.value = formula
        cell.number_format = FMT_NUM1
        cell.border = BORDER

        row += 1

    # 차이(전년) — 올해 - 작년 수식
    this_r = row_refs["올해"]
    last_r = row_refs["작년"]
    normal_r = row_refs["평년"]

    C(ws, row, 1, "차이(전년)", left=True)
    ws.cell(row, 1).border = BORDER
    for col_idx in range(2, 15):  # B ~ N
        cl = get_column_letter(col_idx)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{last_r}=""),"",'
            f'{cl}{this_r}-{cl}{last_r})'
        )
        ws.cell(row, col_idx).value = formula
        ws.cell(row, col_idx).number_format = FMT_NUM1
        ws.cell(row, col_idx).border = BORDER
    row += 1

    # 차이(평년) — 올해 - 평년 수식
    C(ws, row, 1, "차이(평년)", left=True)
    ws.cell(row, 1).border = BORDER
    for col_idx in range(2, 15):
        cl = get_column_letter(col_idx)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{normal_r}=""),"",'
            f'{cl}{this_r}-{cl}{normal_r})'
        )
        ws.cell(row, col_idx).value = formula
        ws.cell(row, col_idx).number_format = FMT_NUM1
        ws.cell(row, col_idx).border = BORDER
    row += 1

    return row


# ---------------------------------------------------------------------------
# 시트 3: 지점별 집계 — VLOOKUP 기반 동적 필터링
# ---------------------------------------------------------------------------

def _build_station_ref_sheet(wb, station_data, ref_year, normal_years):
    """숨긴 _지점별REF 시트 생성.

    지점별 월별 집계를 복합키-값 형식으로 저장한다.
    키 형식: {지표}_{지점번호}_{연도}_{월}
    예: "일평균적산일사_108_2026_1" → 서울(108) 2026년 1월 일평균적산일사

    평년: 각 base_year별로 사전 계산하여 저장.
    키 형식: {지표}_{지점번호}_평년{base_year}_{월}

    열: A=복합키, B=값, D=지점표시명("지점명(번호)"), E=지점번호
    D/E열은 지점 선택 드롭다운 소스로 사용 (255자 제한 우회).

    Parameters
    ----------
    station_data : dict[int, list[dict]]
        이미 동기간 필터된 지점별 일별 데이터.
    """
    ws = wb.create_sheet("_지점별REF")

    metrics = [
        ("일평균적산일사", "sum_gsr", "avg"),
        ("월누적적산일사", "sum_gsr", "cumsum"),
        ("평균일조",       "sum_ss_hr", "avg"),
        ("적산일조",       "sum_ss_hr", "cumsum"),
        ("평균일조율",     "ss_rate", "avg"),
    ]

    # 헤더
    ws.cell(1, 1, "키")
    ws.cell(1, 2, "값")
    ws.cell(1, 4, "지점표시명")
    ws.cell(1, 5, "지점번호")

    # D/E열: 지점 목록 (드롭다운 소스)
    stn_list_row = 2
    first_display = ""
    for stn_id in sorted(station_data.keys()):
        display = f"{stations.station_name(stn_id)}({stn_id})"
        ws.cell(stn_list_row, 4, display)
        ws.cell(stn_list_row, 5, stn_id)
        if stn_list_row == 2:
            first_display = display
        stn_list_row += 1
    n_stations = stn_list_row - 2  # 실제 지점 수

    # A/B열: 키-값 집계 데이터
    row = 2
    for metric_label, metric_key, agg_type in metrics:
        for stn_id in sorted(station_data.keys()):
            parsed = station_data.get(stn_id, [])
            if not parsed:
                continue

            # 단일 지점 → _daily_average_rows 불필요, 직접 집계
            monthly = _monthly_metric_by_year(parsed, metric_key, agg_type)
            all_years = sorted(monthly.keys())

            for year in all_years:
                vals = monthly[year]

                # 월별
                for m in range(1, 13):
                    key = f"{metric_label}_{stn_id}_{year}_{m}"
                    v = vals.get(m)
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(v, 2) if v is not None else "")
                    row += 1

                # 연간 (누적 or 평균)
                total = 0; cnt = 0
                for m in range(1, 13):
                    v = vals.get(m)
                    if v is not None:
                        total += v; cnt += 1
                key = f"{metric_label}_{stn_id}_{year}_연"
                ws.cell(row, 1, key)
                ws.cell(row, 2, round(
                    (total if agg_type == "cumsum" else total / cnt), 2
                ) if cnt else "")
                row += 1

            # 평년: base_year별 사전 계산
            for base_year in all_years:
                n_start = base_year - normal_years
                n_end   = base_year - 1
                if not any(y in monthly for y in range(n_start, n_end + 1)):
                    continue

                normal_vals = _calc_normal_monthly(monthly, n_start, n_end)

                for m in range(1, 13):
                    key = f"{metric_label}_{stn_id}_평년{base_year}_{m}"
                    v = normal_vals.get(m)
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, round(v, 2) if v is not None else "")
                    row += 1

                # 평년 연간
                total = 0; cnt = 0
                for m in range(1, 13):
                    v = normal_vals.get(m)
                    if v is not None:
                        total += v; cnt += 1
                key = f"{metric_label}_{stn_id}_평년{base_year}_연"
                ws.cell(row, 1, key)
                ws.cell(row, 2, round(
                    (total if agg_type == "cumsum" else total / cnt), 2
                ) if cnt else "")
                row += 1

    ws.sheet_state = "hidden"
    return ws, n_stations, first_display


def _write_station_vlookup_block(ws, row_start, title, metric_label):
    """지점별 집계 — 하나의 지표에 대한 VLOOKUP 기반 동적 블록.

    - $B$2: 지점번호 (B1 선택명에서 VLOOKUP으로 추출한 헬퍼 셀)
    - 설정!$B$4: 올해 연도 (동적)

    VLOOKUP 키 예: "일평균적산일사_" & $B$2 & "_" & 설정!$B$4 & "_1"

    Returns
    -------
    int: 다음 블록 시작 행
    """
    apply_title(ws, row_start, 1, title, merge_end_col=14, size=11)

    headers = ["비교"] + [f"{m}월" for m in range(1, 13)] + ["연간"]
    write_header_row(ws, row_start + 1, headers)

    row = row_start + 2

    compare_rows = [
        ('="올해("&설정!$B$4&")"', '설정!$B$4'),
        ('="작년("&설정!$B$5&")"', '설정!$B$5'),
        ('="평년("&(설정!$B$4-설정!$B$6)&"~"&설정!$B$5&")"', '"평년"&설정!$B$4'),
    ]

    row_refs = {}
    for i, (label_formula, year_ref) in enumerate(compare_rows):
        comp_name = ["올해", "작년", "평년"][i]
        row_refs[comp_name] = row

        cell_a = ws.cell(row, 1)
        cell_a.value = label_formula
        cell_a.font = Font(name=FONT_NAME, size=10)
        cell_a.alignment = Alignment("left", "center")
        cell_a.border = BORDER

        # B~M열: 1월~12월
        for m in range(1, 13):
            col_idx = m + 1
            formula = (
                f'=IFERROR(VLOOKUP("{metric_label}_"&$B$2&"_"&{year_ref}&"_{m}",'
                f"'_지점별REF'!$A:$B,2,FALSE),\"\")"
            )
            cell = ws.cell(row, col_idx)
            cell.value = formula
            cell.number_format = FMT_NUM1
            cell.border = BORDER

        # N열: 연간
        formula = (
            f'=IFERROR(VLOOKUP("{metric_label}_"&$B$2&"_"&{year_ref}&"_연",'
            f"'_지점별REF'!$A:$B,2,FALSE),\"\")"
        )
        cell = ws.cell(row, 14)
        cell.value = formula
        cell.number_format = FMT_NUM1
        cell.border = BORDER

        row += 1

    this_r   = row_refs["올해"]
    last_r   = row_refs["작년"]
    normal_r = row_refs["평년"]

    # 차이(전년)
    C(ws, row, 1, "차이(전년)", left=True)
    ws.cell(row, 1).border = BORDER
    for col_idx in range(2, 15):
        cl = get_column_letter(col_idx)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{last_r}=""),"",'
            f'{cl}{this_r}-{cl}{last_r})'
        )
        ws.cell(row, col_idx).value = formula
        ws.cell(row, col_idx).number_format = FMT_NUM1
        ws.cell(row, col_idx).border = BORDER
    row += 1

    # 차이(평년)
    C(ws, row, 1, "차이(평년)", left=True)
    ws.cell(row, 1).border = BORDER
    for col_idx in range(2, 15):
        cl = get_column_letter(col_idx)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{normal_r}=""),"",'
            f'{cl}{this_r}-{cl}{normal_r})'
        )
        ws.cell(row, col_idx).value = formula
        ws.cell(row, col_idx).number_format = FMT_NUM1
        ws.cell(row, col_idx).border = BORDER
    row += 1

    return row


def _sheet_station_summary(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date):
    """지점별 집계 시트 — VLOOKUP 기반 동적 필터링.

    B1 콤보박스에서 지점 선택(전체 없음) → B2 헬퍼 셀로 지점번호 추출
    → 설정!B4 연도 변경 시 데이터 자동 반영.

    지표: 일 평균 적산일사, 월 누적 적산일사, 평균 일조시간,
          적산 일조시간, 평균 일조율
    """
    ws = wb.create_sheet("지점별 집계")

    # 동기간 MM-DD 필터
    start_md, end_md = _period_md_from_dates(start_date, end_date)
    station_data = {}
    for stn_id, parsed in all_parsed.items():
        filtered = _filter_same_period(parsed, start_md, end_md)
        if filtered:
            station_data[stn_id] = filtered

    # 숨긴 REF 시트 생성 (D/E열에 드롭다운 소스 포함)
    _ref_ws, n_stations, first_display = _build_station_ref_sheet(
        wb, station_data, ref_year, normal_years
    )

    # ── Row 1: 지점 선택 콤보 + 비교기간 ──
    H(ws, 1, 1, "지점 선택")
    C(ws, 1, 2, first_display, fill=YEL)

    # 드롭다운: _지점별REF 시트 D열 범위 참조 (255자 제한 우회)
    dv_stn = DataValidation(
        type="list",
        formula1=f"'_지점별REF'!$D$2:$D${n_stations + 1}",
        allow_blank=False,
    )
    dv_stn.prompt = "분석할 지점을 선택하세요"
    dv_stn.promptTitle = "지점 선택"
    ws.add_data_validation(dv_stn)
    dv_stn.add(ws["B1"])

    _write_readonly_period(ws, 1, 3)

    # ── Row 2: 지점번호 헬퍼 셀 (B1 표시명 → 번호 변환) ──
    ws.cell(2, 1).value = "지점번호"
    ws.cell(2, 1).font = Font(name=FONT_NAME, size=9, color="AAAAAA")
    num_cell = ws.cell(2, 2)
    # _지점별REF D/E열로 VLOOKUP: 표시명 → 지점번호
    num_cell.value = "=IFERROR(VLOOKUP(B1,'_지점별REF'!$D:$E,2,FALSE),\"\")"
    num_cell.font = Font(name=FONT_NAME, size=9, color="AAAAAA")
    num_cell.number_format = "0"

    # ── Row 3: 빈 행, Row 4~: 지표 블록 ──
    metrics = [
        ("< 일 평균 적산일사 (MJ/m²) >", "일평균적산일사"),
        ("< 월 누적 적산일사 (MJ/m²) >", "월누적적산일사"),
        ("< 평균 일조시간 (hr) >",        "평균일조"),
        ("< 적산 일조시간 (hr) >",        "적산일조"),
        ("< 평균 일조율 (%) >",           "평균일조율"),
    ]

    row = 4
    for title_text, metric_label in metrics:
        row = _write_station_vlookup_block(ws, row, title_text, metric_label)
        row += 1  # 섹션 간 빈 행

    set_col_widths(ws, [22, 9] + [9] * 12 + [10])


# ---------------------------------------------------------------------------
# 시트 4: 기간별(월별) — VLOOKUP 기반 동적 필터링
# ---------------------------------------------------------------------------

def _write_period_vlookup_block(ws, row_start, title, metric_label, region_cols):
    """기간별(월별) — 하나의 지표에 대한 VLOOKUP 기반 동적 블록.

    _권역별REF 시트를 참조하여 설정!B4 연도 변경 시 자동 반영.
    헤더 행의 권역 이름을 키에 포함시켜 각 열이 해당 권역 데이터를 조회.

    Returns
    -------
    int: 다음 블록 시작 행
    """
    n_regions = len(region_cols)

    # 타이틀
    apply_title(ws, row_start, 1, title,
                merge_end_col=n_regions + 1, size=11)

    # 헤더: 구분 | 전국 | 서울경기 | ... | 제주
    header_row = row_start + 1
    headers = ["구분"] + region_cols
    write_header_row(ws, header_row, headers)

    row = header_row + 1

    # ── 연간 섹션 ──
    compare_rows = [
        ('="올해("&설정!$B$4&")"', '설정!$B$4'),
        ('="작년("&설정!$B$5&")"', '설정!$B$5'),
        ('="평년("&(설정!$B$4-설정!$B$6)&"~"&설정!$B$5&")"', '"평년"&설정!$B$4'),
    ]

    row_refs = {}
    for i, (label_formula, year_ref) in enumerate(compare_rows):
        comp_name = ["올해", "작년", "평년"][i]
        row_refs[comp_name] = row

        cell_a = ws.cell(row, 1)
        cell_a.value = label_formula
        cell_a.font = Font(name=FONT_NAME, size=10)
        cell_a.alignment = Alignment("left", "center")
        cell_a.border = BORDER

        # 각 권역 열: VLOOKUP — 키에 헤더 행 권역명 참조
        for ci in range(2, n_regions + 2):
            cl = get_column_letter(ci)
            formula = (
                f'=IFERROR(VLOOKUP("{metric_label}_"&{cl}${header_row}'
                f'&"_"&{year_ref}&"_연",'
                f"'_권역별REF'!$A:$B,2,FALSE),\"\")"
            )
            cell = ws.cell(row, ci)
            cell.value = formula
            cell.number_format = FMT_NUM1
            cell.border = BORDER

        row += 1

    this_r   = row_refs["올해"]
    last_r   = row_refs["작년"]
    normal_r = row_refs["평년"]

    # 차이(전년)
    C(ws, row, 1, "차이(전년)", left=True)
    ws.cell(row, 1).border = BORDER
    for ci in range(2, n_regions + 2):
        cl = get_column_letter(ci)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{last_r}=""),"",'
            f'{cl}{this_r}-{cl}{last_r})'
        )
        ws.cell(row, ci).value = formula
        ws.cell(row, ci).number_format = FMT_NUM1
        ws.cell(row, ci).border = BORDER
    row += 1

    # 차이(평년)
    C(ws, row, 1, "차이(평년)", left=True)
    ws.cell(row, 1).border = BORDER
    for ci in range(2, n_regions + 2):
        cl = get_column_letter(ci)
        formula = (
            f'=IF(OR({cl}{this_r}="",{cl}{normal_r}=""),"",'
            f'{cl}{this_r}-{cl}{normal_r})'
        )
        ws.cell(row, ci).value = formula
        ws.cell(row, ci).number_format = FMT_NUM1
        ws.cell(row, ci).border = BORDER
    row += 1

    # ── 월별 섹션 ──
    for month in range(1, 13):
        # 월 구분 띠
        H(ws, row, 1, f"── {month}월 ──", fill=LIGHT_BLUE)
        for ci in range(2, n_regions + 2):
            H(ws, row, ci, "", fill=LIGHT_BLUE)
        row += 1

        m_row_refs = {}
        for i, (label_formula, year_ref) in enumerate(compare_rows):
            comp_name = ["올해", "작년", "평년"][i]
            m_row_refs[comp_name] = row

            cell_a = ws.cell(row, 1)
            cell_a.value = label_formula
            cell_a.font = Font(name=FONT_NAME, size=10)
            cell_a.alignment = Alignment("left", "center")
            cell_a.border = BORDER

            for ci in range(2, n_regions + 2):
                cl = get_column_letter(ci)
                formula = (
                    f'=IFERROR(VLOOKUP("{metric_label}_"&{cl}${header_row}'
                    f'&"_"&{year_ref}&"_{month}",'
                    f"'_권역별REF'!$A:$B,2,FALSE),\"\")"
                )
                cell = ws.cell(row, ci)
                cell.value = formula
                cell.number_format = FMT_NUM1
                cell.border = BORDER

            row += 1

        m_this   = m_row_refs["올해"]
        m_last   = m_row_refs["작년"]
        m_normal = m_row_refs["평년"]

        # 차이(전년)
        C(ws, row, 1, "차이(전년)", left=True)
        ws.cell(row, 1).border = BORDER
        for ci in range(2, n_regions + 2):
            cl = get_column_letter(ci)
            formula = (
                f'=IF(OR({cl}{m_this}="",{cl}{m_last}=""),"",'
                f'{cl}{m_this}-{cl}{m_last})'
            )
            ws.cell(row, ci).value = formula
            ws.cell(row, ci).number_format = FMT_NUM1
            ws.cell(row, ci).border = BORDER
        row += 1

        # 차이(평년)
        C(ws, row, 1, "차이(평년)", left=True)
        ws.cell(row, 1).border = BORDER
        for ci in range(2, n_regions + 2):
            cl = get_column_letter(ci)
            formula = (
                f'=IF(OR({cl}{m_this}="",{cl}{m_normal}=""),"",'
                f'{cl}{m_this}-{cl}{m_normal})'
            )
            ws.cell(row, ci).value = formula
            ws.cell(row, ci).number_format = FMT_NUM1
            ws.cell(row, ci).border = BORDER
        row += 1

    return row


def _sheet_period_analysis(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date):
    """기간별(월별) 시트 — VLOOKUP 기반 동적 필터링.

    _권역별REF 시트를 참조하여 설정!B4 연도 변경 시 자동 반영.
    레이아웃:
      열: B~L = 전국/서울경기/.../제주 (11개 권역)
      행: 연간 → 올해/작년/평년/차이(전년)/차이(평년)
          월별 → 1월~12월 각각 올해/작년/평년/차이
    """
    ws = wb.create_sheet("기간별(월별)")

    # 비교기간 — 읽기전용
    _write_readonly_period(ws, 1, 1)

    region_cols = stations.REGION_ORDER  # 전국, 서울경기, ..., 제주

    # 지표별 블록 — _지역별REF VLOOKUP 키와 일치하는 metric_label 사용
    metrics = [
        ("< 일 평균 적산일사 (MJ/m²) >", "평균일사"),
        ("< 월 누적 적산일사 (MJ/m²) >", "적산일사"),
        ("< 평균 일조시간 (hr) >",        "평균일조"),
        ("< 적산 일조시간 (hr) >",        "적산일조"),
        ("< 평균 일조율 (%) >",           "평균일조율"),
    ]

    row = 3
    for title_text, metric_label in metrics:
        row = _write_period_vlookup_block(
            ws, row, title_text, metric_label, region_cols,
        )
        row += 1  # 지표 간 빈 행

    set_col_widths(ws, [18] + [10] * len(region_cols))


# ---------------------------------------------------------------------------
# 시트 5: 년도별 — 지표별 권역 행 + 월별 열 (item 12)
# ---------------------------------------------------------------------------
def _write_yearly_vlookup_block(ws, row_start, title, metric_label,
                                agg_type, data_start_year, ref_year):
    """년도별 — 하나의 지표에 대한 VLOOKUP 기반 동적 블록 (슬라이딩 윈도우).

    _권역별REF 시트를 참조. 설정!B4 연도 변경 시 자동 반영.
    - 슬라이딩 윈도우: 올해-평년기간 ~ 올해 범위만 표시
      조건: OR(year < 설정!$B$4-설정!$B$6, year > 설정!$B$4) 이면 빈칸
    - 미래 대응: ref_year + 5까지 행을 미리 생성 (올해를 미래로 설정 시 대응)
    - C열(평균/누계): 엑셀 AVERAGE/SUM 함수
    - 별도 "올해" 행 없음 — 연도별 행 자체가 조건부로 표시됨

    Parameters
    ----------
    agg_type : "avg" or "cumsum"
    data_start_year : int — 데이터 시작 연도 (예: 2016)
    ref_year : int — 스크립트 실행 시점 기준 연도 (상한 계산용)

    Returns
    -------
    int: 다음 블록 시작 행
    """
    summary_col = "누계" if agg_type == "cumsum" else "평균"
    summary_fn = "SUM" if agg_type == "cumsum" else "AVERAGE"

    n_cols = 15  # A~O

    # 타이틀
    apply_title(ws, row_start, 1, title, merge_end_col=n_cols, size=11)
    row = row_start + 1

    # 헤더
    headers = ["권역", "연도", summary_col] + [f"{m}월" for m in range(1, 13)]
    write_header_row(ws, row, headers)
    row += 1

    # 미래 대응: ref_year + 5까지 행 생성
    max_year = ref_year + 5

    for region in stations.REGION_ORDER:
        # 연도 행 (data_start_year ~ max_year)
        # 슬라이딩 윈도우: 올해-평년기간 ~ 올해 범위 밖이면 빈칸
        # 조건: OR(year < 설정!$B$4-설정!$B$6, year > 설정!$B$4)
        hide_cond = 'OR({y}<설정!$B$4-설정!$B$6,{y}>설정!$B$4)'
        for year in range(data_start_year, max_year + 1):
            cond = hide_cond.format(y=year)

            # 권역 셀: 윈도우 밖이면 빈칸
            cell_a = ws.cell(row, 1)
            cell_a.value = f'=IF({cond},"","{region}")'
            cell_a.font = Font(name=FONT_NAME, size=10)
            cell_a.alignment = Alignment("left", "center")
            cell_a.border = BORDER

            # 연도 셀: 윈도우 밖이면 빈칸
            cell_b = ws.cell(row, 2)
            cell_b.value = f'=IF({cond},"",{year})'
            cell_b.number_format = '0'  # YYYY 정수 표기
            cell_b.font = Font(name=FONT_NAME, size=10)
            cell_b.alignment = Alignment("right", "center")
            cell_b.border = BORDER

            # D~O열: 윈도우 밖이면 빈칸
            for mi, m in enumerate(range(1, 13)):
                ci = mi + 4  # D=4
                formula = (
                    f'=IF({cond},"",IFERROR(VLOOKUP("{metric_label}_{region}_"'
                    f'&$B{row}&"_{m}",'
                    f"'_권역별REF'!$A:$B,2,FALSE),\"\"))"
                )
                cell = ws.cell(row, ci)
                cell.value = formula
                cell.number_format = FMT_NUM1
                cell.border = BORDER

            # C열: 평균 or 누계 — 엑셀 함수 (윈도우 밖이면 "")
            formula_c = (
                f'=IF({cond},"",IF(COUNTA(D{row}:O{row})=0,"",'
                f'{summary_fn}(D{row}:O{row})))'
            )
            cell_c = ws.cell(row, 3)
            cell_c.value = formula_c
            cell_c.number_format = FMT_NUM1
            cell_c.border = BORDER

            row += 1

        # 평년 행 — 설정!$B$4 기준 동적
        C(ws, row, 1, region, left=True)
        ws.cell(row, 1).border = BORDER

        cell_b = ws.cell(row, 2)
        cell_b.value = '="평년("&(설정!$B$4-설정!$B$6)&"~"&설정!$B$5&")"'
        cell_b.font = Font(name=FONT_NAME, size=10)
        cell_b.border = BORDER

        for mi, m in enumerate(range(1, 13)):
            ci = mi + 4
            formula = (
                f'=IFERROR(VLOOKUP("{metric_label}_{region}_"'
                f'&"평년"&설정!$B$4&"_{m}",'
                f"'_권역별REF'!$A:$B,2,FALSE),\"\")"
            )
            cell = ws.cell(row, ci)
            cell.value = formula
            cell.number_format = FMT_NUM1
            cell.border = BORDER

        formula_c = (
            f'=IF(COUNTA(D{row}:O{row})=0,"",'
            f'{summary_fn}(D{row}:O{row}))'
        )
        cell_c = ws.cell(row, 3)
        cell_c.value = formula_c
        cell_c.number_format = FMT_NUM1
        cell_c.border = BORDER
        row += 1

        row += 1  # 권역 간 빈 행

    return row


def _sheet_yearly_analysis(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date):
    """년도별 시트 — VLOOKUP 기반 동적 필터링.

    _권역별REF 시트를 참조하여 설정!B4 연도 변경 시 올해/평년 자동 반영.
    레이아웃:
      열: 권역 | 연도 | 평균(누계) | 1월 | 2월 | ... | 12월
      행: 지표별 그룹 → 권역별 → 연도별 + 평년
    """
    ws = wb.create_sheet("년도별")

    # 데이터 시작 연도 산출
    all_years = set()
    for parsed in all_parsed.values():
        for r in parsed:
            all_years.add(r["year"])
    data_start_year = min(all_years) if all_years else 2016

    # 비교기간 — 읽기전용
    _write_readonly_period(ws, 1, 1)

    # 지표별 블록 — _지역별REF VLOOKUP 키와 일치하는 metric_label 사용
    metrics = [
        ("< 일 평균 적산일사 (MJ/m²) >", "평균일사", "avg"),
        ("< 월 누적 적산일사 (MJ/m²) >", "적산일사", "cumsum"),
        ("< 평균 일조시간 (hr) >",        "평균일조", "avg"),
        ("< 적산 일조시간 (hr) >",        "적산일조", "cumsum"),
        ("< 평균 일조율 (%) >",           "평균일조율", "avg"),
    ]

    row = 3
    for title_text, metric_label, agg_type in metrics:
        row = _write_yearly_vlookup_block(
            ws, row, title_text, metric_label, agg_type,
            data_start_year, ref_year,
        )
        row += 1  # 지표 간 빈 행

    set_col_widths(ws, [10, 22, 10] + [9] * 12)


# ---------------------------------------------------------------------------
# 메인: 전체 프로세스
# ---------------------------------------------------------------------------
def generate_report(
    stn_ids: list[int] | None = None,
    regions: list[str] | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    ref_year: int | None = None,
    normal_years: int = 10,
    output_path: str | None = None,
    use_cache: bool = True,
    verbose: bool = True,
):
    """일사·일조 종합 분석 엑셀을 생성한다."""
    if start_date is None:
        start_date = solar_core.DATA_START
    if end_date is None:
        end_date = solar_core.yesterday()
    if ref_year is None:
        ref_year = date.today().year

    # 대상 지점 결정
    if stn_ids:
        target_stations = stn_ids
    elif regions:
        target_stations = []
        for region in regions:
            target_stations.extend(stations.stations_in_region(region))
    else:
        target_stations = stations.all_station_ids()

    target_regions = regions or stations.REGION_ORDER

    if verbose:
        print(f"=== 일사·일조 종합 분석 ===")
        print(f"기간: {start_date} ~ {end_date}")
        print(f"기준 연도: {ref_year}")
        print(f"평년: {ref_year - normal_years} ~ {ref_year - 1}")
        print(f"대상 지점: {len(target_stations)}개")
        print()

    # 인증키
    keys = solar_core.load_apikeys()
    service_key = keys.get("DATA_GO_KR")
    if not service_key:
        raise ValueError("apikey.txt에 DATA_GO_KR 키가 없습니다.")

    # 데이터 수집
    raw_data = solar_core.fetch_all_stations(
        service_key, target_stations, start_date, end_date,
        use_cache=use_cache, verbose=verbose,
    )

    # 파싱
    all_parsed = {}
    total_rows = 0
    for stn_id, rows in raw_data.items():
        parsed = solar_core.parse_rows(rows)
        all_parsed[stn_id] = parsed
        total_rows += len(parsed)

    if verbose:
        print(f"\n총 {total_rows}건 파싱 완료")
        print("엑셀 생성 중...")

    # 동기간 비교를 위한 기준: 올해 데이터의 실제 종료 MM-DD
    # end_date의 MM-DD를 기준으로 모든 연도의 같은 기간만 비교
    period_start = date(ref_year, start_date.month, start_date.day)
    period_end = end_date

    # 엑셀 생성 — 시트 순서 변경 (item 1: 원데이터 맨 뒤)
    wb = Workbook()

    # 1. 설정
    _sheet_settings(wb, start_date, end_date, ref_year, normal_years,
                    target_regions, target_stations)

    # 2. 권역별
    _sheet_region_summary(wb, all_parsed, ref_year, normal_years,
                          start_date, end_date, target_regions)

    # 3. 지점별 집계
    _sheet_station_summary(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date)

    # 4. 기간별(월별)
    _sheet_period_analysis(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date)

    # 5. 년도별
    _sheet_yearly_analysis(wb, all_parsed, ref_year, normal_years,
                           start_date, end_date)

    # 6. 원데이터 (맨 뒤) — item 1
    _sheet_raw_data(wb, all_parsed)

    # 저장
    if output_path is None:
        output_dir = Path(__file__).parent / "output"
        output_dir.mkdir(exist_ok=True)
        output_path = str(
            output_dir / f"solar_analysis_{ref_year}_{end_date.strftime('%Y%m%d')}.xlsx"
        )

    wb.save(output_path)

    if verbose:
        print(f"\n✓ 저장 완료: {output_path}")

    return output_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="일사·일조 종합 분석 엑셀 생성",
    )
    parser.add_argument(
        "--region", nargs="+", default=None,
        help="분석 대상 권역 (예: 서울경기 충남)",
    )
    parser.add_argument(
        "--station", nargs="+", type=int, default=None,
        help="분석 대상 지점번호 (예: 108 131)",
    )
    parser.add_argument(
        "--year", type=int, default=None,
        help="기준 연도 (기본: 올해)",
    )
    parser.add_argument(
        "--normal-years", type=int, default=10,
        help="평년 산정 기간 (기본: 10년)",
    )
    parser.add_argument(
        "--output", "-o", default=None,
        help="출력 파일 경로",
    )
    parser.add_argument(
        "--no-cache", action="store_true",
        help="CSV 캐시 사용하지 않음 (항상 API 재조회)",
    )
    args = parser.parse_args()

    generate_report(
        stn_ids=args.station,
        regions=args.region,
        ref_year=args.year,
        normal_years=args.normal_years,
        output_path=args.output,
        use_cache=not args.no_cache,
    )


if __name__ == "__main__":
    main()
