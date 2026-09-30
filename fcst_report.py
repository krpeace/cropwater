#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fcst_report.py — 02-Cycle H2 검증 엑셀 (라이브 수식)

시트: 요약 / 일별비교 / 3일누적 / 월별 / 입력진단 / 오차분해 / (격자비교) / Rs계수 / (SKY계수) / 관측 / 설정 / 방법 / 차트자료
  - 하늘상태(SKY) 예보가 있으면 주 방법(★)은 S4(SKY계수 시트의 교차검증 계수), 없으면 S3
  - 값으로 넣는 것: 예보 일 입력(fcst_archive 집계), ASOS 관측 일자료, 오차분해·격자비교(Python 계산)
  - 나머지는 엑셀 수식: Ra·Rs·PM ETo·Kc·ETc·기준선·오차·지표·판정·월별 지표
  - 노란 칸(설정·Rs계수·합격 기준·월별 선행일)을 바꾸면 전체가 다시 계산된다
"""
import datetime as dt, os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

from fao56_core import safe_save
from fcst_archive import PCP_30_50_MM, PCP_GE50_MM, PCP_LT1_MM, SERVICE_RUNS


def run_gaps(runs):
    from cropwater_fcst import run_gaps as _rg
    return _rg(runs)

FONT = "맑은 고딕"; GREEN = "2E6A4C"; BLUE = "2B6E86"; BROWN = "A8681B"; LIGHT = "EFF1EC"; YEL = "FBEED2"
PASS_FILL = "D8EAD3"; FAIL_FILL = "F4C7B8"
thin = Side(style="thin", color="D0D3CC"); BORDER = Border(thin, thin, thin, thin)
F2, F3, PCT, DATE, DTM = "0.00", "0.000", "0%", "yyyy-mm-dd", "yyyy-mm-dd hh:mm"


def H(ws, r, c, v, fill=GREEN, white=True, size=10, wrap=True):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, bold=True, size=size, color=("FFFFFF" if white else "1A1A1A"))
    x.fill = PatternFill("solid", fgColor=fill); x.border = BORDER
    x.alignment = Alignment("center", "center", wrap_text=wrap)
    return x


def C(ws, r, c, v, fmt=None, left=False, fill=None, bold=False, color="1A1A1A", size=10, wrap=False):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, size=size, bold=bold, color=color); x.border = BORDER
    x.alignment = Alignment("left" if left else "center", "center", wrap_text=wrap)
    if fmt: x.number_format = fmt
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x


def T(ws, r, c, v, size=10, bold=False, color="1A1A1A", wrap=False):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, size=size, bold=bold, color=color)
    x.alignment = Alignment("left", "center", wrap_text=wrap)
    return x


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def _disp_width(s):
    """표시 폭 추정: 한글·전각 문자는 2, 나머지는 1"""
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in str(s))


def _wrap_text(s, width):
    """표시 폭 width 이하로 공백에서 나눈 줄 목록. 둘째 줄부터는 두 칸 들여씀"""
    out, cur = [], ""
    for w in s.split(" "):
        cand = f"{cur} {w}" if cur else w
        if cur and _disp_width(cand) > width:
            out.append(cur); cur = "   " + w
        else:
            cur = cand
    return out + [cur] if cur else out


def pm(tmax, tmin, ea, u2, rs, rso, t, es, d, g):
    """FAO-56 식(6) 엑셀 수식 (01-Cycle 계산과정 시트와 같은 형태, G=0)"""
    rnl = (f"0.000000004903*((({tmax}+273.16)^4+({tmin}+273.16)^4)/2)*(0.34-0.14*SQRT({ea}))"
           f"*(1.35*MIN({rs}/{rso},1)-0.35)")
    return f"=(0.408*{d}*(0.77*{rs}-{rnl})+{g}*(900/({t}+273))*{u2}*({es}-{ea}))/({d}+{g}*(1+0.34*{u2}))"


E0 = lambda x: f"0.6108*EXP(17.27*{x}/({x}+237.3))"


# ── 설정 ────────────────────────────────────────────────────────────────
S_ROWS = {}   # 이름 → 셀 주소(설정!$B$n)


def sheet_settings(wb, res):
    ws = wb.create_sheet("설정")
    meta, kp, chk = res["meta"], res["kp"], res["check"]
    st = meta["settings"]
    T(ws, 1, 1, "설정 — 노란 칸을 바꾸면 모든 시트가 다시 계산됩니다", 12, True, GREEN)
    H(ws, 2, 1, "항목"); H(ws, 2, 2, "값"); H(ws, 2, 3, "근거·출처")
    rows = [
        ("[지점]", None, None, None),
        ("site", "지점", f"ASOS {res['stn']} {res.get('stn_name', '')}".strip(),
         f"검증 기준 관측소. 이 지점의 단기예보 격자: {res.get('stn_grid') or '미등록'} (ARCHITECTURE 7장)"),
        ("lat", "위도 (°)", meta["lat"], "Ra·Rso 계산. 01-Cycle 워크북 설정 시트 값 (기상청 지점정보)"),
        ("elev", "고도 (m)", meta["elev"], "Rso, 예보 기압(식7) 계산. 01-Cycle 워크북 설정 시트 값"),
        ("anem", "ASOS 풍속계 높이 (m)", meta.get("anem", 10.0), "관측 u2 환산(식47). 01-Cycle 워크북 설정 시트 값"),
        ("[예보]", None, None, None),
        ("grid", "예보 격자 (nx_ny)", ", ".join(chk["location"]), "과거 단기예보 CSV의 location (지점 격자와 다르면 대표성 오차가 달라짐)"),
        ("fanem", "예보 풍속 높이 (m)", 10.0, "단기예보 WSD는 10 m 풍속으로 가정 → u2 = 0.748 × u10"),
        ("pelev", "예보 기압 P (kPa)", "=101.3*((293-0.0065*$B$6)/293)^5.26", "FAO-56 식(7). 예보에 기압이 없어 고도로 추정"),
        ("rthr", "강수유무 기준 (mm/일)", 1.0, "일강수 ≥ 기준이면 1 (Rs 강수유무 보정의 입력)"),
        ("krs", "FAO-56 kRs (S1)", 0.16, "FAO-56 식(50) 내륙 기본값 (해안 0.19)"),
        ("[작물계수 Kc — 01-Cycle 워크북과 동일]", None, None, None),
        ("crop", "작물", st.get("작물", "사과"), "01-Cycle 워크북 설정 시트"),
        ("scn", "시나리오 번호", kp["scenario"], "FAO-56 Table 12 (참고 표시). 표값은 아래 3칸을 직접 수정"),
        ("kci_t", "표값 Kc_ini", kp["table"][0], "선택 시나리오의 표값"),
        ("kcm_t", "표값 Kc_mid", kp["table"][1], "선택 시나리오의 표값"),
        ("kce_t", "표값 Kc_end", kp["table"][2], "선택 시나리오의 표값"),
        ("bud", "생육 시작일", kp["bud"], "01-Cycle 워크북 설정 시트(발아일)"),
        ("li", "L_ini (일)", kp["L"][0], "FAO-56 Table 11"),
        ("ld", "L_dev (일)", kp["L"][1], "FAO-56 Table 11"),
        ("lm", "L_mid (일)", kp["L"][2], "FAO-56 Table 11"),
        ("ll", "L_late (일)", kp["L"][3], "FAO-56 Table 11"),
        ("h", "생육중기 수고 h (m)", kp["h"], "식(62)·(65)의 (h/3)^0.3 항"),
        ("u2m", "중기 평균 u2 (m/s)", kp["u2_mid"], "01-Cycle 자동집계값(관측 워크북, 생육중기)"),
        ("rhm", "중기 평균 RHmin (%)", kp["rh_mid"], "01-Cycle 자동집계값"),
        ("u2e", "후기 평균 u2 (m/s)", kp["u2_end"], "01-Cycle 값(후기 관측 부족 시 기본값 1.5)"),
        ("rhe", "후기 평균 RHmin (%)", kp["rh_end"], "01-Cycle 값(후기 관측 부족 시 기본값 55)"),
        ("mulch", "멀칭 보정계수", kp["mulch"], "전면 비닐멀칭 0.70~0.90, 없으면 1.00"),
        ("kci", "★ 적용 Kc_ini", None, "Kc_ini × 멀칭 (현지기상 보정 대상 아님)"),
        ("kcm", "★ 적용 Kc_mid", None, "FAO-56 식(62): Kc_mid + [0.04(u2−2) − 0.004(RHmin−45)](h/3)^0.3, u2·RHmin 범위 제한"),
        ("kce", "★ 적용 Kc_end", None, "FAO-56 식(65): 표값 > 0.45일 때만 보정"),
    ]
    r = 3
    for key, label, val, note in rows:
        if label is None:
            H(ws, r, 1, key, fill=LIGHT, white=False); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            r += 1; continue
        S_ROWS[key] = f"설정!$B${r}"
        C(ws, r, 1, label, left=True, fill=LIGHT)
        editable = key in ("lat", "elev", "anem", "fanem", "rthr", "krs", "kci_t", "kcm_t", "kce_t", "bud",
                           "li", "ld", "lm", "ll", "h", "u2m", "rhm", "u2e", "rhe", "mulch")
        fmt = DATE if key == "bud" else (F3 if isinstance(val, float) else None)
        C(ws, r, 2, val, fmt=fmt, fill=(YEL if editable else None))
        C(ws, r, 3, note, left=True, color="555555", size=9)
        r += 1
    assert S_ROWS["elev"] == "설정!$B$6", S_ROWS["elev"]          # pelev 수식이 $B$6을 참조
    s = lambda k: S_ROWS[k].replace("설정!", "")
    ws[s("kci")] = f"={s('kci_t')}*{s('mulch')}"
    ws[s("kcm")] = (f"=({s('kcm_t')}+(0.04*(MEDIAN(1,{s('u2m')},6)-2)-0.004*(MEDIAN(20,{s('rhm')},80)-45))"
                    f"*({s('h')}/3)^0.3)*{s('mulch')}")
    ws[s("kce")] = (f"=IF({s('kce_t')}>0.45,{s('kce_t')}+(0.04*(MEDIAN(1,{s('u2e')},6)-2)-0.004*(MEDIAN(20,{s('rhe')},80)-45))"
                    f"*({s('h')}/3)^0.3,{s('kce_t')})*{s('mulch')}")
    for k in ("kci", "kcm", "kce", "pelev"):
        ws[s(k)].number_format = F3
    widths(ws, {"A": 26, "B": 18, "C": 80})
    ws.freeze_panes = "A3"
    return ws


# ── Rs계수 ──────────────────────────────────────────────────────────────
R_ROWS = {}


def sheet_rscoef(wb, res, obs_rows):
    ws = wb.create_sheet("Rs계수")
    cf = res["coef"]; fitinfo = cf
    yr = res["table"].target.min().year if len(res["table"]) else ""
    T(ws, 1, 1, f"Rs 추정 계수와 H1 재검증 (G2) — 계수는 검증 연도({yr})와 겹치지 않는 해의 관측으로 정함", 12, True, GREEN)
    H(ws, 3, 1, "항목"); H(ws, 3, 2, "값"); H(ws, 3, 3, "설명")
    items = [
        ("a", "a (S3)", cf.get("a"), "Rs/Ra = a + b·√(Tmax−Tmin) + c·강수유무  [THEORY 9장]"),
        ("b", "b (S3)", cf.get("b"), ""),
        ("c", "c (S3)", cf.get("c"), "비 오는 날(≥ 1 mm) 일사 감소분"),
        ("krs2", "지점 kRs (S2)", cf.get("krs"), "식(50)의 kRs를 같은 자료로 맞춘 값"),
        ("src", "출처", cf.get("source", ""), "rs_coef.csv (cropwater_fcst.py calib)"),
        ("fit", "보정 자료", f"{fitinfo.get('fit_start', '')} ~ {fitinfo.get('fit_end', '')} ({fitinfo.get('n', '')}일)",
         fitinfo.get("note", "")),
        ("fitrmse", "보정 자료 Rs RMSE (MJ/m²/일)", fitinfo.get("rmse_rs"), "S3 적합 잔차"),
    ]
    r = 4
    for key, label, val, note in items:
        R_ROWS[key] = f"Rs계수!$B${r}"
        C(ws, r, 1, label, left=True, fill=LIGHT)
        C(ws, r, 2, val, fmt=("0.0000" if isinstance(val, float) else None),
          fill=(YEL if key in ("a", "b", "c", "krs2") else None), left=isinstance(val, str))
        C(ws, r, 3, note, left=True, color="555555", size=9)
        r += 1
    r += 1
    T(ws, r, 1, "H1 재검증 — 관측 입력(기온·습도·풍속·기압)에 Rs만 추정 → 관측 ETo와 비교 (관측 결측일 제외)", 11, True, GREEN); r += 1
    C(ws, r, 1, "검증 시작일", left=True, fill=LIGHT); C(ws, r, 2, res["h1_start"].date(), DATE, fill=YEL); R_ROWS["h1s"] = f"$B${r}"; r += 1
    C(ws, r, 1, "검증 종료일", left=True, fill=LIGHT); C(ws, r, 2, res["h1_end"].date(), DATE, fill=YEL); R_ROWS["h1e"] = f"$B${r}"; r += 1
    C(ws, r, 1, "H1 기준: RMSE 상한 (mm/일)", left=True, fill=LIGHT); C(ws, r, 2, 0.6, F2, fill=YEL); R_ROWS["h1rm"] = f"$B${r}"; r += 1
    C(ws, r, 1, "H1 기준: 합계오차 한계 (±)", left=True, fill=LIGHT); C(ws, r, 2, 0.05, PCT, fill=YEL); R_ROWS["h1se"] = f"$B${r}"; r += 1
    r += 1
    hdr = ["방법", "n", "관측 ETo 평균", "추정 ETo 평균", "MBE", "RMSE", "합계오차", "Rs RMSE", "H1 판정"]
    for j, h in enumerate(hdr, 1):
        H(ws, r, j, h)
    r += 1
    a0, a1 = obs_rows
    rng = lambda col: f"관측!${col}${a0}:${col}${a1}"
    # 관측 결측일(기온·풍속·일사·습도가 하나라도 없음)은 빼고 셈 — Python obs_daily의 관측 ETo 결측 규칙과 같음
    ok = (f"(ISNUMBER({rng('B')})*ISNUMBER({rng('C')})*ISNUMBER({rng('E')})*ISNUMBER({rng('I')})"
          f"*((ISNUMBER({rng('D')})+ISNUMBER({rng('F')})+ISNUMBER({rng('G')}))>0))")
    msk = f"(({rng('A')}>={R_ROWS['h1s']})*({rng('A')}<={R_ROWS['h1e']})*{ok})"
    R_ROWS["h1_first"] = r
    for name, eto, rs in (("S1 식(50) kRs 0.16", "AC", "Z"), ("S2 식(50) 지점 kRs", "AD", "AA"),
                          ("S3 식(50)+강수유무 보정 ★", "AE", "AB")):
        C(ws, r, 1, name, left=True, fill=LIGHT)
        ws.cell(r, 2, f"=SUMPRODUCT({msk}*1)")
        ws.cell(r, 3, f"=SUMPRODUCT({msk}*{rng('X')})/B{r}")
        ws.cell(r, 4, f"=SUMPRODUCT({msk}*{rng(eto)})/B{r}")
        ws.cell(r, 5, f"=D{r}-C{r}")
        ws.cell(r, 6, f"=SQRT(SUMPRODUCT({msk}*({rng(eto)}-{rng('X')})^2)/B{r})")
        ws.cell(r, 7, f"=D{r}/C{r}-1")
        ws.cell(r, 8, f"=SQRT(SUMPRODUCT({msk}*({rng(rs)}-{rng('I')})^2)/B{r})")
        ws.cell(r, 9, f'=IF(AND(F{r}<={R_ROWS["h1rm"]},ABS(G{r})<={R_ROWS["h1se"]}),"충족","미달")')
        for j, fmt in zip(range(2, 10), ["0", F2, F2, F2, F2, "0.0%", F2, None]):
            x = ws.cell(r, j); x.number_format = fmt or "General"; x.font = Font(name=FONT, size=10)
            x.border = BORDER; x.alignment = Alignment("center", "center")
        r += 1
    r += 1
    T(ws, r, 1, "월별 편향 — 관측 입력에 Rs만 추정 (H1 조건). 편향 = 추정 − 관측", 11, True, GREEN); r += 1
    for j, h in enumerate(["월", "일수", "관측 ETo 평균", "S3 ETo 편향 (mm/일)", "S1 ETo 편향 (mm/일)", "관측 Rs 평균",
                           "S3 Rs 편향 (MJ/m²/일)"], 1):
        H(ws, r, j, h)
    r += 1
    for m in range(res["h1_start"].month, res["h1_end"].month + 1):
        mm = f"(({rng('K')}={m})*({rng('A')}>={R_ROWS['h1s']})*({rng('A')}<={R_ROWS['h1e']})*{ok})"
        C(ws, r, 1, m)
        ws.cell(r, 2, f"=SUMPRODUCT({mm}*1)")
        ws.cell(r, 3, f"=IF(B{r}>0,SUMPRODUCT({mm}*{rng('X')})/B{r},0)")
        ws.cell(r, 4, f"=IF(B{r}>0,SUMPRODUCT({mm}*({rng('AE')}-{rng('X')}))/B{r},0)")
        ws.cell(r, 5, f"=IF(B{r}>0,SUMPRODUCT({mm}*({rng('AC')}-{rng('X')}))/B{r},0)")
        ws.cell(r, 6, f"=IF(B{r}>0,SUMPRODUCT({mm}*{rng('I')})/B{r},0)")
        ws.cell(r, 7, f"=IF(B{r}>0,SUMPRODUCT({mm}*({rng('AB')}-{rng('I')}))/B{r},0)")
        for j, fmt in zip(range(2, 8), ["0", F2, F2, F2, F2, F2]):
            x = ws.cell(r, j); x.number_format = fmt; x.font = Font(name=FONT, size=10); x.border = BORDER
            x.alignment = Alignment("center", "center")
        r += 1
    T(ws, r, 1, "월별 시트 ③의 Rs 편향(예보 입력으로 추정)과 비교하면, 예보 입력 탓과 추정식 자체 탓을 나눠 볼 수 있습니다.",
      9, color="555555")
    widths(ws, {"A": 30, "B": 16, "C": 16, "D": 16, "E": 10, "F": 10, "G": 10, "H": 10, "I": 10})
    ws.column_dimensions["C"].width = 16
    return ws


# ── 관측 ────────────────────────────────────────────────────────────────
OBS_COLS = ["일자", "최고기온 Tmax(℃)", "최저기온 Tmin(℃)", "평균습도(%)", "평균풍속 u10(m/s)", "평균증기압(hPa)",
            "평균이슬점(℃)", "현지기압(hPa)", "일사 Rs(MJ/m²)", "일강수(mm)", "월", "연중일 J", "ea(kPa)", "P(kPa)",
            "u2(m/s)", "적위 δ", "일몰각 ωs", "Ra(MJ/m²)", "Rso(MJ/m²)", "T평균(℃)", "es(kPa)", "Δ(kPa/℃)", "γ(kPa/℃)",
            "★ 관측 ETo(mm)", "강수유무", "Rs S1", "Rs S2", "Rs S3", "ETo S1(관측입력)", "ETo S2(관측입력)",
            "ETo S3(관측입력)", "Kc", "관측 ETc(mm)", "직전 7일 평균 ETo"]


def sheet_obs(wb, obs):
    ws = wb.create_sheet("관측")
    for j, h in enumerate(OBS_COLS, 1):
        H(ws, 1, j, h, fill=(GREEN if j <= 10 else BLUE))
    S = S_ROWS; R = R_ROWS
    r0 = 2
    for i, row in enumerate(obs.itertuples()):
        r = r0 + i
        vals = [row.date.date(), row.Tmax, row.Tmin, row.RHmean, row.u10, row.pv, row.td, row.pa, row.Rs, row.rain]
        for j, v in enumerate(vals, 1):
            v = None if (not isinstance(v, dt.date) and pd.isna(v)) else v
            x = ws.cell(r, j, v); x.number_format = DATE if j == 1 else "0.0" if j != 9 else F2
        f = {
            "K": f"=MONTH(A{r})",
            "L": f"=A{r}-DATE(YEAR(A{r}),1,1)+1",
            "M": f'=IF(F{r}<>"",F{r}/10,IF(G{r}<>"",{E0(f"G{r}")},D{r}/100*({E0(f"B{r}")}+{E0(f"C{r}")})/2))',
            "N": f'=IF(H{r}<>"",H{r}/10,{S["pelev"]})',
            "O": f"=E{r}*4.87/LN(67.8*{S['anem']}-5.42)",
            "P": f"=0.409*SIN(2*PI()*L{r}/365-1.39)",
            "Q": f"=ACOS(-TAN(RADIANS({S['lat']}))*TAN(P{r}))",
            "R": f"=(24*60/PI())*0.082*(1+0.033*COS(2*PI()*L{r}/365))*(Q{r}*SIN(RADIANS({S['lat']}))*SIN(P{r})"
                 f"+COS(RADIANS({S['lat']}))*COS(P{r})*SIN(Q{r}))",
            "S": f"=(0.75+0.00002*{S['elev']})*R{r}",
            "T": f"=(B{r}+C{r})/2",
            "U": f"=({E0(f'B{r}')}+{E0(f'C{r}')})/2",
            "V": f"=4098*{E0(f'T{r}')}/(T{r}+237.3)^2",
            "W": f"=0.000665*N{r}",
            "X": pm(f"B{r}", f"C{r}", f"M{r}", f"O{r}", f"I{r}", f"S{r}", f"T{r}", f"U{r}", f"V{r}", f"W{r}"),
            "Y": f"=IF(J{r}>={S['rthr']},1,0)",
            "Z": f"=MIN(MAX({S['krs']}*SQRT(MAX(B{r}-C{r},0))*R{r},0.05*R{r}),S{r})",
            "AA": f"=MIN(MAX({R['krs2']}*SQRT(MAX(B{r}-C{r},0))*R{r},0.05*R{r}),S{r})",
            "AB": f"=MIN(MAX(({R['a']}+{R['b']}*SQRT(MAX(B{r}-C{r},0))+{R['c']}*Y{r})*R{r},0.05*R{r}),S{r})",
            "AC": pm(f"B{r}", f"C{r}", f"M{r}", f"O{r}", f"Z{r}", f"S{r}", f"T{r}", f"U{r}", f"V{r}", f"W{r}"),
            "AD": pm(f"B{r}", f"C{r}", f"M{r}", f"O{r}", f"AA{r}", f"S{r}", f"T{r}", f"U{r}", f"V{r}", f"W{r}"),
            "AE": pm(f"B{r}", f"C{r}", f"M{r}", f"O{r}", f"AB{r}", f"S{r}", f"T{r}", f"U{r}", f"V{r}", f"W{r}"),
        }
        d = f"(A{r}-{S['bud']})"
        li, ld, lm, ll = S["li"], S["ld"], S["lm"], S["ll"]
        ki, km, ke = S["kci"], S["kcm"], S["kce"]
        f["AF"] = (f"=IF({d}<0,0,IF({d}<{li},{ki},IF({d}<{li}+{ld},{ki}+({km}-{ki})*({d}-{li}+1)/{ld},"
                   f"IF({d}<{li}+{ld}+{lm},{km},IF({d}<{li}+{ld}+{lm}+{ll},{km}+({ke}-{km})*({d}-{li}-{ld}-{lm}+1)/{ll},0)))))")
        f["AG"] = f"=AF{r}*X{r}"
        f["AH"] = f"=AVERAGE(X{r - 7}:X{r - 1})" if i >= 7 else ""
        for col, formula in f.items():
            x = ws[f"{col}{r}"]; x.value = formula if formula else None
            x.number_format = "0" if col in ("K", "L", "Y") else F3 if col in ("M", "N", "P", "Q", "U", "V", "AF") else \
                "0.00000" if col == "W" else F2
    r1 = r0 + len(obs) - 1
    for j in range(1, len(OBS_COLS) + 1):
        ws.column_dimensions[CL(j)].width = 11
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "B2"
    return ws, (r0, r1)


# ── 일별비교 ─────────────────────────────────────────────────────────────
DB = {}          # 일별비교 열 키 → 열 문자 (sheet_daily가 채움. 다른 시트는 이 표로 열을 찾는다)
K_ROWS = {}      # SKY계수 시트의 검증용 계수 범위 (sheet_skycoef가 채움)
DB_GROUP = {"A": ("발표·대상일", GREEN), "G": ("예보 일 입력 (fcst_archive 집계값)", GREEN),
            "R": ("Rs 추정·PM 중간값 (수식)", BLUE), "E": ("ETo (mm/일)", BROWN), "X": ("ETo 오차 (−관측)", BROWN),
            "C": ("Kc·ETc (mm/일)", BLUE), "D": ("입력 진단 (관측값·예보 오차)", GREEN)}


def daily_layout(sky):
    """일별비교 열 배치 [(키, 머리글, 묶음)]. sky=True면 하늘상태 입력과 S4(★ 주 방법) 열이 들어간다"""
    L = [("id", "발표번호", "A"), ("rn", "구분", "A"), ("run", "발표시각", "A"), ("rdate", "발표일", "A"),
         ("k", "선행일 k", "A"), ("tgt", "대상일", "A"),
         ("ext", "마지막날(3시간·코드)", "G"), ("tx", "예보 Tmax(℃)", "G"), ("tn", "예보 Tmin(℃)", "G"),
         ("ea", "예보 ea(kPa)", "G"), ("u10", "예보 u10(m/s)", "G"), ("rain", "예보 강수(mm)", "G")]
    if sky:
        L += [("cld", "구름많음 비율(낮)", "G"), ("ovc", "흐림 비율(낮)", "G"), ("pop", "강수확률(낮 평균)", "G")]
    L += [("hrs", "시각 수", "G"), ("fill", "이전 발표로 채운 시각 수", "G"),
          ("flag", "강수유무", "R"), ("ra", "Ra", "R"), ("rso", "Rso", "R")]
    if sky:
        L += [("fold", "S4 교차검증 묶음(대상월)", "R"), ("krow", "S4 계수 행", "R"), ("rs4", "★ Rs S4", "R")]
    L += [("rs3", "Rs S3" if sky else "★ Rs S3", "R"), ("rs1", "Rs S1", "R"), ("u2", "u2(m/s)", "R"),
          ("tm", "T평균", "R"), ("es", "es", "R"), ("dl", "Δ", "R"), ("gm", "γ", "R")]
    if sky:
        L += [("e4", "★ 예보 ETo S4", "E")]
    L += [("e3", "예보 ETo S3" if sky else "★ 예보 ETo S3", "E"), ("e1", "예보 ETo S1", "E"),
          ("eo", "관측 ETo", "E"), ("ep", "지속성", "E"), ("e7", "7일평균", "E")]
    if sky:
        L += [("x4", "오차 S4", "X")]
    L += [("x3", "오차 S3", "X"), ("x1", "오차 S1", "X"), ("xp", "오차 지속성", "X"), ("x7", "오차 7일평균", "X"),
          ("kc", "Kc", "C"), ("etc", "★ 예보 ETc", "C"), ("etco", "관측 ETc", "C"), ("etcp", "지속성 ETc", "C"),
          ("etc7", "7일평균 ETc", "C"), ("xetc", "ETc 오차", "C"), ("xetcp", "지속성 ETc 오차", "C"),
          ("xetc7", "7일평균 ETc 오차", "C"),
          ("otx", "관측 Tmax", "D"), ("otn", "관측 Tmin", "D"), ("oea", "관측 ea", "D"), ("ou", "관측 u10", "D"),
          ("ors", "관측 Rs", "D"), ("oflag", "관측 강수유무", "D"),
          ("dtx", "Tmax 오차", "D"), ("dtn", "Tmin 오차", "D"), ("ddt", "일교차 오차", "D"), ("dea", "ea 오차", "D"),
          ("du", "u10 오차", "D"), ("drs", f"Rs 오차({'S4' if sky else 'S3'})", "D")]
    if sky:
        L += [("drs3", "Rs 오차(S3)", "D")]
    L += [("judge", "강수 판정", "D")]
    return L


def sheet_daily(wb, df, obs_rows, sky=False):
    ws = wb.create_sheet("일별비교")
    layout = daily_layout(sky)
    DB.clear()
    DB.update({key: CL(j) for j, (key, _, _) in enumerate(layout, 1)})
    main = "4" if sky else "3"
    DB.update(eto=DB[f"e{main}"], err=DB[f"x{main}"], rs=DB[f"rs{main}"])
    # 1행: 열 묶음 머리글
    j = 1
    while j <= len(layout):
        g = layout[j - 1][2]
        j1 = j
        while j1 < len(layout) and layout[j1][2] == g:
            j1 += 1
        H(ws, 1, j, DB_GROUP[g][0], fill=DB_GROUP[g][1])
        if j1 > j:
            ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=j1)
        j = j1 + 1
    for j, (_, h, _) in enumerate(layout, 1):
        H(ws, 2, j, h, fill=LIGHT, white=False, size=9)
    S, R = S_ROWS, R_ROWS
    a0, a1 = obs_rows
    ob = lambda col: f"관측!${col}${a0}:${col}${a1}"
    look = lambda col, key: f"INDEX({ob(col)},MATCH({key},{ob('A')},0))"
    r0 = 3
    run_ids = {run: i + 1 for i, run in enumerate(sorted(df.run.unique(), key=lambda t: (t.hour != 2, t)))}
    df = df.sort_values(["run_name", "run", "lead_day"], key=lambda s: s.map({"아침": 0, "저녁": 1}) if s.name == "run_name" else s)
    folds = None
    if sky:
        from cropwater_fcst import month_folds
        folds = df["s4_fold"].astype(str).values if "s4_fold" in df else month_folds(df.target).values
    fmt_val = {"run": DTM, "rdate": DATE, "tgt": DATE, "ea": F3, "u10": F3, "cld": F2, "ovc": F2, "pop": F2}
    fmt_f = {"flag": "0", "oflag": "0", "krow": "0", "es": F3, "dl": F3, "kc": F3, "oea": F3, "dea": F3,
             "gm": "0.00000", "judge": "General"}
    for i, row in enumerate(df.itertuples()):
        r = r0 + i
        c = lambda key: f"{DB[key]}{r}"
        vals = dict(id=run_ids[row.run], rn=row.run_name, run=row.run.to_pydatetime(), rdate=row.run.date(),
                    k=int(row.lead_day), tgt=row.target.date(), ext=("○" if row.ext else ""), tx=row.Tmax, tn=row.Tmin,
                    ea=float(row.ea), u10=float(row.u10), rain=float(row.rain), hrs=int(row.hours), fill=int(row.filled))
        if sky:
            vals.update(cld=float(row.sky_cloudy), ovc=float(row.sky_overcast),
                        pop=(None if pd.isna(getattr(row, "pop", np.nan)) else float(row.pop)), fold=str(folds[i]))
        for key, v in vals.items():
            x = ws[c(key)]; x.value = v
            x.number_format = fmt_val.get(key, "General")
        dT = f"SQRT(MAX({c('tx')}-{c('tn')},0))"
        f = {
            "flag": f"=IF({c('rain')}>={S['rthr']},1,0)",
            "ra": f"={look('R', c('tgt'))}",
            "rso": f"={look('S', c('tgt'))}",
            "rs3": f"=MIN(MAX(({R['a']}+{R['b']}*{dT}+{R['c']}*{c('flag')})*{c('ra')},0.05*{c('ra')}),{c('rso')})",
            "rs1": f"=MIN(MAX({S['krs']}*{dT}*{c('ra')},0.05*{c('ra')}),{c('rso')})",
            "u2": f"={c('u10')}*4.87/LN(67.8*{S['fanem']}-5.42)",
            "tm": f"=({c('tx')}+{c('tn')})/2",
            "es": f"=({E0(c('tx'))}+{E0(c('tn'))})/2",
            "dl": f"=4098*{E0(c('tm'))}/({c('tm')}+237.3)^2",
            "gm": f"=0.000665*{S['pelev']}",
            "e3": pm(c("tx"), c("tn"), c("ea"), c("u2"), c("rs3"), c("rso"), c("tm"), c("es"), c("dl"), c("gm")),
            "e1": pm(c("tx"), c("tn"), c("ea"), c("u2"), c("rs1"), c("rso"), c("tm"), c("es"), c("dl"), c("gm")),
            "eo": f"={look('X', c('tgt'))}",
            "ep": f"={look('X', c('rdate') + '-1')}",
            "e7": f"={look('AH', c('rdate'))}",
            "x3": f"={c('e3')}-{c('eo')}", "x1": f"={c('e1')}-{c('eo')}",
            "xp": f"={c('ep')}-{c('eo')}", "x7": f"={c('e7')}-{c('eo')}",
            "kc": f"={look('AF', c('tgt'))}",
            "etc": f"={c('kc')}*{DB['eto']}{r}", "etco": f"={c('kc')}*{c('eo')}", "etcp": f"={c('kc')}*{c('ep')}",
            "etc7": f"={c('kc')}*{c('e7')}", "xetc": f"={c('etc')}-{c('etco')}", "xetcp": f"={c('etcp')}-{c('etco')}",
            "xetc7": f"={c('etc7')}-{c('etco')}",
            "otx": f"={look('B', c('tgt'))}", "otn": f"={look('C', c('tgt'))}", "oea": f"={look('M', c('tgt'))}",
            "ou": f"={look('E', c('tgt'))}", "ors": f"={look('I', c('tgt'))}", "oflag": f"={look('Y', c('tgt'))}",
            "dtx": f"={c('tx')}-{c('otx')}", "dtn": f"={c('tn')}-{c('otn')}",
            "ddt": f"=({c('tx')}-{c('tn')})-({c('otx')}-{c('otn')})", "dea": f"={c('ea')}-{c('oea')}",
            "du": f"={c('u10')}-{c('ou')}", "drs": f"={DB['rs']}{r}-{c('ors')}",
            "judge": f'=IF({c("flag")}=1,IF({c("oflag")}=1,"적중","오보"),IF({c("oflag")}=1,"놓침","무강수 일치"))',
        }
        if sky:
            K = K_ROWS
            kx = lambda n: f"INDEX({K[n]},{c('krow')})"
            f.update({
                "krow": f'=MATCH({c("fold")}&"|"&{c("k")},{K["key"]},0)',
                "rs4": (f"=MIN(MAX(({kx('a')}+{kx('b')}*{dT}+{kx('c')}*{c('flag')}+{kx('d')}*{c('cld')}"
                        f"+{kx('e')}*{c('ovc')})*{c('ra')},0.05*{c('ra')}),{c('rso')})"),
                "e4": pm(c("tx"), c("tn"), c("ea"), c("u2"), c("rs4"), c("rso"), c("tm"), c("es"), c("dl"), c("gm")),
                "x4": f"={c('e4')}-{c('eo')}",
                "drs3": f"={c('rs3')}-{c('ors')}",
            })
        for key, formula in f.items():
            x = ws[c(key)]; x.value = formula
            x.number_format = fmt_f.get(key, F2)
    r1 = r0 + len(df) - 1
    for j in range(1, len(layout) + 1):
        ws.column_dimensions[CL(j)].width = 10
    for key, w in {"rn": 6, "run": 16, "rdate": 11, "tgt": 11, "id": 7, "k": 7}.items():
        ws.column_dimensions[DB[key]].width = w
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = f"{DB['ext']}3"
    ws.auto_filter.ref = f"A2:{CL(len(layout))}{r1}"
    return ws, (r0, r1), run_ids


# ── SKY계수 (S4 교차검증 계수, Python 적합값) ───────────────────────────────
def sheet_skycoef(wb, res):
    ws = wb.create_sheet("SKY계수")
    T(ws, 1, 1, "S4 하늘상태 반영 Rs 계수 — 예보 입력(일교차·강수유무·구름 비율) → 관측 Rs, 선행일별 (Python 최소제곱 적합값)",
      12, True, GREEN)
    T(ws, 2, 1, "Rs/Ra = a + b·√(Tmax−Tmin) + c·강수유무 + d·구름많음 비율 + e·흐림 비율, [0.05Ra, Rso]로 제한. "
                "구름 비율 = 낮 시간 일사 비중(태양고도 사인)으로 가중한 하늘상태(SKY) 비율 (THEORY 9장)", 9, color="555555")
    fixed = res.get("s4_fixed") is not None
    if fixed:
        f = res["s4_fixed"]
        T(ws, 3, 1, f"① 검증용 계수 — 다른 해 운영 계수를 그대로 적용(계수 고정): {s4_source(res)}. 검증 연도 자료는 계수에 쓰지 않았습니다"
                    "(독립 연도 검증). 일별비교 시트가 '고정|선행일' 키로 이 표를 찾습니다. 노란 칸을 바꾸면 S4 예보가 다시 계산됩니다.",
          9, color=BROWN)
    else:
        T(ws, 3, 1, "① 검증용 계수 — 월 단위 교차검증: 대상월의 행에는 그 달을 뺀 나머지 달로 맞춘 계수를 씁니다. "
                    "일별비교 시트가 '묶음|선행일' 키로 이 표를 찾습니다. 노란 칸을 바꾸면 S4 예보가 다시 계산됩니다.", 9, color=BROWN)
    hdr = ["키", "계수 묶음" if fixed else "검증 달(묶음)", "선행일 k", "적합 행 수(원래 해)" if fixed else "학습 행 수",
           "a", "b", "c (강수유무)", "d (구름많음)", "e (흐림)"]
    for j, h in enumerate(hdr, 1):
        H(ws, 5, j, h, size=9)
    tab = res["s4_table"].sort_values(["fold", "lead_day"]).reset_index(drop=True)
    r0 = 6
    for i, row in tab.iterrows():
        r = r0 + i
        C(ws, r, 1, f"{row.fold}|{int(row.lead_day)}"); C(ws, r, 2, row.fold); C(ws, r, 3, int(row.lead_day))
        C(ws, r, 4, None if _num(pd.to_numeric(row.n, errors="coerce")) is None else int(float(row.n)), "0")
        for j, n in enumerate(("a", "b", "c", "d", "e"), 5):
            C(ws, r, j, float(row[n]), "0.0000", fill=YEL)
    r1 = r0 + len(tab) - 1
    rng = lambda col: f"SKY계수!${col}${r0}:${col}${r1}"
    K_ROWS.update(key=rng("A"), a=rng("E"), b=rng("F"), c=rng("G"), d=rng("H"), e=rng("I"))
    r = r1 + 2
    T(ws, r, 1, ("② 비교 — 이 자료(검증 연도)로 맞춘 계수(참고). ①의 다른 해 계수와 비교해 계수가 해마다 얼마나 달라지는지 봅니다. "
                 "같은 자료로 맞춘 값이라 검증에는 쓰지 않습니다.") if fixed else
                ("② 운영 계수 — 자료 전체로 선행일별 적합(rs_sky_coef.csv, calib-sky). 같은 자료로 맞춘 값이라 검증에는 쓰지 않습니다. "
                 "다른 해 자료로 독립 검증할 때와 운영(G5)에 씁니다."), 9, color=BROWN); r += 1
    hdr = ["선행일 k", "a", "b", "c (강수유무)", "d (구름많음)", "e (흐림)", "적합 행 수", "적합 기간", "Rs RMSE (적합 자료)"]
    for j, h in enumerate(hdr, 1):
        H(ws, r, j, h, size=9)
    r += 1
    for row in res["s4_all"].itertuples():
        C(ws, r, 1, int(row.lead_day))
        for j, n in enumerate(("a", "b", "c", "d", "e"), 2):
            C(ws, r, j, float(getattr(row, n)), "0.0000")
        C(ws, r, 7, int(row.n), "0"); C(ws, r, 8, f"{row.fit_start} ~ {row.fit_end}"); C(ws, r, 9, float(row.rmse_rs), F2)
        r += 1
    r += 1
    for n in ["읽는 법: d·e가 음수 = 구름이 많을수록 Rs가 줄어듦. 가까운 날(D+0)에서 먼 날(D+3)로 갈수록 흐림 계수(e)의 크기가 줄고 "
              "강수유무 계수(c)의 크기가 커지는 경향 → 먼 날의 하늘상태 예보는 덜 믿을 만하다는 뜻입니다.",
              "관측 운량이 없어 계수는 '예보 입력 → 관측 Rs'로 맞췄습니다. 그래서 예보 입력의 계통오차(좁은 일교차, 잦은 비 예보)까지 계수가 흡수합니다.",
              (f"S3(Rs계수 시트)도 검증 연도와 다른 해의 관측으로 정한 계수입니다({res['coef'].get('source', '')})." if fixed else
               f"S3(Rs계수 시트)는 관측 입력으로 정한 계수({res['coef'].get('source', '')})라 이 표와 성격이 다릅니다. "
               "S4 계수의 다른 해 검증은 verify --s4-coef로 합니다(VALIDATION #13).")]:
        T(ws, r, 1, n, 9, color="555555"); r += 1
    widths(ws, {"A": 14, "B": 13, "C": 9, "D": 11, **{CL(j): 12 for j in range(5, 10)}})
    ws.column_dimensions["H"].width = 24
    ws.freeze_panes = "A6"
    return ws


def _ci(col):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(col)


# ── 3일누적 ─────────────────────────────────────────────────────────────
def sheet_cum3(wb, df, db_rows, run_ids, keep=None):
    """keep: 3일 누적에 넣을 (구분, 발표) 집합 — 첫 3개 대상일이 모두 채점 가능한 발표만(cropwater_fcst.cum3와 같음)"""
    ws = wb.create_sheet("3일누적")
    T(ws, 1, 1, "발표별 첫 3개 대상일 누적 (아침 D+0~D+2, 저녁 D+1~D+3) — FAO-56은 추정 Rs 기반 ETo를 여러 날 합계로 쓰도록 권고", 11, True, GREEN)
    hdr = ["발표번호", "구분", "발표시각", "대상기간", "첫 선행일", "예보 ETo 3일합", "관측 ETo 3일합", "지속성 3일합",
           "ETo 오차", "지속성 오차", "예보 ETc 3일합", "관측 ETc 3일합", "지속성 ETc 3일합", "ETc 오차", "지속성 ETc 오차"]
    for j, h in enumerate(hdr, 1):
        H(ws, 2, j, h, size=9)
    d0, d1 = db_rows
    db = lambda col: f"일별비교!${col}${d0}:${col}${d1}"
    runs = df.drop_duplicates("run")[["run", "run_name"]].copy()
    if keep is not None:
        runs = runs[[(n, t) in keep for n, t in zip(runs.run_name, runs.run)]]
    runs["id"] = runs.run.map(run_ids)
    runs = runs.sort_values("id")
    first = {n: min(ls) for n, (_, ls) in SERVICE_RUNS.items()}
    r = 3
    for row in runs.itertuples():
        k0 = first.get(row.run_name, int(df[df.run == row.run].lead_day.min()))
        C(ws, r, 1, row.id); C(ws, r, 2, row.run_name); C(ws, r, 3, row.run.to_pydatetime(), DTM)
        C(ws, r, 4, f"D+{k0}~D+{k0 + 2}"); C(ws, r, 5, k0)
        s = lambda col: (f'=SUMIFS({db(col)},{db(DB["id"])},$A{r},{db(DB["k"])},">="&$E{r},'
                         f'{db(DB["k"])},"<="&($E{r}+2))')
        for j, col in zip(range(6, 9), (DB["eto"], DB["eo"], DB["ep"])):
            C(ws, r, j, s(col), F2)
        C(ws, r, 9, f"=F{r}-G{r}", F2); C(ws, r, 10, f"=H{r}-G{r}", F2)
        for j, col in zip(range(11, 14), (DB["etc"], DB["etco"], DB["etcp"])):
            C(ws, r, j, s(col), F2)
        C(ws, r, 14, f"=K{r}-L{r}", F2); C(ws, r, 15, f"=M{r}-L{r}", F2)
        r += 1
    widths(ws, {CL(j): 11 for j in range(1, 16)}); ws.column_dimensions["C"].width = 16
    ws.freeze_panes = "D3"
    return ws, (3, r - 1)


# ── 입력진단 ─────────────────────────────────────────────────────────────
def sheet_diag(wb, db_rows, groups, main="S3"):
    ws = wb.create_sheet("입력진단")
    T(ws, 1, 1, "입력 진단 — 예보 입력 − 관측 (편향 = 평균, RMSE). 기온 정의 차이: 예보 TMX 09~18시·TMN 03~09시, ASOS 0~24시", 11, True, GREEN)
    hdr = ["구분", "선행일", "n", "Tmax 편향", "Tmax RMSE", "Tmin 편향", "Tmin RMSE", "일교차 편향", "일교차 RMSE",
           "ea 편향", "ea RMSE", "u10 편향", "u10 RMSE", "Rs 편향", "Rs RMSE", "적중", "놓침", "오보", "무강수 일치",
           "POD", "FAR", "CSI"]
    for j, h in enumerate(hdr, 1):
        H(ws, 2, j, h, size=9)
    d0, d1 = db_rows
    db = lambda col: f"일별비교!${col}${d0}:${col}${d1}"
    r = 3
    for rn, k in groups:
        C(ws, r, 1, rn); C(ws, r, 2, k)
        crit = f'{db(DB["rn"])},$A{r},{db(DB["k"])},$B{r}'
        m = f"({db(DB['rn'])}=$A{r})*({db(DB['k'])}=$B{r})"
        C(ws, r, 3, f"=COUNTIFS({crit})", "0")
        for j, key in zip(range(4, 16, 2), ("dtx", "dtn", "ddt", "dea", "du", "drs")):
            col = DB[key]
            fmt = F3 if key == "dea" else F2
            C(ws, r, j, f"=AVERAGEIFS({db(col)},{crit})", fmt)
            C(ws, r, j + 1, f"=SQRT(SUMPRODUCT({m}*{db(col)}^2)/$C{r})", fmt)
        for j, lab in zip(range(16, 20), ("적중", "놓침", "오보", "무강수 일치")):
            C(ws, r, j, f'=COUNTIFS({crit},{db(DB["judge"])},"{lab}")', "0")
        C(ws, r, 20, f"=IF(P{r}+Q{r}>0,P{r}/(P{r}+Q{r}),0)", F2)
        C(ws, r, 21, f"=IF(P{r}+R{r}>0,R{r}/(P{r}+R{r}),0)", F2)
        C(ws, r, 22, f"=IF(P{r}+Q{r}+R{r}>0,P{r}/(P{r}+Q{r}+R{r}),0)", F2)
        r += 1
    r += 1
    notes = ["단위: 기온 ℃, ea kPa, u10 m/s, Rs MJ/m²/일. 편향 = 예보 − 관측 (+면 예보가 큼).",
             (f"Rs 예보는 관측이 아니라 추정한 값({main})입니다. "
              + ("S4 = 식(50)형 + 강수유무 + 하늘상태 구름 비율(SKY계수 시트)." if main == "S4" else "S3 = 식(50) + 강수유무 보정.")),
             "강수: 예보 PCP 일합계와 ASOS 일강수를 1 mm 기준으로 판정. POD = 적중/(적중+놓침), FAR = 오보/(적중+오보), CSI = 적중/(적중+놓침+오보).",
             "아침 D+k와 전날 저녁 D+(k+1)은 같은 대상일입니다. 최고·최저기온은 17시 발표 값이 다음 날 02시 발표까지 유지되는 경우가 많아 두 발표의 기온 진단이 거의 같습니다."]
    for n in notes:
        T(ws, r, 1, n, 9, color="555555"); r += 1
    widths(ws, {CL(j): 9 for j in range(1, 23)})
    return ws


# ── 오차분해 (Python 계산값) ─────────────────────────────────────────────
def sheet_attr(wb, res):
    ws = wb.create_sheet("오차분해")
    T(ws, 1, 1, "오차 분해·보정 탐색·판정 불확실성 — Python(cropwater_fcst.py) 계산값 (수식 아님)", 12, True, GREEN)
    T(ws, 2, 1, "입력을 관측값으로 바꿔 PM을 다시 풀거나 표본을 반복 추출해야 해서 값으로 넣었습니다. "
                "설정·계수를 바꿔도 이 시트는 다시 계산되지 않습니다.", 9, color=BROWN)
    att, names = res["attr"], res["attr_names"]
    r = 4
    T(ws, r, 1, f"① 예보 입력을 하나씩 관측값으로 바꿨을 때의 ETo RMSE (mm/일, {res.get('main', 'S3')})", 11, True, GREEN); r += 1
    hdr = ["구분", "선행일"] + names
    for j, h in enumerate(hdr, 1):
        H(ws, r, j, h, size=9)
    r += 1
    for row in att.itertuples(index=False):
        for j, v in enumerate(row, 1):
            C(ws, r, j, v, F3 if j > 2 else None)
        r += 1
    notes = ["읽는 법: '예보 입력 그대로'보다 많이 줄어드는 입력일수록 그 입력의 예보오차가 ETo 오차에 크게 기여합니다.",
             "'관측 입력 전부 (Rs만 추정)'은 예보가 완벽해도 남는 구조오차(H1 조건, 기압은 고도 추정)입니다.",
             "입력끼리 상관이 있어 줄어든 양의 합이 전체 오차와 같지는 않습니다."]
    for n in notes:
        T(ws, r, 1, n, 9, color="555555"); r += 1
    r += 1
    T(ws, r, 1, "② 지점 보정 탐색 — 월 단위 교차검증 (보정값은 검증 달을 뺀 나머지 달로 추정)", 11, True, GREEN); r += 1
    T(ws, r, 1, f"교차검증 묶음: {', '.join(res['bc_folds'])} (마지막 달 대상일이 10일 미만이면 앞 달에 포함). H2 판정에는 쓰지 않는 탐색 결과입니다.", 9, color="555555"); r += 1
    T(ws, r, 1, "방법: 기온 보정 = Tmax·Tmin에서 평균 편향을 뺌 / 기온+풍속 = 풍속 편향도 뺌 / ETo 비율 = 예보 ETo × (관측 합 ÷ 예보 합)"
                + (" / Rs 계수 재보정 = Rs 식의 a·b·c를 예보 일교차·예보 강수유무와 관측 Rs로 다시 맞춤" if res.get("main", "S3") == "S3" else
                   (". 주 방법 S4는 다른 해 계수를 고정 적용했으므로 'Rs 계수 재보정'은 생략하고(이 해 자료로 다시 맞추면 독립 검증이 아님), "
                    "ETo 비율은 학습 행의 예보로 그대로 구함" if res.get("s4_fixed") is not None else
                    ". 주 방법 S4는 계수가 이미 예보 입력으로 맞춘 교차검증 값이라 'Rs 계수 재보정'은 생략하고, ETo 비율은 검증 달까지 뺀 계수로 "
                    "다시 계산한 학습 행에서 구함(중첩 교차검증)")), 9, color="555555"); r += 1
    bc = res["bc"]
    methods = list(dict.fromkeys(bc.method))
    for metric, label, fmt in (("RMSE", "RMSE (mm/일)", F3), ("MBE", "MBE (mm/일)", F3), ("skill_pers", "지속성 대비 개선율", PCT),
                               ("skill_7d", "7일평균 대비 개선율", PCT)):
        H(ws, r, 1, label, fill=LIGHT, white=False); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        for j, m in enumerate(methods, 3):
            H(ws, r, j, m, size=9)
        r += 1
        pv = bc.pivot_table(index=["run_name", "lead_day"], columns="method", values=metric, sort=False, dropna=False)
        for (rn, k), vals in pv.iterrows():
            C(ws, r, 1, rn); C(ws, r, 2, int(k))
            for j, m in enumerate(methods, 3):
                C(ws, r, j, _num(vals.get(m)), fmt)          # 값을 못 구한 보정(교차검증 학습 자료 부족)은 빈칸
            r += 1
    H(ws, r, 1, "3일 누적 RMSE (mm/3일)", fill=LIGHT, white=False); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for j, m in enumerate(methods, 3):
        H(ws, r, j, m, size=9)
    r += 1
    for rn in dict.fromkeys(bc.run_name):
        C(ws, r, 1, rn); C(ws, r, 2, "첫 3일")
        for j, m in enumerate(methods, 3):
            C(ws, r, j, _num(bc[(bc.method == m) & (bc.run_name == rn)].RMSE_3d.iloc[0]), F2)
        r += 1
    r += 1
    T(ws, r, 1, "③ 전체 기간 입력 편향 (예보 − 관측) — 보정값의 크기", 11, True, GREEN); r += 1
    for j, h in enumerate(["구분", "선행일", "Tmax (℃)", "Tmin (℃)", "u10 (m/s)"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    for row in res["bc_bias"].itertuples(index=False):
        for j, v in enumerate(row, 1):
            C(ws, r, j, v, F2 if j > 2 else None)
        r += 1
    bt = res.get("boot")
    if bt:
        r += 1
        T(ws, r, 1, f"④ 판정의 표본 불확실성 — {bt['block']}일 이동 블록 부트스트랩 {bt['n_boot']:,}회 "
                    f"(발표일 {bt['n_days']}일, 난수 시드 {bt['seed']})", 11, True, GREEN); r += 1
        hdr = ["D+1 RMSE 5%", "D+1 RMSE 중앙값", "D+1 RMSE 95%", "최소 개선율 5%", "최소 개선율 중앙값", "최소 개선율 95%",
               "두 기준 충족 비율"]
        variants = [("보정 없음 (판정 기준)", bt)] + [(f"{n} (탐색)", b) for n, b in (res.get("boot_bc") or {}).items()]
        for name, b in variants:
            H(ws, r, 1, name, fill=LIGHT, white=False); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            for j, h in enumerate(hdr, 3):
                H(ws, r, j, h, size=9)
            r += 1
            for rn, v in b["runs"].items():
                vals = [rn, "", *v["rmse_d1"], *v["min_skill"], v["pass_frac"]]
                for j, x in enumerate(vals, 1):
                    C(ws, r, j, x, PCT if j >= 6 else (F3 if j > 2 else None), bold=(j == 9))
                r += 1
            C(ws, r, 1, "두 발표 모두"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            for j in range(3, 9):
                C(ws, r, j, "")
            C(ws, r, 9, b["pass_all"], PCT, bold=True); r += 1
        for n in [f"발표일을 {bt['block']}일 묶음으로 복원추출해 지표를 다시 계산했습니다. 같은 날의 아침·저녁 발표와 D+1~D+3을 함께 뽑아 "
                  "서로의 상관을 보존합니다. 90% 구간 = 5~95 백분위.",
                  "두 기준 충족 비율 = 반복 중 'D+1~D+3 최소 개선율 ≥ 기준'과 'D+1 RMSE ≤ 기준'을 함께 만족한 비율(기준은 2026-09-29 확정값). "
                  "90%보다 낮으면 점추정이 통과해도 기준과의 여유가 표본 변동보다 작다는 뜻입니다.",
                  "보정 행은 ②의 월 단위 교차검증 보정값을 적용한 예보로 같은 계산을 한 탐색 결과입니다(H2 판정에는 쓰지 않음)."]:
            T(ws, r, 1, n, 9, color="555555"); r += 1
    widths(ws, {"A": 12, "B": 8, **{CL(j): 15 for j in range(3, 11)}})
    return ws


# ── 격자비교 (Python 계산값, --compare 때만) ─────────────────────────────
def sheet_grid_compare(wb, res):
    gc = res["grid_cmp"]
    A, B = ", ".join(res["check"]["location"]), ", ".join(res["compare"]["check"]["location"])
    ws = wb.create_sheet("격자비교")
    T(ws, 1, 1, f"격자 비교 — A {A} (이 검증) vs B {B} — Python(cropwater_fcst.py) 계산값", 12, True, GREEN)
    T(ws, 2, 1, f"같은 관측·Rs 계수·Kc로 계산했습니다. A − B 직접 차이는 같은 발표·선행일끼리 비교합니다(n = {gc['n']}).", 9, color="555555")
    r = 4
    T(ws, r, 1, "① 선행시간별 예보 ETo 성능 (S3, mm/일)", 11, True, GREEN); r += 1
    hdr = ["구분", "선행일", f"A MBE", f"A RMSE", f"A 개선율(지속성)", f"B MBE", f"B RMSE", f"B 개선율(지속성)", "RMSE 차이 A−B"]
    for j, h in enumerate(hdr, 1):
        H(ws, r, j, h, size=9)
    r += 1
    m = gc["metrics"]
    for _, a in m[m.grid == "A"].iterrows():
        b = m[(m.grid == "B") & (m.run_name == a.run_name) & (m.lead_day == a.lead_day)].iloc[0]
        vals = [a.run_name, int(a.lead_day), a.MBE, a.RMSE, a.skill_pers, b.MBE, b.RMSE, b.skill_pers, a.RMSE - b.RMSE]
        for j, v in enumerate(vals, 1):
            C(ws, r, j, v, PCT if j in (5, 8) else (F3 if j > 2 else None))
        r += 1
    r += 1
    T(ws, r, 1, "② 입력 편향 (예보 − ASOS 관측)", 11, True, GREEN); r += 1
    hdr = ["구분", "선행일", "A 최고기온(℃)", "A 최저기온(℃)", "A 풍속(m/s)", "B 최고기온(℃)", "B 최저기온(℃)", "B 풍속(m/s)"]
    for j, h in enumerate(hdr, 1):
        H(ws, r, j, h, size=9)
    r += 1
    dg = gc["diag"]
    for _, a in dg[dg.grid == "A"].iterrows():
        b = dg[(dg.grid == "B") & (dg.run_name == a.run_name) & (dg.lead_day == a.lead_day)].iloc[0]
        vals = [a.run_name, int(a.lead_day), a.Tmax_MBE, a.Tmin_MBE, a.u10_MBE, b.Tmax_MBE, b.Tmin_MBE, b.u10_MBE]
        for j, v in enumerate(vals, 1):
            C(ws, r, j, v, F2 if j > 2 else None)
        r += 1
    r += 1
    T(ws, r, 1, "③ 같은 발표·선행일의 직접 차이 (A − B)", 11, True, GREEN); r += 1
    for j, h in enumerate(["항목", "평균", "표준편차", "값이 다른 비율"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    for _, d in gc["diff"].iterrows():
        C(ws, r, 1, d["item"], left=True); C(ws, r, 2, d["mean"], "+0.00;-0.00;0.00"); C(ws, r, 3, d["sd"], F2)
        C(ws, r, 4, d["frac_diff"], PCT)
        r += 1
    C(ws, r, 1, "강수유무(≥ 1 mm) 일치율", left=True); C(ws, r, 2, gc["rain_agree"], PCT); r += 2
    T(ws, r, 1, "④ 월별 직접 차이 (A − B, 대상일 기준)", 11, True, GREEN); r += 1
    for j, h in enumerate(["월", "n", "최고기온(℃)", "최저기온(℃)", "풍속(m/s)", "예보 ETo(mm/일)"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    for _, d in gc["by_month"].iterrows():
        vals = [int(d.month), int(d.n), d.dTmax, d.dTmin, d.du10, d.dETo]
        for j, v in enumerate(vals, 1):
            C(ws, r, j, v, "+0.00;-0.00;0.00" if j > 2 else "0")
        r += 1
    r += 1
    T(ws, r, 1, "읽는 법: 두 격자는 입력 관측·계수가 같으므로 성능 차이는 모두 예보 격자 차이(대표성)에서 옵니다. "
                "농장 예보는 농장 좌표의 격자로 받아야 합니다(THEORY 9장 ◇ 격자 주의점).", 9, color="555555")
    widths(ws, {"A": 22, "B": 10, **{CL(j): 14 for j in range(3, 10)}})
    return ws


# ── 월별 (수식) ──────────────────────────────────────────────────────────
SUM_CELLS = {}   # 요약 시트 합격 기준 셀 (sheet_summary가 채움, 월별 시트가 참조)


def sheet_month(wb, df, db_rows, main="S3"):
    """대상일 월별 성능(①)·달마다 기준 적용(②, 참고)·입력 편향(③). 모두 일별비교 시트를 참조하는 수식"""
    ws = wb.create_sheet("월별")
    d0, d1 = db_rows
    db = lambda col: f"일별비교!${col}${d0}:${col}${d1}"
    runs = [n for n in SERVICE_RUNS if n in set(df.run_name)]
    months = sorted(int(m) for m in df.target.dt.month.unique())
    T(ws, 1, 1, "월별 성능 — 생육기 안에서 어느 달이 약한가 (대상일 기준 월, 일별비교 시트를 참조하는 수식)", 12, True, GREEN)
    T(ws, 2, 1, "달마다 기준을 적용한 결과는 원인 진단용 참고입니다. H2 판정은 요약 시트(전체 기간)로 합니다. "
                "노란 칸의 선행일을 바꾸면 ①·③이 다시 계산됩니다.", 9, color=BROWN)
    for c0, c1, label in ((1, 3, "선행일 k (①·③)"), (6, 7, "개선율 기준 ≥"), (9, 10, "D+1 RMSE 기준 ≤")):
        for c in range(c0, c1 + 1):
            C(ws, 4, c, label if c == c0 else None, left=True, fill=LIGHT, bold=(c0 == 1))
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c1)
    C(ws, 4, 4, 1, "0", fill=YEL)
    C(ws, 4, 8, f"={SUM_CELLS['skill']}", PCT)
    C(ws, 4, 11, f"={SUM_CELLS['rmse']}", F2)
    T(ws, 4, 12, "← 기준은 요약 시트를 따라감", 9, color="555555")
    K, SK, RM = "$D$4", "$H$4", "$K$4"
    mk = lambda r, k: f"({db(DB['rn'])}=$A{r})*({db(DB['k'])}={k})*(MONTH({db(DB['tgt'])})=$B{r})"
    bad = PatternFill("solid", fgColor=FAIL_FILL)

    # ① 월별 성능
    r = 6
    T(ws, r, 1, f"① 월별 예보 ETo 성능 (선행일 k, {main}, mm/일)", 11, True, GREEN); r += 1
    for j, h in enumerate(["구분", "월", "n", "관측 평균", "예보 평균", "MBE", "RMSE", "상대 RMSE", "지속성 RMSE",
                           "개선율(지속성)", "7일평균 RMSE", "개선율(7일평균)"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    rows1 = {}
    for rn in runs:
        for m in months:
            mask = mk(r, K)
            z = lambda expr, rr=r: f'=IF($C{rr}=0,"-",{expr})'
            C(ws, r, 1, rn); C(ws, r, 2, m)
            C(ws, r, 3, f"=SUMPRODUCT({mask})", "0")
            C(ws, r, 4, z(f"SUMPRODUCT({mask}*{db(DB['eo'])})/$C{r}"), F2)
            C(ws, r, 5, z(f"SUMPRODUCT({mask}*{db(DB['eto'])})/$C{r}"), F2)
            C(ws, r, 6, z(f"SUMPRODUCT({mask}*{db(DB['err'])})/$C{r}"), F2)
            C(ws, r, 7, z(f"SQRT(SUMPRODUCT({mask}*{db(DB['err'])}^2)/$C{r})"), F2)
            C(ws, r, 8, z(f"G{r}/D{r}"), PCT)
            C(ws, r, 9, z(f"SQRT(SUMPRODUCT({mask}*{db(DB['xp'])}^2)/$C{r})"), F2)
            C(ws, r, 10, z(f"1-G{r}/I{r}"), PCT)
            C(ws, r, 11, z(f"SQRT(SUMPRODUCT({mask}*{db(DB['x7'])}^2)/$C{r})"), F2)
            C(ws, r, 12, z(f"1-G{r}/K{r}"), PCT)
            rows1[(rn, m)] = r
            r += 1
    a0, a1 = min(rows1.values()), max(rows1.values())
    ws.conditional_formatting.add(f"G{a0}:G{a1}", FormulaRule(formula=[f"AND({K}=1,ISNUMBER(G{a0}),G{a0}>{RM})"], fill=bad))
    ws.conditional_formatting.add(f"J{a0}:J{a1}", FormulaRule(formula=[f"AND({K}>=1,{K}<=3,ISNUMBER(J{a0}),J{a0}<{SK})"],
                                                              fill=bad))
    T(ws, r, 1, "붉은 칸: 선행일 1의 RMSE가 기준보다 크거나, 선행일 1~3의 개선율이 기준보다 작은 달. "
                "상대 RMSE = RMSE ÷ 관측 평균 (ETo가 큰 여름은 절대 오차도 커지므로 함께 봄).", 9, color="555555")
    chart_at = 7
    r += 2

    # ② 달마다 H2 기준 적용 (참고)
    T(ws, r, 1, "② 달마다 H2 기준을 적용하면 (참고 — H2 판정은 전체 기간)", 11, True, GREEN); r += 1
    for j, h in enumerate(["구분", "월", "D+1 n", "개선율 D+1", "개선율 D+2", "개선율 D+3", "최소 개선율", "D+1 RMSE",
                           "개선율 기준", "D+1 기준", "결과"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    b0 = r
    for rn in runs:
        for m in months:
            C(ws, r, 1, rn); C(ws, r, 2, m)
            C(ws, r, 3, f"=SUMPRODUCT({mk(r, 1)})", "0")
            for j, k in zip((4, 5, 6), (1, 2, 3)):
                a = mk(r, k)
                C(ws, r, j, f'=IF(SUMPRODUCT({a})=0,"-",1-SQRT(SUMPRODUCT({a}*{db(DB["err"])}^2)/SUMPRODUCT({a}*{db(DB["xp"])}^2)))', PCT)
            C(ws, r, 7, f'=IF(COUNT(D{r}:F{r})=0,"-",MIN(D{r}:F{r}))', PCT)
            C(ws, r, 8, f'=IF(C{r}=0,"-",SQRT(SUMPRODUCT({mk(r, 1)}*{db(DB["err"])}^2)/C{r}))', F2)
            C(ws, r, 9, f'=IF(ISNUMBER(G{r}),IF(G{r}>={SK},"충족","미달"),"-")')
            C(ws, r, 10, f'=IF(ISNUMBER(H{r}),IF(H{r}<={RM},"충족","미달"),"-")')
            C(ws, r, 11, f'=IF(AND(I{r}="충족",J{r}="충족"),"충족","미달")', bold=True)
            r += 1
    rng_v = f"I{b0}:K{r - 1}"
    ws.conditional_formatting.add(rng_v, FormulaRule(formula=[f'I{b0}="충족"'], fill=PatternFill("solid", fgColor=PASS_FILL)))
    ws.conditional_formatting.add(rng_v, FormulaRule(formula=[f'I{b0}="미달"'], fill=bad))
    r += 1

    # ③ 월별 입력 편향
    T(ws, r, 1, "③ 월별 입력 편향 (예보 − 관측, 선행일 k)", 11, True, GREEN); r += 1
    for j, h in enumerate(["구분", "월", "n", "Tmax (℃)", "Tmin (℃)", "일교차 (℃)", "ea (kPa)", "u10 (m/s)",
                           f"Rs {main} (MJ/m²/일)", "예보 강수일 비율", "관측 강수일 비율", "오보 (일)", "놓침 (일)"], 1):
        H(ws, r, j, h, size=9)
    r += 1
    for rn in runs:
        for m in months:
            mask = mk(r, K)
            C(ws, r, 1, rn); C(ws, r, 2, m); C(ws, r, 3, f"=SUMPRODUCT({mask})", "0")
            for j, key in zip(range(4, 10), ("dtx", "dtn", "ddt", "dea", "du", "drs")):
                C(ws, r, j, f'=IF($C{r}=0,"-",SUMPRODUCT({mask}*{db(DB[key])})/$C{r})', F3 if key == "dea" else F2)
            C(ws, r, 10, f'=IF($C{r}=0,"-",SUMPRODUCT({mask}*{db(DB["flag"])})/$C{r})', PCT)
            C(ws, r, 11, f'=IF($C{r}=0,"-",SUMPRODUCT({mask}*{db(DB["oflag"])})/$C{r})', PCT)
            C(ws, r, 12, f'=SUMPRODUCT({mask}*({db(DB["judge"])}="오보"))', "0")
            C(ws, r, 13, f'=SUMPRODUCT({mask}*({db(DB["judge"])}="놓침"))', "0")
            r += 1
    r += 1
    for n in [f"편향 = 예보 − 관측 (+면 예보가 큼). Rs는 예보 입력으로 추정한 값({main}) − 관측 Rs.",
              "Rs계수 시트의 월별 표(관측 입력으로 추정한 Rs의 편향)와 비교하면, 예보 입력 탓(일교차·강수유무 예보오차)과 "
              "추정식 자체 탓을 나눠 볼 수 있습니다.",
              "예보 강수일 비율이 관측보다 크면(오보가 많으면) 강수유무 보정(계수 c < 0) 때문에 Rs와 ETo가 과소추정됩니다.",
              "예보 일교차가 관측보다 좁으면(일교차 편향 −) Rs가 작게 추정됩니다. 예보 최고·최저기온은 ASOS 일 최고·최저(0~24시)와 "
              "정의가 달라(낮최고 09~18시, 아침최저 03~09시) 일교차가 좁게 나오는 경향이 있습니다."]:
        T(ws, r, 1, n, 9, color="555555"); r += 1

    # 차트: 월별 RMSE (선행일 k)
    bar = BarChart(); bar.type = "col"; bar.grouping = "clustered"
    bar.title = "월별 RMSE (선행일 k, ETo mm/일)"; bar.y_axis.title = "RMSE (mm/일)"
    for rn, color in zip(runs, (BLUE, BROWN)):
        rr = [rows1[(rn, m)] for m in months]
        bar.add_data(Reference(ws, min_col=7, min_row=rr[0], max_row=rr[-1]), titles_from_data=False)
        s = bar.series[-1]; s.tx = SeriesLabel(v=rn)
        s.graphicalProperties.solidFill = color; s.graphicalProperties.line.solidFill = color
    rr = [rows1[(runs[0], m)] for m in months]
    bar.set_categories(Reference(ws, min_col=2, min_row=rr[0], max_row=rr[-1]))
    bar.height = 7.5; bar.width = 15; bar.legend.position = "b"; bar.y_axis.majorGridlines = None
    ws.add_chart(bar, f"N{chart_at}")
    widths(ws, {"A": 8, "B": 6, "C": 7, **{CL(j): 11 for j in range(4, 14)}})
    ws.freeze_panes = "A5"
    return ws


# ── 차트자료 ─────────────────────────────────────────────────────────────
def sheet_chartdata(wb, df, db_rows, obs_rows):
    ws = wb.create_sheet("차트자료")
    H(ws, 1, 1, "대상일"); H(ws, 1, 2, "관측 ETo"); H(ws, 1, 3, "아침 D+1 예보"); H(ws, 1, 4, "저녁 D+1 예보")
    d0, d1 = db_rows; a0, a1 = obs_rows
    db = lambda col: f"일별비교!${col}${d0}:${col}${d1}"
    m = set(df[(df.run_name == "아침") & (df.lead_day == 1)].target) & set(df[(df.run_name == "저녁") & (df.lead_day == 1)].target)
    dates = sorted(m)
    for i, d in enumerate(dates):
        r = 2 + i
        C(ws, r, 1, d.date(), DATE)
        C(ws, r, 2, f"=INDEX(관측!$X${a0}:$X${a1},MATCH(A{r},관측!$A${a0}:$A${a1},0))", F2)
        C(ws, r, 3, f'=SUMIFS({db(DB["eto"])},{db(DB["rn"])},"아침",{db(DB["k"])},1,{db(DB["tgt"])},A{r})', F2)
        C(ws, r, 4, f'=SUMIFS({db(DB["eto"])},{db(DB["rn"])},"저녁",{db(DB["k"])},1,{db(DB["tgt"])},A{r})', F2)
    widths(ws, {"A": 12, "B": 11, "C": 13, "D": 13})
    return ws, (2, 1 + len(dates))


# ── 요약 ────────────────────────────────────────────────────────────────
def _num(v):
    """숫자 → float, 결측(NaN·None)은 None(빈칸)"""
    return None if v is None or pd.isna(v) else float(v)


SOURCE_NAMES = {"portal": "기상자료개방포털 과거 단기예보", "openapi": "기상청 단기예보 조회서비스(OpenAPI) 응답",
                "element": "단기예보 요소별 CSV(발표·예보 일시 KST, OpenAPI 값 표기)"}


def s4_source(res):
    """고정 S4 계수의 출처 설명: 파일 이름·적합 기간"""
    f = res["s4_fixed"]
    per = f"{f.fit_start.min()} ~ {f.fit_end.max()}" if str(f.fit_start.min()) not in ("", "nan") else "기간 미기재"
    return f"{os.path.basename(str(res.get('s4_coef_path') or 'rs_sky_coef.csv'))}, 적합 {per}"


def source_text(chk):
    """예보 자료 출처 이름 (check_archive의 formats — 없으면 포털)"""
    return " + ".join(SOURCE_NAMES[f] for f in (chk.get("formats") or ["portal"]) if f in SOURCE_NAMES)


def sheet_summary(wb, res, db_rows, c3_rows, groups, chart_rows):
    ws = wb["요약"]
    chk, df = res["check"], res["table"]
    d0, d1 = db_rows
    db = lambda col: f"일별비교!${col}${d0}:${col}${d1}"
    fixed = res.get("s4_fixed") is not None
    T(ws, 1, 1, f"단기예보 기반 ETo·ETc 예측 검증 (H2{', 다른 해 독립 검증' if fixed else ''}) — ASOS {res['stn']} · "
                f"{res['meta']['settings'].get('작물', '사과')}", 14, True, GREEN)
    T(ws, 2, 1, f"02-Cycle 3단계{' · S4 계수 고정(VALIDATION #13)' if fixed else ''} · 작성 {dt.date.today():%Y-%m-%d} · "
                f"cropwater_fcst.py verify · 근거 docs/THEORY.md 9장, 판정 기록 docs/VALIDATION.md", 9, color="555555")
    iss = res["arch"].issues
    n_m, n_e = (int(df[df.run_name == n].run.nunique()) for n in ("아침", "저녁"))
    gaps = run_gaps(res.get("skipped_runs", []))
    dropped = res.get("dropped")
    main = res.get("main", "S3")
    opt = [e for e in ("SKY", "POP") if e in chk.get("coverage", {})]
    opt_txt = ""
    if opt:
        cv = chk["coverage"]
        opt_txt = " · 선택 요소 " + ", ".join(f"{e} {cv[e]['first'][5:13]}시~{cv[e]['last'][5:13]}시" for e in opt)
    rs_txt = (f"S4 = 식(50)형 + 강수유무 + 낮 시간 구름많음·흐림 비율(하늘상태 예보), "
              + (f"다른 해 운영 계수를 선행일별로 그대로 적용({s4_source(res)})" if fixed else "선행일별 계수를 월 단위 교차검증으로 정함")
              + f" — SKY계수 시트. 비교: S3(계수 {res['coef'].get('source', '')}, Rs계수 시트)" if main == "S4" else
              f"식(50) + 강수유무 보정(S3), 계수 {res['coef'].get('source', '')} — Rs계수 시트")
    info = [
        ("예보 자료", f"{source_text(chk)} · 격자 {', '.join(chk['location'])}"
                    f"{'(파일에 격자 정보가 없어 실행 시 지정)' if res.get('grid_given') else ''} · TMX·TMN·TMP·REH·WSD·PCP · "
                    f"{iss[0]:%Y-%m-%d %H}시 ~ {iss[-1]:%Y-%m-%d %H}시 발표 중 6요소가 모두 있는 발표 {chk['issues']}회"
                    f" (한 요소라도 빠진 발표 {len(chk['missing_issues'])}회){opt_txt}"),
        ("서비스 발표", f"아침 02시 {n_m}회 → 오늘~D+3 · 저녁 17시 {n_e}회 → 내일~D+4 (다른 발표는 이전 시각 채움에만 사용)"),
        ("대상일", f"{df.target.min():%Y-%m-%d} ~ {df.target.max():%Y-%m-%d}"),
        ("관측 기준", f"ASOS {res['stn']} 일자료 → FAO-56 PM ETo (관측 Rs, 01-Cycle 규칙과 동일, 차이 < 1e-12 mm)"),
        (f"Rs 추정 (★ {main})", rs_txt),
        ("작물계수", f"{res['meta']['settings'].get('작물', '')} 시나리오 {res['kp']['scenario']}, 생육 시작 {res['kp']['bud']} — 01-Cycle 워크북과 같은 Kc"),
    ]
    if gaps or (dropped is not None and len(dropped)):
        gtxt = ", ".join(f"{a:%m/%d %H}시~{b:%m/%d %H}시 {n}회" if n > 1 else f"{a:%m/%d %H}시" for a, b, n, _ in gaps)
        reasons = dropped.drop_reason.value_counts().to_dict() if dropped is not None and len(dropped) else {}
        info.append(("제외", (f"요소가 빠진 서비스 발표 {sum(g[2] for g in gaps)}회({gtxt})" if gaps else "제외 발표 없음")
                     + ("; " + ", ".join(f"{k} {v}행" for k, v in reasons.items()) if reasons else "")))
    if res.get("compare"):
        info.append(("비교 격자", f"{', '.join(res['compare']['check']['location'])} — 같은 관측·계수·Kc로 계산해 격자비교 시트에 정리"))
    r = 4
    for k, v in info:
        C(ws, r, 1, k, left=True, fill=LIGHT, bold=True); x = C(ws, r, 2, v, left=True, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
        lines = max(1, -(-_disp_width(v) // 150))       # 병합 칸은 자동 맞춤이 안 되므로 줄 수만큼 높이를 직접 정함
        if lines > 1:
            ws.row_dimensions[r].height = 14.5 * lines
        r += 1
    r += 1
    # 합격 기준 (수정 가능)
    T(ws, r, 1, "H2 합격 기준 (노란 칸 수정 가능, 2026-09-29 확정값)", 11, True, GREEN); r += 1
    C(ws, r, 1, "지속성 대비 개선율 ≥", left=True, fill=LIGHT); C(ws, r, 2, 0.30, PCT, fill=YEL)
    C(ws, r, 3, "D+1~D+3 모두 (두 발표 각각)", left=True, color="555555"); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    skill_cell = f"$B${r}"; r += 1
    C(ws, r, 1, "D+1 RMSE ≤ (mm/일)", left=True, fill=LIGHT); C(ws, r, 2, 1.0, F2, fill=YEL)
    C(ws, r, 3, "두 발표 각각", left=True, color="555555"); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    rmse_cell = f"$B${r}"; r += 2
    SUM_CELLS.update(skill=f"요약!{skill_cell}", rmse=f"요약!{rmse_cell}")
    verdict_hdr = r
    for j, h in enumerate(["구분", "D+1~D+3 최소 개선율", "D+1 RMSE", "개선율 기준", "D+1 기준", "판정"], 1):
        H(ws, r, j, h)
    r += 1
    verdict_rows = {}
    for rn in dict.fromkeys(g[0] for g in groups):
        verdict_rows[rn] = r; r += 1
    overall_row = r; r += 2
    # 선행시간별 ETo 지표
    T(ws, r, 1, f"선행시간별 예보 ETo 성능 (★ {main}, mm/일) — 모든 값은 일별비교 시트를 참조하는 수식", 11, True, GREEN); r += 1
    met_hdr = ["", "구분", "선행일", "n", "관측 평균", "예보 평균", "MBE", "RMSE", "R²", "지속성 RMSE", "개선율(지속성)",
               "7일평균 RMSE", "개선율(7일평균)", "H2 판정 대상"]
    eto_first = r + 1
    rows_eto = _metric_table(ws, r, met_hdr, groups, db, DB["eto"], DB["eo"], DB["err"], DB["xp"], DB["x7"])
    r = rows_eto[-1] + 2
    if main == "S4":
        T(ws, r, 1, "Rs 추정 방법 비교 — 같은 대상일, 예보 ETo RMSE (mm/일)와 지속성 대비 개선율. "
                    f"S4 = 하늘상태 반영(★), S3 = 식(50)+강수유무({res['coef'].get('fit_start', '')[:4]} 계수), S1 = 식(50) kRs 0.16",
          11, True, GREEN); r += 1
        r = _method_table(ws, r, groups, db) + 1
    T(ws, r, 1, "선행시간별 예보 ETc (Kc × ETo, mm/일)", 11, True, GREEN); r += 1
    rows_etc = _metric_table(ws, r, met_hdr, groups, db, DB["etc"], DB["etco"], DB["xetc"], DB["xetcp"], DB["xetc7"])
    r = rows_etc[-1] + 2
    lead_row = {g: rr for g, rr in zip(groups, rows_eto)}
    for rn, vr in verdict_rows.items():
        ks = [lead_row[(rn, k)] for k in (1, 2, 3) if (rn, k) in lead_row]
        C(ws, vr, 1, rn)
        C(ws, vr, 2, "=MIN(" + ",".join(f"K{x}" for x in ks) + ")", PCT)
        C(ws, vr, 3, f"=H{lead_row[(rn, 1)]}", F2)
        C(ws, vr, 4, f'=IF(B{vr}>={skill_cell},"충족","미달")')
        C(ws, vr, 5, f'=IF(C{vr}<={rmse_cell},"충족","미달")')
        C(ws, vr, 6, f'=IF(AND(D{vr}="충족",E{vr}="충족"),"통과","미달")', bold=True)
    C(ws, overall_row, 1, f"H2 종합 (대상일 {df.target.min():%m/%d}~{df.target.max():%m/%d})", bold=True, fill=LIGHT)
    ws.merge_cells(start_row=overall_row, start_column=1, end_row=overall_row, end_column=5)
    conds = ",".join(f'F{vr}="통과"' for vr in verdict_rows.values())
    C(ws, overall_row, 6, f'=IF(AND({conds}),"통과","미달")', bold=True)
    rng_v = f"D{verdict_hdr + 1}:F{overall_row}"
    ws.conditional_formatting.add(rng_v, FormulaRule(formula=[f'OR(D{verdict_hdr + 1}="통과",D{verdict_hdr + 1}="충족")'],
                                                     fill=PatternFill("solid", fgColor=PASS_FILL)))
    ws.conditional_formatting.add(rng_v, FormulaRule(formula=[f'D{verdict_hdr + 1}="미달"'],
                                                     fill=PatternFill("solid", fgColor=FAIL_FILL)))
    # 3일 누적
    T(ws, r, 1, "3일 누적 (발표별 첫 3개 대상일) — 3일누적 시트 참조", 11, True, GREEN); r += 1
    c0, c1 = c3_rows
    cr = lambda col: f"'3일누적'!${col}${c0}:${col}${c1}"
    hdr3 = ["", "구분", "대상기간", "n", "관측 3일합 평균", "ETo MBE", "ETo RMSE", "상대 RMSE", "지속성 RMSE", "개선율(지속성)",
            "ETc MBE", "ETc RMSE", "ETc 지속성 RMSE", "ETc 개선율"]
    for j, h in enumerate(hdr3, 1):
        if h: H(ws, r, j, h, size=9)
    r += 1
    c3_row = {}
    for rn in dict.fromkeys(g[0] for g in groups):
        c3_row[rn] = r
        m = f"({cr('B')}=$B{r})"
        C(ws, r, 2, rn)
        C(ws, r, 3, "D+0~D+2" if rn == "아침" else "D+1~D+3")
        C(ws, r, 4, f"=COUNTIF({cr('B')},$B{r})", "0")
        C(ws, r, 5, f"=SUMPRODUCT({m}*{cr('G')})/D{r}", F2)
        C(ws, r, 6, f"=SUMPRODUCT({m}*{cr('I')})/D{r}", F2)
        C(ws, r, 7, f"=SQRT(SUMPRODUCT({m}*{cr('I')}^2)/D{r})", F2)
        C(ws, r, 8, f"=G{r}/E{r}", PCT)
        C(ws, r, 9, f"=SQRT(SUMPRODUCT({m}*{cr('J')}^2)/D{r})", F2)
        C(ws, r, 10, f"=1-G{r}/I{r}", PCT)
        C(ws, r, 11, f"=SUMPRODUCT({m}*{cr('N')})/D{r}", F2)
        C(ws, r, 12, f"=SQRT(SUMPRODUCT({m}*{cr('N')}^2)/D{r})", F2)
        C(ws, r, 13, f"=SQRT(SUMPRODUCT({m}*{cr('O')}^2)/D{r})", F2)
        C(ws, r, 14, f"=1-L{r}/M{r}", PCT)
        r += 1
    r += 1
    # 주요 발견 (수식으로 연결된 문장 + 해석)
    T(ws, r, 1, "주요 발견", 11, True, GREEN); r += 1
    va, ve = verdict_rows["아침"], verdict_rows["저녁"]
    di = lambda rn, k, col: f"입력진단!{col}{3 + groups.index((rn, k))}"
    lines = [
        f'="① 판정: D+1~D+3 지속성 대비 개선율 최소 아침 "&TEXT(B{va},"0%")&" · 저녁 "&TEXT(B{ve},"0%")'
        f'&", D+1 RMSE 아침 "&TEXT(C{va},"0.00")&" · 저녁 "&TEXT(C{ve},"0.00")&" mm/일 → "&F{overall_row}',
        f'="② 편향: 예보 ETo가 관측보다 평균 "&TEXT(AVERAGE(G{rows_eto[0]}:G{rows_eto[-1]}),"+0.00;-0.00")'
        f'&" mm/일 (−면 과소). 3일 누적 편향은 아침 "&TEXT(F{c3_row["아침"]},"+0.0;-0.0")&" mm, 저녁 "&TEXT(F{c3_row["저녁"]},"+0.0;-0.0")&" mm"',
        f'="③ 입력 편향(아침 D+1, 예보−관측): 최고기온 "&TEXT({di("아침", 1, "D")},"+0.0;-0.0")&"℃, 최저기온 "'
        f'&TEXT({di("아침", 1, "F")},"+0.0;-0.0")&"℃, 풍속 "&TEXT({di("아침", 1, "L")},"+0.00;-0.00")&" m/s, ea "'
        f'&TEXT({di("아침", 1, "J")},"+0.00;-0.00")&" kPa (입력진단 시트)"',
    ]
    for txt in res.get("findings", []):
        lines.append(txt)
    for ln in lines:
        # 긴 해석 문장은 여러 행으로 나눔(병합 칸의 줄바꿈·행 높이는 엑셀과 LibreOffice가 다르게 그려 차트와 겹침)
        for part in ([ln] if ln.startswith("=") else _wrap_text(ln, 160)):
            T(ws, r, 1, part, 10); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14); r += 1
    # 차트 (주요 발견 아래). LibreOffice는 글꼴에 맞춰 행 높이를 늘린 뒤에도 차트를 처음 위치에 두므로
    # 위쪽 행 수의 5%만큼 빈 행을 더 둬 글과 겹치지 않게 한다(엑셀은 차트가 행을 따라 움직여 영향 없음)
    r += 1 + max(2, -(-r // 20))
    bar = BarChart(); bar.type = "col"; bar.grouping = "clustered"
    bar.title = "선행시간별 RMSE (ETo, mm/일)"; bar.y_axis.title = "RMSE (mm/일)"
    cats = Reference(ws, min_col=1, min_row=rows_eto[0], max_row=rows_eto[-1])
    for col, name, color in ((8, f"예보({main})", GREEN), (10, "지속성", BROWN), (12, "7일평균", "9AA39A")):
        bar.add_data(Reference(ws, min_col=col, min_row=rows_eto[0], max_row=rows_eto[-1]), titles_from_data=False)
        srs = bar.series[-1]; srs.tx = SeriesLabel(v=name)
        srs.graphicalProperties.solidFill = color; srs.graphicalProperties.line.solidFill = color
    bar.set_categories(cats); bar.height = 8; bar.width = 17; bar.legend.position = "b"
    bar.y_axis.majorGridlines = None
    ws.add_chart(bar, f"A{r}")
    ln = LineChart(); ln.title = "관측 ETo와 D+1 예보 ETo (mm/일)"; ln.y_axis.title = "ETo (mm/일)"
    cw = wb["차트자료"]; q0, q1 = chart_rows
    for col, name, color, w in ((2, "관측", "1A1A1A", 22000), (3, "아침 D+1 예보", BLUE, 15000), (4, "저녁 D+1 예보", BROWN, 15000)):
        ln.add_data(Reference(cw, min_col=col, min_row=q0, max_row=q1), titles_from_data=False)
        srs = ln.series[-1]; srs.tx = SeriesLabel(v=name); srs.smooth = False
        srs.graphicalProperties.line.solidFill = color; srs.graphicalProperties.line.width = w
    ln.set_categories(Reference(cw, min_col=1, min_row=q0, max_row=q1))
    ln.x_axis.number_format = "mm-dd"; ln.height = 8; ln.width = 24; ln.legend.position = "b"
    ws.add_chart(ln, f"G{r}")
    widths(ws, {"A": 20, "B": 16, "C": 13, "D": 12, "E": 12, "F": 12, "G": 10, "H": 10, "I": 10, "J": 12, "K": 13,
                "L": 12, "M": 14, "N": 12})
    return ws


def _method_table(ws, r, groups, db):
    """방법 비교 표(수식): 발표 × 선행일별 RMSE(S4·S3·S1), 개선율, S4 − S3. 반환: 다음 빈 행"""
    hdr = ["", "구분", "선행일", "n", "RMSE S4 ★", "RMSE S3", "RMSE S1", "S4 − S3", "개선율 S4", "개선율 S3",
           "MBE S4", "MBE S3"]
    for j, h in enumerate(hdr, 1):
        if h:
            H(ws, r, j, h, size=9)
    r += 1
    r0 = r
    for rn, k in groups:
        m = f"({db(DB['rn'])}=$B{r})*({db(DB['k'])}=$C{r})"
        rm = lambda key: f"SQRT(SUMPRODUCT({m}*{db(DB[key])}^2)/$D{r})"
        C(ws, r, 1, f"{rn} D+{k}", fill=LIGHT); C(ws, r, 2, rn); C(ws, r, 3, k)
        C(ws, r, 4, f"=SUMPRODUCT({m})", "0")
        C(ws, r, 5, f"={rm('x4')}", F2); C(ws, r, 6, f"={rm('x3')}", F2); C(ws, r, 7, f"={rm('x1')}", F2)
        C(ws, r, 8, f"=E{r}-F{r}", "+0.00;-0.00;0.00")
        C(ws, r, 9, f"=1-E{r}/{rm('xp')}", PCT); C(ws, r, 10, f"=1-F{r}/{rm('xp')}", PCT)
        C(ws, r, 11, f"=SUMPRODUCT({m}*{db(DB['x4'])})/$D{r}", F2)
        C(ws, r, 12, f"=SUMPRODUCT({m}*{db(DB['x3'])})/$D{r}", F2)
        r += 1
    ws.conditional_formatting.add(f"H{r0}:H{r - 1}", FormulaRule(formula=[f"H{r0}<0"], fill=PatternFill("solid", fgColor=PASS_FILL)))
    ws.conditional_formatting.add(f"H{r0}:H{r - 1}", FormulaRule(formula=[f"H{r0}>0"], fill=PatternFill("solid", fgColor=FAIL_FILL)))
    return r


def _metric_table(ws, r, hdr, groups, db, fcol, ocol, ecol, pcol, mcol):
    for j, h in enumerate(hdr, 1):
        if h: H(ws, r, j, h, size=9)
    r += 1
    rows = []
    for rn, k in groups:
        crit = f'{db(DB["rn"])},$B{r},{db(DB["k"])},$C{r}'
        m = f"({db(DB['rn'])}=$B{r})*({db(DB['k'])}=$C{r})"
        f, o = db(fcol), db(ocol)
        C(ws, r, 1, f"{rn} D+{k}", fill=LIGHT)
        C(ws, r, 2, rn); C(ws, r, 3, k)
        C(ws, r, 4, f"=COUNTIFS({crit})", "0")
        C(ws, r, 5, f"=AVERAGEIFS({o},{crit})", F2)
        C(ws, r, 6, f"=AVERAGEIFS({f},{crit})", F2)
        C(ws, r, 7, f"=AVERAGEIFS({db(ecol)},{crit})", F2)
        C(ws, r, 8, f"=SQRT(SUMPRODUCT({m}*{db(ecol)}^2)/D{r})", F2)
        C(ws, r, 9, (f"=(D{r}*SUMPRODUCT({m}*{f}*{o})-SUMPRODUCT({m}*{f})*SUMPRODUCT({m}*{o}))^2"
                     f"/((D{r}*SUMPRODUCT({m}*{f}^2)-SUMPRODUCT({m}*{f})^2)*(D{r}*SUMPRODUCT({m}*{o}^2)-SUMPRODUCT({m}*{o})^2))"), F2)
        C(ws, r, 10, f"=SQRT(SUMPRODUCT({m}*{db(pcol)}^2)/D{r})", F2)
        C(ws, r, 11, f"=1-H{r}/J{r}", PCT)
        C(ws, r, 12, f"=SQRT(SUMPRODUCT({m}*{db(mcol)}^2)/D{r})", F2)
        C(ws, r, 13, f"=1-H{r}/L{r}", PCT)
        C(ws, r, 14, "○" if k in (1, 2, 3) else "참고")
        rows.append(r); r += 1
    return rows


# ── 방법 ────────────────────────────────────────────────────────────────
def sheet_method(wb, res):
    df, loc, sg = res["table"], res["check"]["location"], res.get("stn_grid")
    fmts = res["check"].get("formats") or ["portal"]
    api_like = any(f in fmts for f in ("openapi", "element"))
    bt = res.get("boot")
    if sg and loc == [sg]:
        grid_note = f"예보 격자 {sg} = ASOS {res['stn']} 관측소 격자. 5 km 격자 대표값과 지점 관측의 차이(대표성 오차)는 남음."
    else:
        grid_note = (f"예보 격자 {', '.join(loc)}가 ASOS {res['stn']} 관측소 격자({sg or '미등록'})와 다름 → "
                     f"대표성 오차가 달라질 수 있음.")
    if res.get("grid_given"):
        grid_note += " (파일에 격자 정보가 없어 실행 시 지정한 격자)"
    ws = wb.create_sheet("방법")
    T(ws, 1, 1, "방법과 정의 (근거: docs/THEORY.md 9장, 설계: docs/ARCHITECTURE.md 7장)", 12, True, GREEN)
    items = [
        ("자료", "예보", " ".join(t for f, t in (
            ("portal", "기상자료개방포털 '단기예보(격자)' 과거자료 CSV(요소별 파일): 발표일시 = CSV의 UTC 일·시 + 9시간, "
                       "여러 달을 한 파일로 받으면 'Start : YYYYMMDD' 행마다 연월을 바꿔 읽음."),
            ("openapi", "단기예보 조회서비스(OpenAPI) 응답 CSV(한 파일에 모든 요소): 발표·예보 일시는 KST 그대로."),
            ("element", "요소별 KST CSV(발표일·발표시각·예보일·예보시각·값, 요소마다 1파일): 일시는 KST 그대로, 값 표기는 OpenAPI와 같음. "
                        "요소는 값으로 판별, 격자 정보는 파일에 없어 실행 시 지정."),
            ("api", f"강수 문자열 '강수없음' 0 · '1mm 미만' {PCP_LT1_MM:g} · '30.0~50.0mm' {PCP_30_50_MM:g} · "
                    f"'50.0mm 이상' {PCP_GE50_MM:g} mm/h.")) if f in fmts or (f == "api" and api_like))
                        + " 필수 요소 TMX·TMN·TMP·REH·WSD·PCP, 선택 요소 SKY(하늘상태)·POP(강수확률)."),
        ("", "관측", "ASOS 일자료(01-Cycle cropwater_station.py 출력 워크북 원데이터 시트). 관측 ETo는 01-Cycle 계산과정 시트와 같은 식."),
        ("서비스", "발표", "아침 02시 발표 → 오늘(D+0)~D+3, 저녁 17시 발표 → 내일(D+1)~D+4. 검증에는 이 두 발표만 사용(다른 발표를 섞으면 성능이 부풀려짐)."),
        ("일 입력", "Tmax·Tmin", "해당 발표의 TMX(낮최고 09:01~18:00)·TMN(아침최저 03:01~09:00). ASOS 일최고·최저(0~24시)와 정의가 다름."),
        ("", "ea", "시간별 e°(TMP) × REH/100 의 일평균 (ASOS 평균증기압과 같은 개념)."),
        ("", "u2", "WSD 일평균(10 m) × 4.87/ln(67.8×10−5.42) = 0.748 × u10  [식47]"),
        ("", "강수", "PCP 일합계. 기준(설정, 1 mm) 이상이면 강수유무 1."),
        ("", "하늘상태", "SKY 코드 1 맑음·3 구름많음·4 흐림(그 밖의 값은 결측). 시각마다 태양고도의 사인값(밤 0)으로 가중해 낮 시간의 "
                      "구름많음 비율·흐림 비율(0~1)을 만듦 = 일사가 많은 한낮의 하늘상태가 더 크게 반영됨. 서비스 발표 자체에 SKY가 없으면 비워 둠. "
                      "강수확률(POP)은 같은 가중의 낮 평균(참고 값, Rs 추정에는 쓰지 않음)."),
        ("", "지나간 시각 채움", "발표 시점에 이미 지난 시각은 그 이전의 가장 최근 발표 값으로 채움. "
                                 + ("포털 과거자료는 발표 6시간 뒤부터 들어 있음(02시 발표 → 08시부터) → 아침 D+0의 00~07시 = 전날 17·20·23시 발표. "
                                    if "portal" in fmts else "")
                                 + ("OpenAPI 자료는 발표 1시간 뒤부터 들어 있음(02시 발표 → 03시부터) → 아침 D+0의 00~02시 = 전날 23시 발표"
                                    + (" (운영과 같음)." if "portal" not in fmts else ".") if api_like else
                                    "(API는 발표 1시간 뒤부터 제공되므로 운영에서는 00~02시만 채움)")),
        ("", "결측", "값이 ±900 이상(예: −999.9)이면 결측. 필수 6요소 중 하나라도 없는 서비스 발표는 통째로 제외(다른 발표로 대신하지 않음). "
                    "주 방법이 S4면 하늘상태가 없는 행, 그리고 대상일 관측이 없는 행도 제외. 3일 누적은 첫 3개 대상일이 모두 있을 때만."),
        ("", "마지막 날", "아침 D+3·저녁 D+4는 00시(1시간) + 03~21시(3시간 간격) 8개 시각. 풍속·강수는 코드값 → WSD 1: 같은 발표 직전 정량일 평균(최대 3.9), 2: 6.5, 3: 11 m/s / PCP 1: 1.5, 2: 9, 3: 20 mm/h × 3시간."),
        ("예보 ETo", "PM", "FAO-56 식(6), G = 0. 기압은 식(7) 고도 추정. Ra·Rso는 지점 위도·고도로 계산."),
        ("", "Rs", ("S4(★): Rs/Ra = a + b√(Tmax−Tmin) + c·강수유무 + d·구름많음 비율 + e·흐림 비율, [0.05Ra, Rso]로 제한. "
                    "계수는 선행일별로 '예보 입력 → 관측 Rs' 최소제곱. "
                    + (f"이 검증은 다른 해 운영 계수를 그대로 적용(계수 고정, {s4_source(res)}) — 독립 연도 검증. "
                       if res.get("s4_fixed") is not None else
                       "검증은 월 단위 교차검증(대상월을 뺀 나머지 달로 맞춘 계수, SKY계수 시트). ")
                    + f"비교용 S3: Rs/Ra = a + b√(Tmax−Tmin) + c·강수유무, 계수 {res['coef'].get('source', '')}. S1: kRs 0.16.")
                   if res.get("main") == "S4" else
                   f"S3: Rs/Ra = a + b√(Tmax−Tmin) + c·강수유무, [0.05Ra, Rso]로 제한. 계수는 검증 연도와 겹치지 않는 해의 관측으로 결정"
                   f"({res['coef'].get('source', '')}). 비교용 S1: kRs 0.16."),
        ("ETc", "Kc", "01-Cycle 워크북과 같은 Kc(시나리오 표값 + 식62·65 현지기상 보정, 식66 선형 보간). 예보·관측·기준선 모두 대상일 Kc를 곱함."),
        ("기준선", "지속성", "발표일 전날(가장 최근의 완결된 관측일) 관측 ETo를 모든 선행일에 사용. 02시에는 전날 자료가 아직 공개 전일 수 있어 실제 운영보다 기준선에 유리한 가정(판정에 보수적)."),
        ("", "7일평균", "발표일 전 7일(D−7~D−1) 관측 ETo 평균."),
        ("지표", "MBE·RMSE·R²", "오차 = 예보 − 관측. MBE +면 과대추정. R²는 예보·관측 상관계수의 제곱."),
        ("", "개선율", "1 − RMSE_예보 / RMSE_기준선. 0보다 크면 기준선보다 좋음."),
        ("", "3일 누적", "발표별 첫 3개 대상일 합계(아침 D+0~D+2, 저녁 D+1~D+3). FAO-56은 추정 Rs 기반 ETo를 여러 날 합계로 쓰도록 권고."),
        ("", "월별", "대상일의 월로 나눈 지표(월별 시트, 수식). 달마다 H2 기준을 적용한 결과는 원인 진단용 참고이며, H2 판정은 전체 기간으로 함."),
        ("", "불확실성", (f"{bt['block']}일 이동 블록 부트스트랩 {bt['n_boot']:,}회(오차분해 ④): 발표일을 {bt['block']}일 묶음으로 "
                        "복원추출해 D+1 RMSE·최소 개선율의 90% 구간과 두 기준을 모두 충족한 비율을 구함. 예보 오차는 날씨가 며칠 "
                        "이어져 서로 상관이 있으므로 하루 단위로 뽑지 않음. 교차검증한 보정(오차분해 ②)을 적용한 예보도 같은 방법으로 "
                        "계산(탐색, 판정 미사용).") if bt else "계산하지 않음(분석 생략)"),
        ("판정", "H2", "두 발표 각각 D+1~D+3 개선율 ≥ 기준, D+1 RMSE ≤ 상한 (요약 시트 노란 칸). 아침 D+0·저녁 D+4는 참고."),
        ("한계", "격자", grid_note),
        ("", "기간", f"대상일 {df.target.min():%Y-%m-%d} ~ {df.target.max():%Y-%m-%d}. "
                    + ("생육기(4~9월) 전체를 포함합니다(자료 공백으로 뺀 발표·행은 요약 시트 '제외' 행)."
                       if set(range(4, 10)) <= set(df.target.dt.month)
                       else "생육기(4~9월) 중 일부만 포함 → 나머지 달의 예보 자료로 확인.")),
        ("", "연도·지점", f"{'·'.join(str(y) for y in sorted(set(df.target.dt.year)))}년, ASOS {res['stn']} 한 지점의 결과. "
                          "다른 해·다른 지점에서의 성능은 검정하지 않음."),
        ("", "강수", "예보 강수는 강수유무(Rs 보정)에만 쓰였고, 유효강수·물수지 영향은 다음 단계(G4)에서 검증."),
    ]
    H(ws, 3, 1, "구분"); H(ws, 3, 2, "항목"); H(ws, 3, 3, "내용")
    r = 4
    for a, b, c in items:
        C(ws, r, 1, a, fill=(LIGHT if a else None), bold=bool(a)); C(ws, r, 2, b, left=True)
        C(ws, r, 3, c, left=True, wrap=True); ws.row_dimensions[r].height = 30 if len(c) > 70 else 18
        r += 1
    widths(ws, {"A": 10, "B": 16, "C": 120})
    return ws


# ── 빌드 ────────────────────────────────────────────────────────────────
def build_verify_workbook(res, out):
    df = res["table"]
    groups = [(rn, int(k)) for rn in ("아침", "저녁") for k in sorted(df[df.run_name == rn].lead_day.unique())]
    first_run = df.run.min().normalize()
    obs = res["obs"]
    obs = obs[(obs.date >= first_run - pd.Timedelta(days=7))].reset_index(drop=True)
    wb = Workbook()
    wb.active.title = "요약"
    sheet_settings(wb, res)
    # Rs계수는 관측 행 범위를 알아야 하므로, 관측 시트 행 범위를 먼저 계산
    obs_rows = (2, 1 + len(obs))
    sheet_rscoef(wb, res, obs_rows)
    _, obs_rows2 = sheet_obs(wb, obs)
    assert obs_rows2 == obs_rows
    main = res.get("main", "S3")
    if main == "S4":
        sheet_skycoef(wb, res)               # 일별비교의 S4 수식이 이 시트의 계수 범위를 참조
    _, db_rows, run_ids = sheet_daily(wb, df, obs_rows, sky=(main == "S4"))
    _, c3_rows = sheet_cum3(wb, df, db_rows, run_ids, res.get("cum3_runs"))
    sheet_diag(wb, db_rows, groups, main)
    sheet_attr(wb, res)
    if res.get("grid_cmp"):
        sheet_grid_compare(wb, res)
    _, chart_rows = sheet_chartdata(wb, df, db_rows, obs_rows)
    sheet_summary(wb, res, db_rows, c3_rows, groups, chart_rows)
    sheet_month(wb, df, db_rows, main)      # 요약 시트의 합격 기준 셀을 참조하므로 요약 다음에 만든다
    sheet_method(wb, res)
    order = ["요약", "일별비교", "3일누적", "월별", "입력진단", "오차분해"] + (["격자비교"] if res.get("grid_cmp") else []) + \
            ["Rs계수"] + (["SKY계수"] if main == "S4" else []) + ["관측", "설정", "방법", "차트자료"]
    wb._sheets = [wb[n] for n in order]
    for n in ("요약", "Rs계수", "SKY계수", "설정", "방법", "오차분해", "입력진단", "3일누적", "월별", "격자비교"):
        if n in wb.sheetnames:
            wb[n].sheet_view.showGridLines = False
    wb["요약"].sheet_properties.tabColor = GREEN
    for n in [x for x in ("요약", "월별", "Rs계수", "SKY계수", "입력진단", "오차분해", "격자비교", "방법", "설정") if x in wb.sheetnames]:
        w = wb[n]; w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    return safe_save(wb, out)
